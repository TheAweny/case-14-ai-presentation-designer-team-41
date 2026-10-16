"""Builds examples/template.pptx, the template of the example run config.

Plain python-pptx (MIT) on its own default theme: a cover, a section divider and the usual
content layouts, each with sample text, so the engine finds a composition on every slide. It
stands in for the template a person uploads; the service treats it like any other.

    uv run --no-project --with python-pptx python examples/make-template.py
"""

from pathlib import Path

from pptx import Presentation

SLIDES = [
    (0, ["Название презентации", "Подзаголовок: для кого и зачем"]),
    (2, ["Раздел", "Короткое пояснение к разделу"]),
    (1, ["Заголовок слайда", "Первый тезис\nВторой тезис\nТретий тезис"]),
    (3, ["Сравнение двух вариантов", "Что было и почему", "Что стало и почему"]),
    (4, ["Сравнение с подписями", "До", "Как было", "После", "Как стало"]),
    (7, ["Тезис с пояснением", "Главная мысль слайда в двух строках текста", "Подпись к тезису"]),
    (0, ["Спасибо", "Контакты и следующий шаг"]),
]


def main() -> None:
    # The default theme lays its placeholders out for 4:3, so the slides keep that size.
    deck = Presentation()
    for layout, texts in SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[layout])
        placeholders = sorted(slide.placeholders, key=lambda p: p.placeholder_format.idx)
        for placeholder, text in zip(placeholders, texts, strict=False):
            placeholder.text = text
    out = Path(__file__).with_name("template.pptx")
    deck.save(out)
    print(out)


if __name__ == "__main__":
    main()
