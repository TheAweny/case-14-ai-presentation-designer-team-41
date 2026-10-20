"""Rebuild CC0 brand assets: uv run --with reportlab python examples/brandbooks/generate.py."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).parent


def generate():
    tokens = {
        "colors": {
            "primary": {"$type": "color", "$value": "#1746A2", "role": "primary"},
            "accent": {"$type": "color", "$value": "#E85D3F", "role": "accent"},
            "background": {"$type": "color", "$value": "#F7F8FC", "role": "background"},
            "text": {"$type": "color", "$value": "#17213A", "role": "text"},
            "link": {"$type": "color", "$value": "{colors.primary}"},
        },
        "typography": {
            "heading": {"$type": "fontFamily", "$value": "Arial", "role": "heading"},
            "body": {"$type": "fontFamily", "$value": "Arial", "role": "body"},
            "headingSize": {"$type": "fontSize", "$value": "36pt", "role": "heading"},
            "bodySize": {"$type": "fontSize", "$value": "20pt", "role": "body"},
        },
        "spacing": {"gutter": {"$type": "spacing", "$value": {"value": 24, "unit": "px"}}},
    }
    (ROOT / "tokens.json").write_text(json.dumps(tokens, indent=2) + "\n")
    (ROOT / "symbol.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 120 120">'
        '<rect x="12" y="12" width="96" height="96" rx="24" fill="#1746A2"/>'
        '<circle cx="60" cy="60" r="24" fill="#E85D3F"/></svg>\n'
    )
    pdf = Canvas(str(ROOT / "guideline.pdf"), pagesize=(720, 480), invariant=1)
    pdf.setFillColor(HexColor("#F7F8FC"))
    pdf.rect(0, 0, 720, 480, stroke=0, fill=1)
    pdf.setFillColor(HexColor("#17213A"))
    pdf.setFont("Helvetica-Bold", 30)
    pdf.drawString(48, 416, "Northstar / Brand guide")
    pdf.setFont("Helvetica", 14)
    for y, line in zip(
        range(360, 100, -36),
        [
            "Primary color: HEX #1746A2",
            "Accent color: HEX #E85D3F",
            "Background color: HEX #F7F8FC",
            "Text color: HEX #17213A",
            "Font family: Arial",
            "Use the symbol only on a clear, high-contrast background.",
            "Declare placement in the slide brief; do not infer it from this sample.",
        ],
        strict=False,
    ):
        pdf.drawString(48, y, line)
    pdf.showPage()
    pdf.save()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F8FC")
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(10), Inches(2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Northstar: clear ideas, native slides"
    run.font.name, run.font.size = "Arial", Pt(36)
    run.font.color.rgb = RGBColor.from_string("1746A2")
    deck.save(str(ROOT / "slide-example.pptx"))
    with zipfile.ZipFile(ROOT / "northstar-brandbook.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        names = ["tokens.json", "symbol.svg", "guideline.pdf", "slide-example.pptx"]
        if (ROOT / "northstar.fig").is_file():
            names.append("northstar.fig")
        for name in names:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, (ROOT / name).read_bytes())


if __name__ == "__main__":
    generate()
