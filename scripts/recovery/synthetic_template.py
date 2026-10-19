"""Synthetic unfamiliar template for isolated integration runs; no customer material."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def create(destination: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
    layouts = [
        [(0.6, 0.5, 12, 1, 32), (0.6, 2, 5.6, 4.5, 20), (7, 2, 5.6, 4.5, 20)],
        [(0.6, 0.5, 12, 1, 32), (0.6, 2, 12, 2, 20), (0.6, 4.5, 12, 2, 20)],
        [(0.6, 0.6, 4, 5.8, 32), (5, 0.6, 7.6, 2.7, 20), (5, 3.9, 7.6, 2.7, 20)],
    ]
    for layout in layouts:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 247, 243)
        for index, (x, y, width, height, size) in enumerate(layout):
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            run = shape.text_frame.paragraphs[0].add_run()
            run.text = (
                "Исследование пилота"
                if index == 0
                else (
                    "Проверенный результат исследования описывает наблюдения команды, "
                    "ограничения пилота и следующий этап проверки устойчивости сервиса."
                )
            )
            run.font.name, run.font.size = "Play", Pt(size)
            run.font.color.rgb = RGBColor(32, 45, 60)
    destination.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(destination))


if __name__ == "__main__":
    import sys

    create(Path(sys.argv[1]))
