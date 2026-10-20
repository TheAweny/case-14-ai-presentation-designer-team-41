"""FastAPI application of the engine."""

import base64
import hashlib
import re
import subprocess
import tempfile
import time
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pydantic import BaseModel

from engine import __version__, checks, dataset, export, layout, parser, render, visuals
from engine.api._brandbook import effective_composition_request
from engine.api._composition_audit import composition_slides
from engine.api._profiles import UnknownProfileError, deterministic_profile
from engine.api._visuals import build_payload, capability_payload, compare_payload
from engine.contracts import (
    AuditReport,
    AuditRequest,
    DeckExportRequest,
    DeckExportResponse,
    DeckImportRequest,
    DeckImportResponse,
    DeckPlan,
    EditableDeckRequest,
    EditableDeckResponse,
    HealthCheck,
    LayoutVariantsRequest,
    LayoutVariantsResponse,
    MaterialExtractRequest,
    MaterialExtractResponse,
    SlideIR,
    Status,
    TemplatePreviewsRequest,
    TemplatePreviewsResponse,
    VisualBuildRequest,
    VisualBuildResponse,
    VisualCapability,
    VisualCompareRequest,
    VisualCompareResponse,
)
from engine.contracts._generated import (
    BrandbookParseRequest,
    BrandbookProfile,
    BrandbookTemplateRequest,
    BrandbookTemplateResponse,
    CompositionCheckResponse,
    CompositionExportRequest,
    CompositionExportResponse,
    DatasetExtractRequest,
    GenerationProfile,
)
from engine.export.brandbook import apply_brandbook
from engine.export.compositions import export_compositions
from engine.layout.text_fit import check_composition_text, fitting_scale
from engine.parser._errors import TemplateParseError
from engine.parser.brandbook import BrandbookError, parse_brandbook
from engine.parser.compositions import extract_compositions
from engine.parser.materials import MaterialError, extract_material

_STARTED_AT = time.monotonic()


def bundle_previews(previews: list[Path], output: Path) -> Path:
    """Write the rendered slides into one portable PNG archive."""
    if not previews:
        raise ValueError("Cannot create a PNG archive without rendered slides.")
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for preview in previews:
            archive.write(preview, preview.name)
    return output


class TemplateProfileRequest(BaseModel):
    template_path: str


app = FastAPI(
    title="Deck engine",
    version=__version__,
    description=(
        "Deterministic document engine: template parsing, layout, visuals, export, render and "
        "the deterministic audit checks. It never calls a model."
    ),
)


@app.get("/health", response_model=HealthCheck, summary="Liveness of the engine")
def health() -> HealthCheck:
    """Report that the engine process is up, using the shared health contract."""
    return HealthCheck(
        status=Status.ok,
        service="engine",
        version=__version__,
        uptimeSeconds=time.monotonic() - _STARTED_AT,
    )


@app.get(
    "/v1/visuals/capabilities",
    response_model=VisualCapability,
    summary="What the deterministic visual layer can place",
)
def visual_capabilities() -> VisualCapability:
    """Capability adapter for a future semantic planner. This process does not call a model."""
    return VisualCapability.model_validate(capability_payload())


@app.post(
    "/v1/visuals",
    response_model=VisualBuildResponse,
    response_model_exclude_unset=True,
    summary="Place a KPI block, chart or table as SlideIR",
)
def visual_build(request: VisualBuildRequest) -> VisualBuildResponse:
    """Build native SlideIR for one visual from structured data and a template profile."""
    try:
        payload = build_payload(request)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return VisualBuildResponse.model_validate(payload)


@app.post(
    "/v1/visuals/compare",
    response_model=VisualCompareResponse,
    summary="Compare structural signatures of visual variants",
)
def visual_compare(request: VisualCompareRequest) -> VisualCompareResponse:
    """Three variants must differ. A cosmetic control must not count as a new design."""
    return VisualCompareResponse.model_validate(compare_payload(request))


@app.post(
    "/v1/template-profile",
    response_model=dict,
    summary="Parse a saved PPTX template and return its profile",
)
def template_profile(request: TemplateProfileRequest) -> dict:
    """Return the deterministic parser output for a saved template file."""
    path = Path(request.template_path).resolve()
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Template PPTX was not found.")
    try:
        design_system = parser.design_system_of(path)
        # Use the public parser entrypoint so the canonical IR digest is included in the payload.
        template_ir = parser.parse_template(path)
    except (OSError, ValueError, TemplateParseError) as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    preview_data_url: str | None = None
    try:
        with tempfile.TemporaryDirectory(prefix="template-preview-") as directory:
            rendered = render.render_to_png(
                path, Path(directory), dpi=96, first_slide_only=True, timeout_seconds=20
            )
            if rendered:
                preview_data_url = "data:image/png;base64," + base64.b64encode(
                    rendered[0].read_bytes()
                ).decode("ascii")
    except (OSError, subprocess.SubprocessError, ValueError):
        # A profile remains useful when a headless renderer is unavailable.
        preview_data_url = None
    return {
        "designSystem": design_system,
        "templateIr": template_ir,
        "previewDataUrl": preview_data_url,
    }


@app.post(
    "/v1/dataset-extract",
    response_model=dataset.Dataset,
    summary="Read the tables of a data file with their cell references and diagnostics",
)
def dataset_extract(request: DatasetExtractRequest) -> dataset.Dataset:
    """Return every table of a CSV, TSV, XLSX, JSON or text-table file, typed and unabridged.

    Values keep the text the file held; a value is derived only where the text allows one reading,
    and what is ambiguous, unreadable or too large is said in `diagnostics`. No model is called.
    """
    path = Path(request.path).resolve()
    if not path.is_file():
        raise HTTPException(status_code=404, detail="File not found.")
    return dataset.extract_dataset(path)


@app.post(
    "/v1/material-extract",
    response_model=MaterialExtractResponse,
    summary="Read a slice of an uploaded text or spreadsheet file",
)
def material_extract(request: MaterialExtractRequest) -> MaterialExtractResponse:
    """Return characters or cells from a saved document. The engine does not call a model."""
    try:
        payload = extract_material(request)
    except MaterialError as error:
        raise HTTPException(status_code=error.status, detail=error.detail) from error
    return MaterialExtractResponse.model_validate(payload)


#: English Metric Units per inch: how OOXML measures a slide.
EMU_PER_INCH = 914400


@app.post(
    "/v1/template-previews",
    response_model=TemplatePreviewsResponse,
    summary="Draw every slide of a saved template as a PNG of the requested width",
)
def template_previews(request: TemplatePreviewsRequest) -> TemplatePreviewsResponse:
    """Pictures of the template's sample slides, for the library to show them as they are."""
    path = Path(request.templatePath).resolve()
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Template PPTX was not found.")
    try:
        slide_width = Presentation(str(path)).slide_width or EMU_PER_INCH * 10
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        raise HTTPException(status_code=422, detail=f"PPTX cannot be read: {error}") from error
    dpi = max(1, round(request.widthPx * EMU_PER_INCH / slide_width))
    try:
        images = render.render_to_png(
            path, Path(request.outputDir).resolve(), dpi=dpi, timeout_seconds=900
        )
    except (OSError, subprocess.SubprocessError, ValueError) as error:
        raise HTTPException(
            status_code=503, detail=f"The template could not be drawn: {error}"
        ) from error
    return TemplatePreviewsResponse.model_validate(
        {"schemaVersion": 1, "previews": [image.name for image in images]}
    )


@app.post(
    "/v1/deck-import",
    response_model=DeckImportResponse,
    # SlideIR tells an absent field from a null one; unset optional fields must stay absent.
    response_model_exclude_unset=True,
    summary="Read a finished PPTX back as SlideIR, the way its slides are drawn",
)
def deck_import(request: DeckImportRequest) -> DeckImportResponse:
    """Every slide of a deck as editable elements, with the pictures it draws written out.

    The editor opens a generated deck from this, so it shows what the file and its PDF show:
    what the master and layout put on each slide, text in the size, colour and typeface it is
    drawn in, the picture behind it and the outline of every shape.
    """
    path = Path(request.pptxPath).resolve()
    if not path.is_file():
        raise HTTPException(status_code=422, detail=f"PPTX was not found: {request.pptxPath}")
    media_dir = Path(request.mediaDir).resolve() if request.mediaDir else None
    try:
        slides = parser.import_deck(path, media_dir)
    except TemplateParseError as error:
        raise HTTPException(status_code=422, detail=f"PPTX cannot be read: {error}") from error
    return DeckImportResponse.model_validate({"schemaVersion": 1, "slides": slides})


@app.post(
    "/v1/editable-decks",
    response_model=EditableDeckResponse,
    summary="Build one editable PPTX from a local template and fixture content",
)
def build_editable_deck(request: EditableDeckRequest) -> EditableDeckResponse:
    """Join deterministic parser, layout and export layers for the MVP vertical slice."""
    template_path = Path(request.templatePath).resolve()
    if not template_path.is_file():
        raise HTTPException(status_code=404, detail="Template PPTX was not found.")

    design_system, patterns = parser.template_profile_of(template_path)
    plan = DeckPlan.model_validate(
        {
            "schemaVersion": 1,
            "planId": f"plan-{request.runId}",
            "briefId": request.runId,
            "title": request.slides[0].title,
            "language": "ru",
            "slides": [
                {
                    "slideId": slide.slideId,
                    "headline": slide.title,
                    "subhead": "",
                    "intent": "statement",
                    "blocks": [{"kind": "paragraph", "text": slide.body}],
                    "speakerNotes": "",
                    "sources": [],
                }
                for index, slide in enumerate(request.slides, start=1)
            ],
        }
    )
    output_path = Path(request.outputPath).resolve()
    slide_ir = _styled(layout.build_slide_ir(plan, design_system, patterns), design_system)
    export.export_pptx(slide_ir, design_system, template_path, output_path, visuals.draw_visual)
    return EditableDeckResponse.model_validate(
        {
            "schemaVersion": 1,
            "outputPath": str(output_path),
            "designSystem": design_system,
            "slideCount": len(slide_ir),
        }
    )


@app.post(
    "/v1/generation-profile",
    response_model=GenerationProfile,
    # Optional slot fields such as prompt stay absent when unset: TypeScript rejects a null.
    response_model_exclude_unset=True,
)
def generation_profile(request: TemplateProfileRequest) -> dict:
    """Prepare exact reusable compositions before the timed generation request."""
    path = Path(request.template_path).resolve()
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Template not found.")
    compositions = extract_compositions(path)
    requested = sorted(
        {slot["font"] for pattern in compositions for slot in pattern["slots"] if slot["font"]}
    )
    try:
        missing = render.missing_fonts(set(requested))
        fonts = {
            "state": "missing" if missing else "ready",
            "requested": requested,
            "missing": missing,
        }
    except (OSError, subprocess.SubprocessError):
        fonts = {"state": "unknown", "requested": requested, "missing": []}
    return {
        "schemaVersion": 1,
        "templateHash": hashlib.sha256(path.read_bytes()).hexdigest(),
        "designSystem": parser.design_system_of(path),
        "compositions": compositions,
        "fontReadiness": fonts,
    }


@app.post("/v1/brandbook-profile", response_model=BrandbookProfile)
def brandbook_profile(request: BrandbookParseRequest) -> dict:
    """I inventory and parse saved brand sources without interpreting document instructions."""
    try:
        return parse_brandbook(Path(request.path).resolve(), Path(request.outputDir).resolve())
    except BrandbookError as error:
        raise HTTPException(status_code=error.status, detail=error.detail) from error


@app.post(
    "/v1/brandbook-template",
    response_model=BrandbookTemplateResponse,
    response_model_exclude_unset=True,
)
def brandbook_template(request: BrandbookTemplateRequest) -> dict:
    """I prepare native editable styles before profiling, planning and text measurement."""
    source = Path(request.templatePath).resolve()
    target = Path(request.outputPath).resolve()
    if not source.is_file():
        raise HTTPException(status_code=404, detail="Template not found.")
    if source == target:
        raise HTTPException(
            status_code=422, detail="Use a distinct effective template output path."
        )
    try:
        warnings = apply_brandbook(source, target, request.brandbook.model_dump(mode="json"))
        profile = generation_profile(TemplateProfileRequest(template_path=str(target)))
    except (ValueError, OSError, zipfile.BadZipFile) as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return {
        "schemaVersion": 1,
        "templatePath": str(target),
        "profile": profile,
        "designSystem": profile["designSystem"],
        "warnings": warnings,
    }


@app.post("/v1/composition-check", response_model=CompositionCheckResponse)
def composition_check(request: CompositionExportRequest) -> dict:
    """Check every binding against the source typography before writing a file."""
    if request.brandbook is not None:
        with effective_composition_request(request) as effective:
            return composition_check(effective)
    source = Path(request.templatePath).resolve()
    if not source.is_file():
        raise HTTPException(status_code=404, detail="Template not found.")
    patterns = {p["patternId"]: p for p in extract_compositions(source)}
    findings = []
    for index, slide in enumerate(request.slides):
        pattern = patterns.get(slide.patternId)
        if pattern is None:
            raise HTTPException(status_code=422, detail="Unknown template composition.")
        texts = [t.model_dump() for t in slide.texts]
        if sorted(t["slotId"] for t in texts) != sorted(s["slotId"] for s in pattern["slots"]):
            raise HTTPException(status_code=422, detail="Bind every template slot exactly once.")
        findings.extend({**f, "slideIndex": index} for f in check_composition_text(pattern, texts))
    return {"schemaVersion": 1, "findings": findings}


@app.post(
    "/v1/composition-decks",
    response_model=CompositionExportResponse,
    # TemplateIR tells an absent paragraph field from a null one. Pydantic defaults such as
    # bullet and marginEmu must stay absent, or the TypeScript schema rejects the export.
    response_model_exclude_unset=True,
)
def composition_deck(request: CompositionExportRequest) -> dict:
    """Join independent document layers for a template-faithful editable export.

    A template family that is not installed does not stop the export. The file keeps that
    family name; the preview may draw a substitute, and the response names the gap.
    """
    if request.brandbook is not None:
        with effective_composition_request(request) as effective:
            return composition_deck(effective)
    source = Path(request.templatePath).resolve()
    output = Path(request.outputPath).resolve()
    findings = composition_check(request)["findings"]
    if findings and not getattr(request, "allowOverflow", False):
        raise HTTPException(status_code=422, detail=findings)
    missing: list[str] = []
    try:
        patterns = extract_compositions(source)
        selected = {slide.patternId for slide in request.slides}
        missing = render.missing_fonts(
            {
                slot["font"]
                for pattern in patterns
                if pattern["patternId"] in selected
                for slot in pattern["slots"]
            }
        )
        design_system = parser.design_system_of(source)
        # Copy the rewriter could not shorten is set smaller until it fits its box: spilling over
        # the shapes below it would cost the whole slide, a smaller font costs only that slot.
        by_id = {pattern["patternId"]: pattern for pattern in patterns}
        font_scales: dict[tuple[int, str], float] = {}
        for finding in findings if getattr(request, "allowOverflow", False) else []:
            slide = request.slides[finding["slideIndex"]]
            slot = next(
                s for s in by_id[slide.patternId]["slots"] if s["slotId"] == finding["slotId"]
            )
            text = next(t.text for t in slide.texts if t.slotId == finding["slotId"])
            font_scales[(finding["slideIndex"], finding["slotId"])] = fitting_scale(text, slot)

        def draw(shapes: Any, element: dict[str, Any]) -> None:
            if element["kind"] == "native":
                built = visuals.build_visual(
                    {
                        "schemaVersion": 1,
                        "designSystem": design_system,
                        "variant": element.get("variant") or "balanced",
                        "elementId": f"visual-{element['slotId']}",
                        "box": element["box"],
                        "language": "ru",
                        "block": element["block"],
                    }
                )
                if not built["fit"]["fits"]:
                    raise ValueError(f"Native visual overflow: {built['fit']['overflow']}")
                for native in built["elements"]:
                    export.draw_native_element(shapes, native, design_system, visuals.draw_visual)
                return
            styled = visuals.style_visuals([{"elements": [element]}], design_system)
            visuals.draw_visual(shapes, styled[0]["elements"][0], design_system)

        export_compositions(
            source,
            output,
            [s.model_dump(mode="json") for s in request.slides],
            patterns,
            draw,
            allow_overflow=getattr(request, "allowOverflow", False),
            font_scales=font_scales,
        )
        previews = render.render_to_png(output, output.parent / "previews", dpi=110)
        ir = parser.parse_template(output)
        with zipfile.ZipFile(output) as archive:
            media_dir = output.parent / "media"
            media_dir.mkdir(exist_ok=True)
            for entry in archive.namelist():
                if entry.startswith("ppt/media/") and not entry.endswith("/"):
                    (media_dir / Path(entry).name).write_bytes(archive.read(entry))
        bundle_previews(previews, output.parent / "slides.zip")
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    warnings: list[str] = []
    if missing:
        warnings.append(
            "Fonts not installed on the render server: "
            f"{', '.join(missing)}. "
            "The .pptx keeps the template fonts; previews use substitutes."
        )
    return {
        "schemaVersion": 1,
        "outputPath": str(output),
        "templateIr": ir,
        "previews": [str(p) for p in previews],
        "pdfPath": str(output.parent / "previews" / f"{output.stem}.pdf"),
        "warnings": warnings,
        "missingFonts": missing,
    }


def _media_sizes(slides: list[dict[str, Any]], stated: dict[str, Any]) -> dict[str, Any]:
    """Pixel sizes of the pictures: those the caller stated, then files the engine can open."""
    sizes = dict(stated)
    for slide in slides:
        for element in slide["elements"]:
            part = element.get("mediaPart")
            if element["kind"] != "image" or not part or part in sizes:
                continue
            path = Path(part)
            if not path.is_file():
                continue
            try:
                with Image.open(path) as picture:
                    sizes[part] = {"width": picture.width, "height": picture.height}
            except (OSError, UnidentifiedImageError):
                continue
    return sizes


@app.post(
    "/v1/audit",
    response_model=AuditReport,
    # SlideIR tells an absent field from a null one; unset optional fields must stay absent.
    response_model_exclude_unset=True,
    summary="Verify generated slides against the template and the audit profile",
)
def audit(request: AuditRequest) -> AuditReport:
    """Measure every slide of a generated deck and report what breaks the requirements.

    The input is whatever a generator produced, as SlideIR, with the design system of the template
    it was asked to follow. Nothing in it is trusted: layout, template fidelity and density are
    measured here, and each finding says what is wrong, why, and how to fix it when that can be
    computed.
    """
    payload = request.model_dump(mode="json", exclude_unset=True)
    profile_name = request.profile or "default"
    try:
        profile = deterministic_profile(profile_name)
    except UnknownProfileError as error:
        raise HTTPException(status_code=422, detail=f"Unknown audit profile {error}.") from error
    payload["mediaSizes"] = _media_sizes(payload["slides"], payload.get("mediaSizes") or {})
    return AuditReport.model_validate(checks.audit_deck(payload, profile_name, profile))


@app.post(
    "/v1/layouts",
    response_model=LayoutVariantsResponse,
    # SlideIR tells an absent field from a null one; unset optional fields must stay absent.
    response_model_exclude_unset=True,
    summary="Lay one deck plan out as the three variants on a local template",
)
def build_layouts(request: LayoutVariantsRequest) -> LayoutVariantsResponse:
    """Build the requested variants of one plan against the profile of one template.

    The plan is reused as it is: every variant carries the same words, and the variants differ
    in the layouts they take from the template and in how copy sits in them.
    """
    template_path = Path(request.templatePath).resolve()
    if not template_path.is_file():
        raise HTTPException(status_code=404, detail="Template PPTX was not found.")
    design_system, patterns = parser.template_profile_of(template_path)
    plan = DeckPlan.model_validate(request.plan.model_dump(mode="json"))
    wanted = [str(getattr(v, "value", v)) for v in request.variants or layout.VARIANTS]
    variants = [
        {
            "variant": variant,
            "slides": [
                slide.model_dump(mode="json", exclude_unset=True)
                for slide in _styled(
                    layout.build_slide_ir(plan, design_system, patterns, variant), design_system
                )
            ],
        }
        for variant in layout.VARIANTS
        if variant in wanted
    ]
    return LayoutVariantsResponse.model_validate(
        {
            "schemaVersion": 1,
            "designSystem": design_system,
            "patternCatalog": patterns,
            "variants": variants,
        }
    )


def _styled(slides: list[SlideIR], design_system: dict[str, Any]) -> list[SlideIR]:
    """Slides with their tables and charts given the template's look by the visuals layer."""
    dumped = [slide.model_dump(mode="json", exclude_unset=True) for slide in slides]
    return [SlideIR.model_validate(s) for s in visuals.style_visuals(dumped, design_system)]


def _export_media(request: DeckExportRequest) -> dict[str, bytes]:
    """Read only referenced uploads contained in this export's own media directory."""
    entries = request.media or []
    referenced = {
        element.mediaPart
        for slide in request.slides
        for element in slide.elements
        if element.kind == "image" and element.mediaPart.startswith("upload/")
    }
    keys = [entry.mediaPart for entry in entries]
    if (
        len(entries) > 32
        or len(set(keys)) != len(keys)
        or set(keys) != referenced
        or any(re.fullmatch(r"upload/[\w.-]+\.(?:png|jpe?g)", key) is None for key in keys)
        or any(".." in key for key in keys)
    ):
        raise ValueError("Export media must bind each referenced upload part exactly once.")
    if not entries:
        return {}
    output = Path(request.outputDir).resolve()
    media = output / "media"
    if media.resolve(strict=True) != media:
        raise ValueError("Export media directory must be under the output directory.")
    files = []
    total = 0
    # Validate every path and size before opening any file for its contents.
    for entry in entries:
        file = Path(entry.path).resolve(strict=True)
        if file == media or not file.is_relative_to(media) or not file.is_file():
            raise ValueError("Export media file must be under the output media directory.")
        size = file.stat().st_size
        total += size
        if not 0 < size <= 1024 * 1024 or total > 3 * 1024 * 1024:
            raise ValueError("Export media exceeds 1 MiB per file or 3 MiB total.")
        files.append((entry.mediaPart, file))
    result = {}
    total = 0
    for key, file in files:
        with file.open("rb") as stream:
            blob = stream.read(1024 * 1024 + 1)
        total += len(blob)
        if not 0 < len(blob) <= 1024 * 1024 or total > 3 * 1024 * 1024:
            raise ValueError("Export media exceeds 1 MiB per file or 3 MiB total.")
        expected = "PNG" if key.endswith(".png") else "JPEG"
        signature = (
            blob.startswith(b"\x89PNG\r\n\x1a\n")
            if expected == "PNG"
            else (blob.startswith(b"\xff\xd8\xff") and blob.endswith(b"\xff\xd9"))
        )
        if not signature:
            raise ValueError("Export media must have a matching PNG or JPEG signature.")
        with Image.open(BytesIO(blob)) as picture:
            if (
                picture.format != expected
                or not 1 <= picture.width <= 8192
                or not 1 <= picture.height <= 8192
                or picture.width * picture.height > 40_000_000
            ):
                raise ValueError("Export media has an invalid format or dimensions.")
            picture.verify()
        with Image.open(BytesIO(blob)) as picture:
            picture.load()
        result[key] = blob
    return result


@app.post(
    "/v1/exports",
    response_model=DeckExportResponse,
    # SlideIR tells an absent field from a null one; unset optional fields must stay absent.
    response_model_exclude_unset=True,
    summary="Write slides onto their template as native PPTX, and PDF from it",
)
def export_deck(request: DeckExportRequest) -> DeckExportResponse:
    """Write one variant's slides as files, tables and charts as native objects."""
    # Keep every downstream operation (styling, fit, native export, render) on the same
    # physically branded template. Older generated contracts simply have no brandbook attribute.
    binding = getattr(request, "brandbook", None)
    if binding is not None:
        with tempfile.TemporaryDirectory(prefix="brandbook-export-") as directory:
            effective = Path(directory) / "template.pptx"
            apply_brandbook(
                Path(request.templatePath).resolve(),
                effective,
                binding.model_dump(mode="json") if hasattr(binding, "model_dump") else binding,
            )
            return export_deck(
                request.model_copy(update={"templatePath": str(effective), "brandbook": None})
            )
    template_path = Path(request.templatePath).resolve()
    if not template_path.is_file():
        raise HTTPException(status_code=404, detail="Template PPTX was not found.")
    try:
        external_media = _export_media(request)
    except (OSError, ValueError, SyntaxError, Image.DecompressionBombError) as error:
        raise HTTPException(status_code=422, detail=f"Invalid export media: {error}") from error
    design_system, _ = parser.template_profile_of(template_path)
    variant = str(getattr(request.variant, "value", request.variant))
    slides = [
        SlideIR.model_validate(slide.model_dump(mode="json", exclude_unset=True))
        for slide in request.slides
    ]
    formats = [str(getattr(f, "value", f)) for f in request.formats or ["pptx", "pdf"]]
    # A PPTX download does not need PNG previews, and a PDF download needs no second pass:
    # the API sets this per requested format so an export pays only for what it returns.
    previews = request.previews is not False
    output_dir = Path(request.outputDir).resolve()
    pptx_path = export.export_pptx(
        _styled(slides, design_system),
        design_system,
        template_path,
        output_dir / f"{variant}.pptx",
        visuals.draw_visual,
        external_media=external_media,
    )
    files = [{"format": "pptx", "path": str(pptx_path)}] if "pptx" in formats else []
    if "pdf" in formats:
        try:
            pdf_path = export.export_pdf(pptx_path, output_dir / f"{variant}.pdf")
        except RuntimeError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
        files.append({"format": "pdf", "path": str(pdf_path)})
        if previews:
            # The previews reuse the PDF the download already paid for: one LibreOffice pass.
            try:
                render.rasterize_pdf(pdf_path, output_dir / "previews", dpi=110)
            except (ValueError, FileNotFoundError, subprocess.SubprocessError) as error:
                raise HTTPException(status_code=503, detail=str(error)) from error
    elif previews:
        # HTML and PNG downloads are built from the previews; the PDF is only the render stage.
        try:
            render.render_to_png(pptx_path, output_dir / "previews", dpi=110)
        except (ValueError, FileNotFoundError, subprocess.SubprocessError) as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
    return DeckExportResponse.model_validate(
        {"schemaVersion": 1, "variant": variant, "files": files}
    )


@app.post(
    "/v1/composition-audit",
    response_model=AuditReport,
    # SlideIR tells an absent field from a null one; unset optional fields must stay absent.
    response_model_exclude_unset=True,
)
def composition_audit(request: CompositionExportRequest) -> AuditReport:
    """Audit a deck exported from compositions, measuring the written file.

    Only the slots are judged: the rest of each slide is the template's own design. The
    `composition` profile leaves out what the typographic fit check already covers exactly and
    what would flag the template's own geometry.
    """
    if request.brandbook is not None:
        with effective_composition_request(request) as effective:
            return composition_audit(effective)
    source = Path(request.templatePath).resolve()
    output = Path(request.outputPath).resolve()
    if not source.is_file() or not output.is_file():
        raise HTTPException(status_code=404, detail="Template or exported deck not found.")
    patterns = {p["patternId"]: p for p in extract_compositions(source)}
    slides = composition_slides(
        parser.parse_template(output), [s.model_dump(mode="json") for s in request.slides], patterns
    )
    # Native blocks can create several text/chart/table shapes. Read those from
    # the final file rather than reconstructing them from the requested values.
    if any(s.visual is not None and s.visual.kind == "native" for s in request.slides):
        actual = parser.import_deck(output)
        for index, composed in enumerate(request.slides):
            if composed.visual is None or composed.visual.kind != "native":
                continue
            slots = patterns[composed.patternId]["slots"]
            bounds = [slot["box"] for slot in slots if slot.get("box")]
            read = actual[index]
            read["slideId"] = slides[index]["slideId"]
            read["patternId"] = composed.patternId
            for element in read["elements"]:
                box = element["box"]
                centre = (box["x"] + box["cx"] / 2, box["y"] + box["cy"] / 2)
                if not any(
                    b["x"] <= centre[0] <= b["x"] + b["cx"]
                    and b["y"] <= centre[1] <= b["y"] + b["cy"]
                    for b in bounds
                ):
                    element["role"] = "decoration"
            slides[index] = read
    payload: dict[str, Any] = {
        "schemaVersion": 1,
        "deckId": output.parent.name,
        "slides": slides,
        # File-backed checks must inspect the actual exported media, not the requested image
        # or its full-frame average. Internal orchestration data is not a wire contract.
        "pptxPath": str(output),
        # An empty slot's shape is taken off the slide at export, so only filled ones remain.
        "editableTextShapeIds": [
            [
                binding.slotId
                for binding in slide.texts
                if binding.text.strip()
                and (slide.visual is None or binding.slotId != slide.visual.slotId)
            ]
            for slide in request.slides
        ],
        "designSystem": parser.design_system_of(source),
        "patternCatalog": None,
        # What each slot showed in the template, so copy left from the sample is told apart.
        "slotSamples": {
            f"s{index + 1}": {
                f"slot-{slot['slotId']}": slot["sample"]
                for slot in patterns.get(slide.patternId, {"slots": []})["slots"]
            }
            for index, slide in enumerate(request.slides)
        },
    }
    if request.today is not None:
        payload["today"] = request.today.isoformat()
    if request.briefYears is not None:
        payload["knownYears"] = request.briefYears
    try:
        profile = deterministic_profile("composition")
    except UnknownProfileError:
        profile = {}
    report = checks.audit_deck(payload, "composition", profile)
    # The file check names OOXML shapes; composition SlideIR names those same editable
    # shapes slot-<id>. Preserve an actionable target for clients consuming this endpoint.
    for finding in report["findings"]:
        if finding["checkId"] == "image-text-readability":
            element_id = finding["target"].get("elementId")
            if element_id:
                finding["target"]["elementId"] = f"slot-{element_id.rsplit('-', 1)[-1]}"
    return AuditReport.model_validate(report)
