"""Visuals layer: native charts, tables and figures, styled from the template.

The layout places a table or a chart as data in a box. This layer gives it the template's look,
in palette roles so a deck stays on brand when the palette is edited, and draws it into a
presentation as native PowerPoint objects: a real table and a real chart with its data sheet,
never a picture of one, so both stay editable.
"""

from itertools import cycle
from typing import Any, cast

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.shapes.shapetree import SlideShapes
from pptx.util import Emu, Pt

from engine.visuals.build import build_visual, capabilities
from engine.visuals.contact import contact_sheet_svg
from engine.visuals.signature import compare_subjects, structural_signature

__all__ = [
    "build_visual",
    "capabilities",
    "compare_subjects",
    "contact_sheet_svg",
    "draw_visual",
    "structural_signature",
    "style_visuals",
]

_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

#: Palette roles a table takes its look from, when the template has them.
_HEADER_FILL = "accent"
_HEADER_TEXT = "text-on-accent"
_BAND_FILL = "surface"
_RULE_COLORS = ("line", "text-muted")

#: Series take the palette's colours in this order, skipping the roles a template lacks.
_SERIES_ROLES = ("accent", "accent-alt", "text", "text-muted", "line", "positive", "negative")

#: Tables with more body rows than this are banded, so a row can be followed across.
_BAND_FROM_ROWS = 3

_HAIRLINE_EMU = 9_525

#: WCAG contrast body text keeps against a band fill, the same floor the audit applies.
_MIN_CONTRAST = 4.5

_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
}


def style_visuals(
    slides: list[dict[str, Any]], design_system: dict[str, Any]
) -> list[dict[str, Any]]:
    """Give every table of a deck the template's look, stated in palette roles.

    Args:
        slides: Slide IR in its JSON form.
        design_system: Design system of the template the deck follows.

    Returns:
        The same slides, tables styled. Nothing else is touched, and a look a user already set
        on a table is kept.
    """
    roles = {entry["role"] for entry in design_system["palette"]}
    palette = _palette(design_system)
    text = palette.get(_body_style(design_system)["colorRole"]) or palette.get("text")
    if _BAND_FILL in palette and text and _contrast(text, palette[_BAND_FILL]) < _MIN_CONTRAST:
        # Banding with a fill the body text cannot be read on would hide every other row.
        roles.discard(_BAND_FILL)
    return [
        {
            **slide,
            "elements": [
                _style_table(element, roles) if element["kind"] == "table" else element
                for element in slide["elements"]
            ],
        }
        for slide in slides
    ]


def _contrast(a: str, b: str) -> float:
    def luminance(hex_colour: str) -> float:
        channels = [int(hex_colour[i : i + 2], 16) / 255 for i in (0, 2, 4)]
        lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
        return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]

    la, lb = luminance(a), luminance(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def _style_table(element: dict[str, Any], roles: set[str]) -> dict[str, Any]:
    look: dict[str, Any] = {"headerRow": True}
    if len(element["rows"]) > _BAND_FROM_ROWS and _BAND_FILL in roles:
        look |= {"bandedRows": True, "bandFill": _BAND_FILL}
    rule = next((role for role in _RULE_COLORS if role in roles), None)
    if rule is not None:
        look["borders"] = {
            "color": rule,
            "widthEmu": _HAIRLINE_EMU,
            "horizontal": True,
            "vertical": False,
            "outer": False,
        }
    if {_HEADER_FILL, _HEADER_TEXT} <= roles:
        header = {"fill": _HEADER_FILL, "color": _HEADER_TEXT, "bold": True}
        look["cellStyles"] = [[header for _ in element["headers"]]]
    return {**look, **element}


def draw_visual(
    shapes: SlideShapes, element: dict[str, Any], design_system: dict[str, Any]
) -> None:
    """Draw one table or chart of the slide IR onto a slide as a native object.

    Args:
        shapes: Shape tree of the slide being written.
        element: A table or chart element in its JSON form. Other kinds are ignored.
        design_system: Design system that resolves palette roles and supplies the body font.
    """
    if element["kind"] == "table":
        _draw_table(shapes, element, design_system)
    elif element["kind"] == "chart":
        _draw_chart(shapes, element, design_system)


def _palette(design_system: dict[str, Any]) -> dict[str, str]:
    return {entry["role"]: entry["hex"] for entry in design_system["palette"]}


def _resolve(value: str | None, palette: dict[str, str]) -> RGBColor | None:
    if not value:
        return None
    hex_value = palette.get(value, value)
    try:
        return RGBColor.from_string(hex_value)
    except ValueError:
        return None


def _body_style(design_system: dict[str, Any]) -> dict[str, Any]:
    scale = {style["role"]: style for style in design_system["typeScale"]}
    return scale.get("body") or scale.get("bullet") or design_system["typeScale"][0]


def _draw_table(
    shapes: SlideShapes, element: dict[str, Any], design_system: dict[str, Any]
) -> None:
    box = element["box"]
    headers, rows = element["headers"], element["rows"]
    frame = shapes.add_table(
        len(rows) + 1, len(headers), Emu(box["x"]), Emu(box["y"]), Emu(box["cx"]), Emu(box["cy"])
    )
    table = frame.table
    # add_table applies PowerPoint's stock table style, whose own banding would paint every row
    # we leave unfilled in a colour that is not the template's; fills are set cell by cell below.
    table.horz_banding = False
    table.first_row = True
    palette = _palette(design_system)
    style = _body_style(design_system)
    # Rows as tall as two lines of body text rather than stretched over the whole box: a short
    # table in a tall slot reads as a table, not as a grid of empty cells.
    row_height = min(box["cy"] // (len(rows) + 1), int(float(style["sizePt"]) * 2.4 * 12_700))
    for row in table.rows:
        row.height = Emu(row_height)
    styles = element.get("cellStyles") or []
    header_style = styles[0][0] if styles and styles[0] and styles[0][0] else {}
    for r, values in enumerate([headers, *rows]):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = str(value)
            run = cell.text_frame.paragraphs[0].runs[0] if cell.text else None
            fill = _resolve(header_style.get("fill"), palette) if r == 0 else None
            if r > 0 and element.get("bandedRows") and r % 2 == 0:
                fill = _resolve(element.get("bandFill"), palette)
            if fill is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
            elif r > 0:
                cell.fill.background()
            own = styles[r][c] if r < len(styles) and c < len(styles[r]) else None
            align = own.get("align") if own else None
            if align in _ALIGN:
                cell.text_frame.paragraphs[0].alignment = _ALIGN[align]
            if run is None:
                continue
            run.font.name = style["font"]
            run.font.size = Pt(float(style["sizePt"]))
            run.font.bold = bool(r == 0 and header_style.get("bold", True))
            colour = _resolve(header_style.get("color") if r == 0 else style["colorRole"], palette)
            if colour is not None:
                run.font.color.rgb = colour


def _draw_chart(
    shapes: SlideShapes, element: dict[str, Any], design_system: dict[str, Any]
) -> None:
    box = element["box"]
    kind = element["chartType"]
    if kind == "line" and _unequal_positions(element):
        _draw_xy_line(shapes, element, design_system)
        return
    data = CategoryChartData()
    data.categories = element["categories"]
    for series in element["series"]:
        data.add_series(series["name"], _chart_values(series))
    # python-pptx types add_chart as taking ChartData and returning a Chart; it takes any chart
    # data and returns the graphic frame that holds the chart.
    frame = cast(
        Any,
        shapes.add_chart(
            _CHART_TYPES[kind],
            Emu(box["x"]),
            Emu(box["y"]),
            Emu(box["cx"]),
            Emu(box["cy"]),
            cast(Any, data),
        ),
    )
    chart = frame.chart
    style = _body_style(design_system)
    chart.font.name = style["font"]
    chart.font.size = Pt(float(style["sizePt"]))
    several = len(element["series"]) > 1
    chart.has_legend = several or kind == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    palette = _palette(design_system)
    colours = [RGBColor.from_string(palette[role]) for role in _SERIES_ROLES if role in palette]
    plot = chart.plots[0]
    _label_values(plot, kind, style, palette, element.get("groundHex"))
    if kind == "pie":
        for point, colour in zip(plot.series[0].points, cycle(colours), strict=False):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = colour
        return
    for series, colour in zip(plot.series, cycle(colours), strict=False):
        fill = series.format.line if kind == "line" else series.format.fill
        if kind == "line":
            fill.color.rgb = colour
        else:
            fill.solid()
            fill.fore_color.rgb = colour
    for axis, title in (
        (chart.category_axis, element["xTitle"]),
        (chart.value_axis, element["yTitle"]),
    ):
        if title:
            axis.has_title = True
            axis.axis_title.text_frame.text = title
    rule = next((palette[role] for role in _RULE_COLORS if role in palette), None)
    if rule is not None:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(rule)
        chart.value_axis.major_gridlines.format.line.width = Emu(_HAIRLINE_EMU)


def _chart_values(series: dict[str, Any]) -> list[float | None]:
    """None stays a gap. A measured zero stays zero."""
    return [None if value is None else float(value) for value in series["values"]]


def _unequal_positions(element: dict[str, Any]) -> bool:
    positions = element.get("positions")
    categories = element.get("categories") or []
    if not isinstance(positions, list) or len(positions) != len(categories) or len(positions) < 3:
        return False
    gaps = [
        float(positions[index + 1]) - float(positions[index]) for index in range(len(positions) - 1)
    ]
    if any(gap <= 0 for gap in gaps):
        return False
    return max(gaps) - min(gaps) > 1e-9


def _draw_xy_line(
    shapes: SlideShapes, element: dict[str, Any], design_system: dict[str, Any]
) -> None:
    """A line on the category positions, so unequal steps are not drawn as equal slots."""
    box = element["box"]
    data = XyChartData()
    positions = [float(value) for value in element["positions"]]
    for series in element["series"]:
        added = data.add_series(series["name"])
        for position, value in zip(positions, _chart_values(series), strict=False):
            added.add_data_point(position, value)
    frame = cast(
        Any,
        shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER_LINES,
            Emu(box["x"]),
            Emu(box["y"]),
            Emu(box["cx"]),
            Emu(box["cy"]),
            cast(Any, data),
        ),
    )
    chart = frame.chart
    style = _body_style(design_system)
    chart.font.name = style["font"]
    chart.font.size = Pt(float(style["sizePt"]))
    chart.has_legend = len(element["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    palette = _palette(design_system)
    colours = [RGBColor.from_string(palette[role]) for role in _SERIES_ROLES if role in palette]
    for series, colour in zip(chart.series, cycle(colours), strict=False):
        series.format.line.color.rgb = colour
    for axis, title in (
        (chart.category_axis, element.get("xTitle") or ""),
        (chart.value_axis, element.get("yTitle") or ""),
    ):
        if title:
            axis.has_title = True
            axis.axis_title.text_frame.text = title


def _luminance(hex_value: str) -> float:
    channels = []
    for index in (0, 2, 4):
        c = int(hex_value[index : index + 2], 16) / 255
        channels.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


#: WCAG AA for normal text: a value printed over the plot must be read, not guessed.
_LABEL_CONTRAST = 4.5


def _label_colour(
    role: str | None, palette: dict[str, str], ground: str | None = None
) -> RGBColor | None:
    """The text colour of the profile that reads on the chart's surface.

    A value label sits on the card the chart fills. The body style's own colour comes first; one
    too faint on that surface (a light caption grey, white text of a dark slide) gives way to the
    profile's text colour with the most contrast.
    """
    ground = ground or palette.get("surface") or palette.get("background") or "FFFFFF"
    candidates = [
        palette.get(role or "", role or ""),
        *(palette.get(r) for r in ("text", "text-muted")),
    ]
    valid: list[str] = []
    for value in candidates:
        if value and len(value) == 6:
            try:
                int(value, 16)
            except ValueError:
                continue
            valid.append(value)
    if not valid:
        return None
    good = [value for value in valid if _contrast(value, ground) >= _LABEL_CONTRAST]
    chosen = good[0] if good else max(valid, key=lambda value: _contrast(value, ground))
    if _contrast(chosen, ground) < _LABEL_CONTRAST:
        chosen = "000000" if _luminance(ground) > 0.18 else "FFFFFF"
    return RGBColor.from_string(chosen)


def _label_values(
    plot: Any,
    kind: str,
    style: dict[str, Any],
    palette: dict[str, str],
    ground: str | None = None,
) -> None:
    """Print every value on its mark: a chart on a slide is read, not measured against an axis.

    A pie shows each slice's value and its share of the whole.
    """
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_value = True
    labels.number_format = "General"
    labels.number_format_is_linked = False
    labels.font.size = Pt(float(style["sizePt"]))
    colour = _label_colour(style["colorRole"], palette, ground)
    if colour is not None and kind != "pie":
        labels.font.color.rgb = colour
    if kind == "pie":
        labels.show_percentage = True
        labels.position = XL_LABEL_POSITION.BEST_FIT
    elif kind in ("column", "bar"):
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
    elif kind == "line":
        labels.position = XL_LABEL_POSITION.ABOVE
