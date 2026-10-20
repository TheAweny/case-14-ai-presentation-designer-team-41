"""Conservative file audit of image/text intersections, without rendering text into the sample.

Only axis-aligned rectangles, positive crop and opaque sRGB paints are measured. Unsupported
paints/transforms remain unknown; an opaque rectangle above them can make them irrelevant.
The entire inset text frame is a conservative text region, not a claim of glyph positioning.
All media pixels contributing to that region are inspected, with explicit work limits.
"""

from __future__ import annotations

import math
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from itertools import pairwise
from pathlib import Path
from typing import Any

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from engine.checks._geometry import Rect
from engine.checks._rules import _Emitter
from engine.checks.registry import PPTX_CHECKS

__all__ = ["audit_image_readability"]

_CHECK = PPTX_CHECKS[0]
_MAX_MEDIA_PIXELS = 16_000_000
_MAX_REGION_PIXELS = 1_000_000
_MAX_CELLS = 256
_LINEAR = tuple(v / 3294.6 if v <= 10 else ((v / 255 + 0.055) / 1.055) ** 2.4 for v in range(256))
_FILLS = frozenset({"solidFill", "noFill", "gradFill", "blipFill", "pattFill", "grpFill"})


@dataclass
class _PixelBudget:
    remaining: int = _MAX_REGION_PIXELS


class _UnknownError(ValueError):
    """A concrete OOXML/media feature that this bounded audit cannot measure."""


def _tag(node: Any) -> str:
    return str(node.tag).rsplit("}", 1)[-1]


def _fill_node(node: Any) -> Any:
    return next((child for child in node if _tag(child) in _FILLS), None)


def _colour(fill: Any) -> tuple[int, int, int]:
    if fill is None or _tag(fill) != "solidFill":
        raise _UnknownError("non-solid or inherited colour")
    children = list(fill)
    if len(children) != 1 or _tag(children[0]) != "srgbClr":
        raise _UnknownError("theme/system colour requires resolution")
    colour = children[0]
    if any(_tag(c) != "alpha" or c.get("val") != "100000" for c in colour):
        raise _UnknownError("transparent or transformed colour")
    value = colour.get("val", "")
    if len(value) != 6:
        raise _UnknownError("invalid sRGB colour")
    try:
        return int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16)
    except ValueError as error:
        raise _UnknownError("invalid sRGB colour") from error


def _luminance(rgb: tuple[int, ...]) -> float:
    return 0.2126 * _LINEAR[rgb[0]] + 0.7152 * _LINEAR[rgb[1]] + 0.0722 * _LINEAR[rgb[2]]


@dataclass(frozen=True)
class _Transform:
    sx: float = 1
    sy: float = 1
    tx: float = 0
    ty: float = 0
    reason: str = ""

    def box(self, x: float, y: float, cx: float, cy: float) -> Rect:
        left, top = math.floor(self.tx + x * self.sx), math.floor(self.ty + y * self.sy)
        right = math.ceil(self.tx + (x + cx) * self.sx)
        bottom = math.ceil(self.ty + (y + cy) * self.sy)
        return Rect(left, top, right - left, bottom - top)


@dataclass
class _Layer:
    shape: Any
    box: Rect
    transform: _Transform
    element_id: str
    reason: str = ""

    @property
    def picture(self) -> bool:
        return (
            self.shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}
            or self.shape._element.find("{*}spPr/{*}blipFill") is not None
        )

    @property
    def text(self) -> bool:
        return self.shape.has_text_frame and bool(self.shape.text.strip())


def _unsupported_transform(xfrm: Any) -> bool:
    return xfrm is not None and (
        int(xfrm.get("rot", "0")) % 21600000 != 0
        or xfrm.get("flipH", "0") in {"1", "true"}
        or xfrm.get("flipV", "0") in {"1", "true"}
    )


def _flatten(
    shapes: Any, prefix: str, slide_box: Rect, transform: _Transform | None = None
) -> Iterator[_Layer]:
    transform = transform or _Transform()
    for shape in shapes:
        element_id = f"{prefix}-{shape.shape_id}"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.find("{*}grpSpPr/{*}xfrm")
            reason = transform.reason
            props = shape._element.find("{*}grpSpPr")
            effects = props is not None and any(
                _tag(child) in {"effectLst", "effectDag", "scene3d", "sp3d"} for child in props
            )
            if xfrm is None or _unsupported_transform(xfrm) or effects:
                reason = "group effects, rotation/flip or absent group transform"
                child = _Transform(reason=reason)
            else:
                nodes = [xfrm.find("{*}" + name) for name in ("off", "ext", "chOff", "chExt")]
                if any(n is None for n in nodes):
                    child = _Transform(reason="incomplete group transform")
                else:
                    off, ext, ch_off, ch_ext = nodes
                    cw, ch = int(ch_ext.get("cx", 0)), int(ch_ext.get("cy", 0))
                    if cw <= 0 or ch <= 0:
                        child = _Transform(reason="degenerate group transform")
                    else:
                        sx, sy = int(ext.get("cx")) / cw, int(ext.get("cy")) / ch
                        child = _Transform(
                            transform.sx * sx,
                            transform.sy * sy,
                            transform.tx
                            + transform.sx * (int(off.get("x")) - sx * int(ch_off.get("x"))),
                            transform.ty
                            + transform.sy * (int(off.get("y")) - sy * int(ch_off.get("y"))),
                            reason,
                        )
            yield from _flatten(shape.shapes, element_id, slide_box, child)
            continue
        reason = transform.reason
        if _unsupported_transform(shape._element.find("{*}spPr/{*}xfrm")):
            reason = "shape rotation/flip"
        values = (shape.left, shape.top, shape.width, shape.height)
        if any(v is None for v in values):
            reason = "unresolved inherited geometry"
        box = slide_box if reason else transform.box(*values)
        yield _Layer(shape, box, transform, element_id, reason)


def _drawing_properties(layer: _Layer) -> Any:
    if layer.reason:
        raise _UnknownError(layer.reason)
    if layer.shape.is_placeholder:
        raise _UnknownError("unresolved placeholder drawing inheritance")
    props = layer.shape._element.find("{*}spPr")
    if props is None:
        raise _UnknownError("non-shape drawing object")
    if any(
        props.find("{*}" + name) is not None
        for name in ("effectLst", "effectDag", "scene3d", "sp3d")
    ):
        raise _UnknownError("drawing effects or 3D transform")
    effect_ref = layer.shape._element.find("{*}style/{*}effectRef")
    if effect_ref is not None and effect_ref.get("idx", "0") != "0":
        raise _UnknownError("unresolved theme effect reference")
    return props


def _rectangular(layer: _Layer) -> None:
    props = _drawing_properties(layer)
    geom = props.find("{*}prstGeom")
    if geom is None or geom.get("prst") != "rect":
        raise _UnknownError("non-rectangular geometry")


def _paint(layer: _Layer, *, own_text: bool = False) -> tuple[int, int, int] | None:
    props = _drawing_properties(layer)
    fill = _fill_node(props)
    line = props.find("{*}ln")
    if line is not None and line.find("{*}noFill") is None:
        raise _UnknownError("visible shape outline")
    line_ref = layer.shape._element.find("{*}style/{*}lnRef")
    if line is None and line_ref is not None and line_ref.get("idx", "0") != "0":
        raise _UnknownError("unresolved theme outline reference")
    if fill is not None and _tag(fill) == "noFill":
        nonvisual = layer.shape._element.find("{*}nvSpPr/{*}cNvSpPr")
        if own_text and nonvisual is not None and nonvisual.get("txBox") in {"1", "true"}:
            return None
    _rectangular(layer)
    if fill is not None and _tag(fill) == "noFill":
        return None
    return _colour(fill)


def _text_region(layer: _Layer) -> Rect:
    if layer.reason:
        raise _UnknownError(layer.reason)
    frame = layer.shape.text_frame
    body = frame._txBody.bodyPr
    if body.get("vert", "horz") != "horz" or int(body.get("rot", "0")) != 0:
        raise _UnknownError("vertical or rotated text")
    if body.find("{*}spAutoFit") is not None:
        raise _UnknownError("text frame can grow outside its stored bounds")
    if int(body.get("numCol", "1")) != 1:
        raise _UnknownError("multi-column text")
    if any(body.find("{*}" + name) is not None for name in ("prstTxWarp", "scene3d", "sp3d")):
        raise _UnknownError("text warp or 3D transform")
    return layer.transform.box(
        layer.shape.left + frame.margin_left,
        layer.shape.top + frame.margin_top,
        layer.shape.width - frame.margin_left - frame.margin_right,
        layer.shape.height - frame.margin_top - frame.margin_bottom,
    )


def _text_colours(layer: _Layer) -> set[tuple[int, int, int]]:
    colours: set[tuple[int, int, int]] = set()
    for paragraph in layer.shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            if run._r.find("{*}rPr/{*}effectLst") is not None:
                raise _UnknownError("text effects")
            fill = None
            for props in (
                run._r.find("{*}rPr"),
                paragraph._p.find("{*}pPr/{*}defRPr"),
                layer.shape.text_frame._txBody.find(
                    f"{{*}}lstStyle/{{*}}lvl{paragraph.level + 1}pPr/{{*}}defRPr"
                ),
            ):
                if props is not None:
                    if any(
                        props.find("{*}" + effect) is not None
                        for effect in (
                            "effectLst",
                            "effectDag",
                            "ln",
                            "scene3d",
                            "sp3d",
                            "highlight",
                        )
                    ):
                        raise _UnknownError("text effects or outline")
                    fill = _fill_node(props)
                    if fill is not None:
                        break
            colours.add(_colour(fill))
        if paragraph._p.find("{*}fld") is not None:
            raise _UnknownError("dynamic text field")
    if not colours:
        raise _UnknownError("unresolved text colour")
    return colours


def _image_range(
    layer: _Layer, cell: Rect, cache: dict[str, Image.Image], budget: _PixelBudget
) -> tuple[float, float]:
    _rectangular(layer)
    blip_fill = layer.shape._element.find("{*}blipFill")
    if blip_fill is None:
        raise _UnknownError("missing picture fill")
    blip = blip_fill.find("{*}blip")
    if blip is None or len(blip):
        raise _UnknownError("picture opacity/effects")
    stretch = blip_fill.find("{*}stretch/{*}fillRect")
    if stretch is None or any(int(v) for v in stretch.attrib.values()):
        raise _UnknownError("picture tile or fill transform")
    try:
        image_ref = layer.shape.image
        key = image_ref.sha1
        if key not in cache:
            with Image.open(BytesIO(image_ref.blob)) as source:
                decoded = sum(i.width * i.height for i in cache.values())
                if decoded + source.width * source.height > _MAX_MEDIA_PIXELS:
                    raise _UnknownError("media exceeds decoded pixel budget")
                if (
                    source.format not in {"PNG", "JPEG", "BMP"}
                    or getattr(source, "n_frames", 1) != 1
                ):
                    raise _UnknownError("unsupported or animated media format")
                if source.getexif().get(274, 1) != 1:
                    raise _UnknownError("EXIF image orientation")
                if source.mode in {"CMYK", "LAB"} or source.info.get("icc_profile"):
                    raise _UnknownError("media colour profile requires colour management")
                cache[key] = source.convert("RGBA")
        image = cache[key]
    except (OSError, ValueError, KeyError, AttributeError, UnidentifiedImageError) as error:
        if isinstance(error, _UnknownError):
            raise
        raise _UnknownError(f"unreadable/linked media: {type(error).__name__}") from error
    left, top, right, bottom = (
        layer.shape.crop_left,
        layer.shape.crop_top,
        layer.shape.crop_right,
        layer.shape.crop_bottom,
    )
    if min(left, top, right, bottom) < 0 or left + right >= 1 or top + bottom >= 1:
        raise _UnknownError("negative or degenerate image crop")
    box = layer.box
    x1 = image.width * (left + (cell.x - box.x) / box.cx * (1 - left - right))
    x2 = image.width * (left + (cell.right - box.x) / box.cx * (1 - left - right))
    y1 = image.height * (top + (cell.y - box.y) / box.cy * (1 - top - bottom))
    y2 = image.height * (top + (cell.bottom - box.y) / box.cy * (1 - top - bottom))
    roi = (
        max(0, math.floor(x1) - 1),
        max(0, math.floor(y1) - 1),
        min(image.width, math.ceil(x2) + 1),
        min(image.height, math.ceil(y2) + 1),
    )
    budget.remaining -= (roi[2] - roi[0]) * (roi[3] - roi[1])
    if budget.remaining < 0:
        raise _UnknownError("local media region exceeds pixel budget")
    sample = image.crop(roi)
    if sample.getchannel("A").getextrema()[0] != 255:
        raise _UnknownError("transparent media needs compositing")
    low, high = 1.0, 0.0
    pixels = sample.tobytes()
    for offset in range(0, len(pixels), 4):
        value = _luminance((pixels[offset], pixels[offset + 1], pixels[offset + 2]))
        low, high = min(low, value), max(high, value)
    return low, high


def _cells(region: Rect, layers: list[_Layer]) -> Iterator[Rect]:
    intersections = [b for layer in layers if (b := region.intersection(layer.box)) is not None]
    xs = sorted({region.x, region.right, *(v for b in intersections for v in (b.x, b.right))})
    ys = sorted({region.y, region.bottom, *(v for b in intersections for v in (b.y, b.bottom))})
    if (len(xs) - 1) * (len(ys) - 1) > _MAX_CELLS:
        raise _UnknownError("overlap partition exceeds cell budget")
    for x1, x2 in pairwise(xs):
        for y1, y2 in pairwise(ys):
            yield Rect(x1, y1, x2 - x1, y2 - y1)


def _measure(
    text: _Layer, region: Rect, layers: list[_Layer], cache: dict[str, Image.Image]
) -> tuple[float, float]:
    colours = [_luminance(rgb) for rgb in _text_colours(text)]
    low, high = 1.0, 0.0
    budget = _PixelBudget(_MAX_REGION_PIXELS)
    for cell in _cells(region, layers):
        for layer in reversed(layers):
            if layer.box.intersection(cell) is None:
                continue
            if layer.picture:
                cell_low, cell_high = _image_range(layer, cell, cache, budget)
            else:
                if layer.text and layer is not text:
                    raise _UnknownError(f"background contains other text: {layer.element_id}")
                colour = _paint(layer, own_text=layer is text)
                if colour is None:
                    continue
                cell_low = cell_high = _luminance(colour)
            low, high = min(low, cell_low), max(high, cell_high)
            break
        else:
            raise _UnknownError("image covers only part of text frame; unresolved slide background")
    # Every background luminance in the observed interval is possible. This is deliberately
    # conservative when several media or partial panels create disjoint luminance ranges.
    ratios = [
        1.0
        if low <= colour <= high
        else min((max(colour, value) + 0.05) / (min(colour, value) + 0.05) for value in (low, high))
        for colour in colours
    ]
    return min(ratios), high - low


def audit_image_readability(
    pptx_path: str | Path,
    *,
    variant: str = "balanced",
    slide_ids: Sequence[str] | None = None,
    editable_text_shape_ids: Sequence[Sequence[str]] | None = None,
    profile: Mapping[str, Mapping[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Return AuditFinding dictionaries for an already saved PPTX; never modify the file.

    Target ids use ``slide-<shape_id>`` (nested group ids are included). Supply slide_ids to
    match a caller's report. editable_text_shape_ids selects decimal OOXML cNvPr ids per slide;
    absent means all slide text, an empty per-slide list means none. Unknowns are major
    findings with ``status=unknown; reason=...``
    in evidence.measured, since the existing contract has no unknown/pass status field.
    Merge findings before calculating report counts, category summaries, score and passed.
    Invalid files/arguments raise, rather than masquerading as successfully audited decks.
    """
    if variant not in {"dense", "balanced", "airy"}:
        raise ValueError("Unknown deck variant")
    configured = dict((profile or {}).get(_CHECK.id, {}))
    if not configured.pop("enabled", True):
        return []
    settings = {**_CHECK.defaults, **configured}
    min_ratio, max_range = float(settings["minContrastRatio"]), float(settings["maxLuminanceRange"])
    if not 1 <= min_ratio <= 21 or not 0 <= max_range <= 1:
        raise ValueError("Invalid image readability thresholds")
    deck = Presentation(str(pptx_path))
    if slide_ids is not None and len(slide_ids) != len(deck.slides):
        raise ValueError("slide_ids must match the saved deck")
    if slide_ids is not None and any(not value for value in slide_ids):
        raise ValueError("slide_ids must be nonempty")
    if editable_text_shape_ids is not None and (
        len(editable_text_shape_ids) != len(deck.slides)
        or any(
            not isinstance(value, str) or not value.isdecimal()
            for ids in editable_text_shape_ids
            for value in ids
        )
    ):
        raise ValueError("editable_text_shape_ids must contain decimal shape ids per slide")
    if deck.slide_width is None or deck.slide_height is None:
        raise ValueError("Missing slide size")
    slide_box = Rect(0, 0, int(deck.slide_width), int(deck.slide_height))
    findings: list[dict[str, Any]] = []
    for index, slide in enumerate(deck.slides):
        # Cache only one slide at a time so a deck cannot retain all decoded media at once.
        cache: dict[str, Image.Image] = {}
        layers: list[_Layer] = []
        background_picture = False
        for owner, prefix in (
            (slide.slide_layout.slide_master, "master"),
            (slide.slide_layout, "layout"),
            (slide, "slide"),
        ):
            if prefix == "master" and any(
                owner_element.get("showMasterSp", "1") in {"0", "false"}
                for owner_element in (slide._element, slide.slide_layout._element)
            ):
                continue
            background_picture |= (
                owner._element.find("{*}cSld/{*}bg/{*}bgPr/{*}blipFill") is not None
            )
            shapes = [s for s in owner.shapes if prefix == "slide" or not s.is_placeholder]
            layers.extend(_flatten(shapes, prefix, slide_box))
        if editable_text_shape_ids is not None:
            actual_ids = {
                str(layer.shape.shape_id)
                for layer in layers
                if layer.element_id.startswith("slide-") and layer.shape.has_text_frame
            }
            if not set(editable_text_shape_ids[index]) <= actual_ids:
                raise ValueError(f"Unknown editable text shape id on slide {index + 1}")
        emitter = _Emitter(
            _CHECK,
            {
                "slideId": slide_ids[index] if slide_ids is not None else f"s{index + 1}",
                "variant": variant,
            },
        )
        for position, text in enumerate(layers):
            if not text.text or not text.element_id.startswith("slide-"):
                continue
            if (
                editable_text_shape_ids is not None
                and str(text.shape.shape_id) not in (editable_text_shape_ids[index])
            ):
                continue
            pictures = [
                layer for layer in layers if layer.picture and text.box.intersection(layer.box)
            ]
            if not pictures and not background_picture:
                continue
            measured = ""
            try:
                region = _text_region(text).intersection(slide_box)
                if region is None or region.area <= 0:
                    raise _UnknownError("empty or off-slide text region")
                # Export can reserve safe room by increasing insets. A picture touching the
                # outer shape but not the actual text region is not a text/image background.
                pictures = [layer for layer in pictures if region.intersection(layer.box)]
                if not pictures and not background_picture:
                    continue
                for foreground in layers[position + 1 :]:
                    if not foreground.box.intersection(region):
                        continue
                    if foreground.picture:
                        if foreground.reason:
                            raise _UnknownError(foreground.reason)
                        raise _UnknownError(
                            f"foreground picture occludes text: {foreground.element_id}"
                        )
                    if foreground.text or _paint(foreground) is not None:
                        raise _UnknownError(
                            f"foreground shape occludes text: {foreground.element_id}"
                        )
                ratio, spread = _measure(text, region, layers[: position + 1], cache)
                if ratio >= min_ratio and spread <= max_range:
                    continue
                measured = f"status=fail; minContrastRatio={ratio:.3f}; luminanceRange={spread:.3f}"
                title = "Текст недостаточно защищён от изображения"
            except _UnknownError as error:
                measured = f"status=unknown; reason={error}"
                title = "Читаемость текста поверх изображения не измерена"
            emitter.add(
                title=title,
                detail=(
                    f"{measured}. Изображения: {', '.join(p.element_id for p in pictures)}. "
                    "Выберите совместимую композицию либо существующую защиту шаблона "
                    "через разрешённые операции и повторите аудит; сохраните содержание."
                ),
                element={"elementId": text.element_id, "box": text.box.as_box()},
                measured=measured,
                expected=(
                    f"minContrastRatio>={min_ratio}; luminanceRange<={max_range}; status=measured"
                ),
                quote=text.shape.text,
            )
        findings.extend(emitter.findings)
        for image in cache.values():
            image.close()
    return findings
