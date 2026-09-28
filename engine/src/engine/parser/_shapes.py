"""Walking the shape tree of one slide, layout or master.

Groups are flattened rather than nested. A consumer that lays out a deck cares where a shape sits
on the page and in what order it paints, not how the designer happened to group things; the group
chain is kept in ``Shape.path`` for the cases that do care.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from engine.parser._model import (
    Box,
    Paragraph,
    ParseWarning,
    Placeholder,
    Run,
    Shape,
    WarningCode,
)
from engine.parser._style import fill_of, geometry_of, has_shadow, stroke_of

_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

#: Shape kinds we describe explicitly. Anything else is recorded under its own enum name and
#: raises a warning, so an unsupported object is visible rather than silently dropped.
_KNOWN_KINDS = {
    MSO_SHAPE_TYPE.AUTO_SHAPE: "auto-shape",
    MSO_SHAPE_TYPE.TEXT_BOX: "text-box",
    MSO_SHAPE_TYPE.PICTURE: "picture",
    MSO_SHAPE_TYPE.TABLE: "table",
    MSO_SHAPE_TYPE.CHART: "chart",
    MSO_SHAPE_TYPE.GROUP: "group",
    MSO_SHAPE_TYPE.LINE: "line",
    MSO_SHAPE_TYPE.FREEFORM: "freeform",
    MSO_SHAPE_TYPE.PLACEHOLDER: "placeholder",
    MSO_SHAPE_TYPE.MEDIA: "media",
    MSO_SHAPE_TYPE.DIAGRAM: "diagram",
}


def _kind_of(shape: Any, location: str, warnings: list[ParseWarning]) -> str:
    """Name the shape kind, warning once for anything we do not model explicitly."""
    try:
        shape_type = shape.shape_type
    except (AttributeError, ValueError, NotImplementedError):
        shape_type = None

    known = _KNOWN_KINDS.get(shape_type) if shape_type is not None else None
    if known is not None:
        return known

    name = getattr(shape_type, "name", None) or "unknown"
    warnings.append(
        ParseWarning(
            code=WarningCode.UNSUPPORTED_SHAPE,
            message=f"shape kind {name} is recorded but not modelled",
            location=location,
        )
    )
    return str(name).lower().replace("_", "-")


def _box_of(shape: Any) -> Box | None:
    """Geometry in EMU, or nothing when the shape inherits its position from a layout."""
    left, top = shape.left, shape.top
    width, height = shape.width, shape.height
    if left is None or top is None or width is None or height is None:
        return None
    return Box(x=int(left), y=int(top), cx=int(width), cy=int(height))


class _Transform:
    """Affine scale and translation from a group child space to slide EMU."""

    def __init__(
        self,
        scale_x: float = 1,
        scale_y: float = 1,
        offset_x: float = 0,
        offset_y: float = 0,
    ) -> None:
        self.scale_x = scale_x
        self.scale_y = scale_y
        self.offset_x = offset_x
        self.offset_y = offset_y

    def box(self, value: Box | None) -> Box | None:
        if value is None:
            return None
        return Box(
            x=round(self.offset_x + value.x * self.scale_x),
            y=round(self.offset_y + value.y * self.scale_y),
            cx=round(value.cx * self.scale_x),
            cy=round(value.cy * self.scale_y),
        )

    def group_children(self, shape: Any) -> _Transform:
        """Compose this transform with the OOXML child coordinate system of one group."""
        try:
            xfrm = shape._element.grpSpPr.xfrm
            off, ext = xfrm.off, xfrm.ext
            child_off, child_ext = xfrm.chOff, xfrm.chExt
            scale_x = int(ext.cx) / int(child_ext.cx)
            scale_y = int(ext.cy) / int(child_ext.cy)
            return _Transform(
                self.scale_x * scale_x,
                self.scale_y * scale_y,
                self.offset_x + self.scale_x * (int(off.x) - int(child_off.x) * scale_x),
                self.offset_y + self.scale_y * (int(off.y) - int(child_off.y) * scale_y),
            )
        except (AttributeError, TypeError, ValueError, ZeroDivisionError):
            return self


def _flips(shape: Any) -> tuple[bool, bool]:
    """Horizontal and vertical flip, which python-pptx does not surface on every shape type."""
    try:
        xfrm = shape._element.spPr.xfrm
    except AttributeError:
        return False, False
    if xfrm is None:
        return False, False
    return bool(xfrm.get("flipH") == "1"), bool(xfrm.get("flipV") == "1")


def _placeholder_of(shape: Any) -> Placeholder | None:
    """Placeholder identity when the shape inherits from a layout, otherwise nothing."""
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        fmt = shape.placeholder_format
        type_name = getattr(fmt.type, "name", None) or str(fmt.type)
        return Placeholder(type=str(type_name).lower().replace("_", "-"), idx=fmt.idx)
    except (AttributeError, ValueError):
        return None


def _color_of(run: Any) -> str:
    """Hex colour of a run when it is stated outright.

    A run whose colour comes from the theme is left empty on purpose: substituting the resolved
    value here would hide the fact that the slide inherits it, and the design-system layer needs
    that distinction to tell a brand colour from a one-off override.
    """
    try:
        color = run.font.color
        if color is None or color.type is None:
            return ""
        rgb = getattr(color, "rgb", None)
        return str(rgb).upper() if rgb is not None else ""
    except (AttributeError, ValueError):
        return ""


def _paragraphs_of(shape: Any) -> tuple[Paragraph, ...]:
    """Text of a shape, paragraph by paragraph, with the formatting stated on each run."""
    if not getattr(shape, "has_text_frame", False):
        return ()

    paragraphs: list[Paragraph] = []
    for para in shape.text_frame.paragraphs:
        runs: list[Run] = []
        for run in para.runs:
            size = run.font.size
            runs.append(
                Run(
                    text=run.text,
                    font=run.font.name or "",
                    size_pt=round(Emu(size).pt, 2) if size is not None else None,
                    bold=run.font.bold,
                    italic=run.font.italic,
                    color=_color_of(run),
                )
            )
        paragraphs.append(
            Paragraph(level=para.level or 0, runs=tuple(runs), line_spacing=_line_spacing(para))
        )
    return tuple(paragraphs)


def _line_spacing(paragraph: Any) -> float | None:
    """Line spacing as a multiple, when the paragraph states it as a percentage.

    python-pptx returns a float for a percentage and a length for exact spacing. Exact spacing is
    tied to one font size and says nothing about the rhythm of the template, so it is left out.
    """
    try:
        spacing = paragraph.line_spacing
    except (AttributeError, ValueError):
        return None
    if isinstance(spacing, float):
        return round(spacing, 3)
    return None


def _table_size_of(shape: Any) -> tuple[int, int] | None:
    if not getattr(shape, "has_table", False):
        return None
    table = shape.table
    return len(table.rows), len(table.columns)


def _chart_of(shape: Any, warnings: list[ParseWarning], location: str) -> tuple[str, str]:
    """Chart part name and chart type, empty when the shape holds no chart."""
    if not getattr(shape, "has_chart", False):
        return "", ""
    try:
        chart = shape.chart
        chart_type = getattr(chart.chart_type, "name", "") or ""
        part = str(chart.part.partname)
        return part, str(chart_type).lower().replace("_", "-")
    except (AttributeError, ValueError, KeyError) as exc:
        warnings.append(
            ParseWarning(
                code=WarningCode.UNSUPPORTED_SHAPE,
                message=f"chart could not be described: {exc}",
                location=location,
            )
        )
        return "", ""


def _media_sha1(shape: Any, warnings: list[ParseWarning], location: str) -> str:
    """Digest of the embedded image, used later to name the media part it draws."""
    try:
        return str(shape.image.sha1)
    except (AttributeError, ValueError, KeyError, NotImplementedError):
        warnings.append(
            ParseWarning(
                code=WarningCode.UNREADABLE_MEDIA,
                message="picture does not carry an embedded image; it may be linked",
                location=location,
            )
        )
        return ""


def collect_shapes(
    container: Any,
    location: str,
    warnings: list[ParseWarning],
    sha_to_part: dict[str, str],
    path: tuple[int, ...] = (),
    transform: _Transform | None = None,
) -> tuple[Shape, ...]:
    """Flatten a shape container into an ordered tuple.

    Order is the drawing order of the file, which is also the paint order, so ``z`` is simply the
    index within the parent.
    """
    collected: list[Shape] = []
    transform = transform or _Transform()
    for index, shape in enumerate(container):
        shape_id = int(getattr(shape, "shape_id", 0) or 0)
        where = f"{location}#{shape_id}"
        kind = _kind_of(shape, where, warnings)
        flip_h, flip_v = _flips(shape)
        chart_part, chart_type = _chart_of(shape, warnings, where)

        media_part = ""
        if kind == "picture":
            digest = _media_sha1(shape, warnings, where)
            media_part = sha_to_part.get(digest, "")

        box = transform.box(_box_of(shape))
        props = shape._element.find(f"{_P}spPr")
        geometry, radius = geometry_of(props, box.cx if box else None, box.cy if box else None)

        collected.append(
            Shape(
                shape_id=shape_id,
                name=str(getattr(shape, "name", "") or ""),
                kind=kind,
                box=box,
                rotation=float(getattr(shape, "rotation", 0.0) or 0.0),
                flip_h=flip_h,
                flip_v=flip_v,
                z=index,
                path=path,
                placeholder=_placeholder_of(shape),
                paragraphs=_paragraphs_of(shape),
                media_part=media_part,
                table_size=_table_size_of(shape),
                chart_part=chart_part,
                chart_type=chart_type,
                fill=fill_of(props),
                stroke=stroke_of(props),
                geometry=geometry,
                corner_radius_emu=radius,
                has_shadow=has_shadow(props),
            )
        )

        if kind == "group":
            collected.extend(
                collect_shapes(
                    shape.shapes,
                    location,
                    warnings,
                    sha_to_part,
                    (*path, shape_id),
                    transform.group_children(shape),
                )
            )

    return tuple(collected)
