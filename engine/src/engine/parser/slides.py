"""Reading a finished deck back as SlideIR: what each slide looks like, as editable elements.

TemplateIR describes a file as it is written, which keeps inheritance visible and leaves most
values unresolved. An editor that opens a generated deck needs the opposite: every slide as it is
drawn. This module resolves what PowerPoint resolves before it paints a slide:

* the graphics of the master and the layout underneath the slide's own shapes, unless a part
  hides them (``showMasterSp="0"``);
* a placeholder's position, size and text style from the layout and master placeholders it
  follows, and from the master's ``p:txStyles``;
* scheme colours through the colour map in force and their transforms; theme font references;
* shape style references (``p:style``) against the theme's format scheme;
* the outline of freeforms and preset shapes, as SVG path data;
* the background, first stated on the slide, the layout or the master;
* the number a slide-number field shows.

The result is a list of SlideIR payloads, each validated against the generated contract. Nothing
here branches on a template: every value comes out of the file.
"""

from __future__ import annotations

import logging
import math
from collections.abc import Callable, Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import lxml.etree as etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from engine.contracts import SlideIR
from engine.parser._deck_frames import chart_fields, table_fields
from engine.parser._deck_style import (
    Colors,
    Paint,
    Theme,
    fill_child,
    paint_of,
    paragraph_props,
    read_theme,
    run_props,
    style_fill,
    style_line,
)
from engine.parser._geometry import adjustments_of, custom_geometry, preset_geometry
from engine.parser._media import mean_color

logger = logging.getLogger(__name__)

_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

SLIDE_IR_SCHEMA_VERSION = 1

_TITLE_TYPES = frozenset(("title", "ctrTitle"))
#: Placeholders whose text PowerPoint styles like ordinary shapes, from ``p:otherStyle``.
_OTHER_STYLE_TYPES = frozenset(("dt", "ftr", "sldNum", "hdr"))
_ROLE_OF_PLACEHOLDER = {
    "title": "title",
    "ctrTitle": "title",
    "subTitle": "subtitle",
    "sldNum": "page-number",
    "ftr": "footer",
    "dt": "footer",
}
_LINE_PRESETS = frozenset(("line", "straightConnector1"))
_NAMED_PRESETS = {"rect": "rect", "ellipse": "ellipse", "diamond": "diamond"}
_DASH = {
    "solid": "solid",
    "dot": "dot",
    "sysDot": "dot",
    "dash": "dash",
    "sysDash": "dash",
    "lgDash": "dash",
    "dashDot": "dash",
    "lgDashDot": "dash",
    "lgDashDotDot": "dash",
    "sysDashDot": "dash",
    "sysDashDotDot": "dash",
}
_ANCHOR = {"t": "top", "ctr": "middle", "b": "bottom", "just": "top", "dist": "middle"}
#: ``a:bodyPr`` insets when nothing states them: 0.1 inch across, 0.05 inch down.
_DEFAULT_INSETS = {"l": 91440, "t": 45720, "r": 91440, "b": 45720}
_INSET_ATTRS = {"l": "lIns", "t": "tIns", "r": "rIns", "b": "bIns"}
#: Size of text nothing sizes, in points: the OOXML default.
_DEFAULT_SIZE_PT = 18.0
#: Width of an outline that states none, in EMU: PowerPoint's 0.75 pt.
_DEFAULT_LINE_EMU = 9525
#: A zero-width outline is a hairline; it still has to show.
_HAIRLINE_EMU = 3175
#: Default corner of a rounded rectangle, in 100 000ths of the shorter side.
_ROUND_RECT_ADJ = 16667
#: A picture of the master or layout covering less of the slide than this reads as a logo.
_LOGO_SHARE = 0.05
#: Single line spacing as a multiple of the type size, which spacing in points is measured against.
_SINGLE_LINE = 1.2


# ----------------------------------------------------------------------------------- structures


@dataclass(frozen=True)
class _Placeholder:
    type: str
    idx: int | None
    element: Any


@dataclass
class _Master:
    element: Any
    part: Any
    theme: Theme
    color_map: dict[str, str]
    placeholders: list[_Placeholder]
    styles: dict[str, Any]


@dataclass
class _Layout:
    element: Any
    part: Any
    placeholders: list[_Placeholder]
    color_map: dict[str, str] | None


@dataclass(frozen=True)
class _Origin:
    """The part a shape is drawn from: ``master``, ``layout`` or ``slide``."""

    kind: str
    part: Any


@dataclass(frozen=True)
class _Group:
    off: tuple[float, float]
    ext: tuple[float, float]
    child_off: tuple[float, float]
    child_ext: tuple[float, float]
    rotation: float
    flip_h: bool
    flip_v: bool


@dataclass
class _Placement:
    x: float
    y: float
    cx: float
    cy: float
    rotation: float = 0.0
    flip_h: bool = False
    flip_v: bool = False

    def box(self) -> dict[str, int]:
        return {"x": round(self.x), "y": round(self.y), "cx": round(self.cx), "cy": round(self.cy)}


@dataclass(frozen=True)
class _Scope:
    """Where the walk is: the groups around a shape, outermost first, and the group fill."""

    groups: tuple[_Group, ...] = ()
    group_ids: tuple[str, ...] = ()
    group_paint: Paint | None = None


@dataclass
class _Slide:
    index: int
    number: int
    part: Any
    element: Any
    layout: _Layout
    master: _Master
    colors: Colors
    default_style: Any
    color_value: Callable[[str], str]
    size: tuple[int, int]
    media: dict[str, bytes]
    ids: set[str] = field(default_factory=set)

    @property
    def theme(self) -> Theme:
        return self.master.theme

    def unique(self, base: str) -> str:
        candidate, counter = base, 2
        while candidate in self.ids:
            candidate = f"{base}-{counter}"
            counter += 1
        self.ids.add(candidate)
        return candidate


# -------------------------------------------------------------------------------------- helpers


def _local(element: Any) -> str:
    return etree.QName(element).localname if isinstance(element.tag, str) else ""


def _int(value: str | None, default: int | None = None) -> int | None:
    if value is None:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def _truthy(value: str | None) -> bool:
    return value in ("1", "true", "on")


def _placeholder_ref(shape: Any) -> Any:
    return shape.find(f"*/{_P}nvPr/{_P}ph")


def _placeholders(tree: Any) -> list[_Placeholder]:
    found = []
    for shape in tree:
        ph = _placeholder_ref(shape) if isinstance(shape.tag, str) else None
        if ph is not None:
            found.append(_Placeholder(ph.get("type", "obj"), _int(ph.get("idx")), shape))
    return found


def _family(kind: str) -> str:
    if kind in _TITLE_TYPES:
        return "title"
    if kind in _OTHER_STYLE_TYPES:
        return kind
    return "body"


def _match(
    kind: str, idx: int | None, candidates: Iterable[_Placeholder], by_idx: bool
) -> _Placeholder | None:
    """The placeholder a placeholder follows: the same index, then the same type or family."""
    candidates = list(candidates)
    if by_idx and idx is not None:
        for candidate in candidates:
            if candidate.idx == idx:
                return candidate
    for candidate in candidates:
        if candidate.type == kind:
            return candidate
    family = _family(kind)
    for candidate in candidates:
        if _family(candidate.type) == family:
            return candidate
    return None


def _color_map(element: Any, override: bool) -> dict[str, str] | None:
    if override:
        mapping = element.find(f"{_P}clrMapOvr/{_A}overrideClrMapping")
    else:
        mapping = element.find(f"{_P}clrMap")
    return dict(mapping.attrib) if mapping is not None else None


def _tree(element: Any) -> Any:
    return element.find(f"{_P}cSld/{_P}spTree")


def _xfrm(xfrm: Any) -> _Placement | None:
    if xfrm is None:
        return None
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    if off is None or ext is None:
        return None
    return _Placement(
        x=float(_int(off.get("x"), 0) or 0),
        y=float(_int(off.get("y"), 0) or 0),
        cx=float(_int(ext.get("cx"), 0) or 0),
        cy=float(_int(ext.get("cy"), 0) or 0),
        rotation=(_int(xfrm.get("rot"), 0) or 0) / 60000,
        flip_h=_truthy(xfrm.get("flipH")),
        flip_v=_truthy(xfrm.get("flipV")),
    )


def _group_of(xfrm: Any) -> _Group | None:
    placed = _xfrm(xfrm)
    if placed is None:
        return None
    child_off, child_ext = xfrm.find(f"{_A}chOff"), xfrm.find(f"{_A}chExt")
    return _Group(
        off=(placed.x, placed.y),
        ext=(placed.cx, placed.cy),
        child_off=(
            (float(_int(child_off.get("x"), 0) or 0), float(_int(child_off.get("y"), 0) or 0))
            if child_off is not None
            else (placed.x, placed.y)
        ),
        child_ext=(
            (float(_int(child_ext.get("cx"), 0) or 0), float(_int(child_ext.get("cy"), 0) or 0))
            if child_ext is not None
            else (placed.cx, placed.cy)
        ),
        rotation=placed.rotation,
        flip_h=placed.flip_h,
        flip_v=placed.flip_v,
    )


def _in_group(placed: _Placement, group: _Group) -> _Placement:
    """Where a shape placed in a group's child space sits in the space around the group."""
    scale_x = group.ext[0] / group.child_ext[0] if group.child_ext[0] else 1.0
    scale_y = group.ext[1] / group.child_ext[1] if group.child_ext[1] else 1.0
    cx, cy = placed.cx * scale_x, placed.cy * scale_y
    centre_x = group.off[0] + (placed.x - group.child_off[0]) * scale_x + cx / 2
    centre_y = group.off[1] + (placed.y - group.child_off[1]) * scale_y + cy / 2
    rotation, flip_h, flip_v = placed.rotation, placed.flip_h, placed.flip_v
    group_x, group_y = group.off[0] + group.ext[0] / 2, group.off[1] + group.ext[1] / 2
    if group.flip_h:
        centre_x = 2 * group_x - centre_x
        flip_h, rotation = not flip_h, -rotation
    if group.flip_v:
        centre_y = 2 * group_y - centre_y
        flip_v, rotation = not flip_v, -rotation
    if group.rotation:
        angle = math.radians(group.rotation)
        dx, dy = centre_x - group_x, centre_y - group_y
        centre_x = group_x + dx * math.cos(angle) - dy * math.sin(angle)
        centre_y = group_y + dx * math.sin(angle) + dy * math.cos(angle)
        rotation += group.rotation
    return _Placement(centre_x - cx / 2, centre_y - cy / 2, cx, cy, rotation, flip_h, flip_v)


def _placed(placed: _Placement, scope: _Scope) -> _Placement:
    for group in reversed(scope.groups):
        placed = _in_group(placed, group)
    return placed


def _look(placed: _Placement, scope: _Scope, origin: _Origin) -> dict[str, Any]:
    look: dict[str, Any] = {}
    rotation = round(placed.rotation % 360, 3) % 360
    if rotation:
        look["rotation"] = rotation
    if scope.group_ids:
        look["groupIds"] = list(scope.group_ids)
    if origin.kind != "slide":
        # Graphics of the master and the layout cannot be picked on a slide in PowerPoint either.
        # Export that clones the source slide already paints them through the layout, so it must
        # not draw this copy a second time on top of the slide.
        look["locked"] = True
        look["sourcePart"] = origin.kind
    return look


def _media_part(part: Any, rid: str | None, slide: _Slide) -> str | None:
    """Part name of a related picture, remembering its bytes for the media directory."""
    if not rid:
        return None
    try:
        media = part.related_part(rid)
    except (KeyError, AttributeError):
        return None
    name = str(media.partname)
    slide.media.setdefault(name, media.blob)
    return name


# ------------------------------------------------------------------------------------------ text


def _font_ref(style: Any, slide: _Slide) -> dict[str, Any]:
    reference = style.find(f"{_A}fontRef") if style is not None else None
    if reference is None:
        return {}
    props: dict[str, Any] = {}
    face = {"major": slide.theme.major, "minor": slide.theme.minor}.get(
        reference.get("idx", ""), ""
    )
    if face:
        props["font"] = face
    color = slide.colors.first(reference)
    if color is not None:
        props["color"] = color[0]
    return props


@dataclass
class _Body:
    anchor: str
    insets: dict[str, int]
    font_scale: float
    spacing_reduction: float


def _body(bodies: Iterable[Any]) -> _Body:
    """``a:bodyPr`` as drawn: each attribute from the nearest part that states it."""
    bodies = [body for body in bodies if body is not None]
    anchor = next((b.get("anchor") for b in bodies if b.get("anchor")), "t")
    insets = {}
    for key, attr in _INSET_ATTRS.items():
        stated = next((b.get(attr) for b in bodies if b.get(attr) is not None), None)
        insets[key] = max(0, _int(stated, _DEFAULT_INSETS[key]) or 0)
    font_scale, reduction = 1.0, 0.0
    for body in bodies:
        autofit = next(
            (c for c in body if _local(c) in ("normAutofit", "spAutoFit", "noAutofit")), None
        )
        if autofit is None:
            continue
        if _local(autofit) == "normAutofit":
            font_scale = (_int(autofit.get("fontScale"), 100000) or 100000) / 100000
            reduction = (_int(autofit.get("lnSpcReduction"), 0) or 0) / 100000
        break
    return _Body(_ANCHOR.get(anchor, "top"), insets, font_scale, reduction)


_ROMAN = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"))
_ROMAN += ((40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))


def _roman(number: int) -> str:
    out = []
    for value, letters in _ROMAN:
        while number >= value:
            out.append(letters)
            number -= value
    return "".join(out)


def _alpha(number: int) -> str:
    letters = ""
    while number > 0:
        number, rest = divmod(number - 1, 26)
        letters = chr(ord("a") + rest) + letters
    return letters


def _autonumber(scheme: str, number: int) -> str:
    """The label ``a:buAutoNum`` of a scheme gives the n-th paragraph, such as ``1.`` or ``b)``."""
    if scheme.startswith("alphaLc"):
        label = _alpha(number)
    elif scheme.startswith("alphaUc"):
        label = _alpha(number).upper()
    elif scheme.startswith("romanLc"):
        label = _roman(number)
    elif scheme.startswith("romanUc"):
        label = _roman(number).upper()
    else:
        label = str(number)
    if scheme.endswith("ParenBoth"):
        return f"({label})"
    if scheme.endswith("ParenR"):
        return f"{label})"
    if scheme.endswith("Plain"):
        return label
    if scheme.endswith("Minus"):
        return f"- {label} -"
    return f"{label}."


def _run(text: str, props: dict[str, Any], slide: _Slide, font_scale: float) -> dict[str, Any]:
    size = float(props.get("size", _DEFAULT_SIZE_PT)) * font_scale
    return {
        "text": text,
        "font": props.get("font") or slide.theme.minor,
        "sizePt": round(size, 2),
        "bold": bool(props.get("b", False)),
        "italic": bool(props.get("i", False)),
        "underline": bool(props.get("u", False)),
        "color": props.get("color") or slide.colors.scheme("tx1") or "000000",
        "letterSpacingPt": float(props.get("spc", 0.0)),
    }


def _paragraphs(body: Any, sources: list[Any], slide: _Slide, fit: _Body) -> list[dict[str, Any]]:
    """Paragraphs of a text body with every property resolved, lowest source first."""
    paragraphs: list[dict[str, Any]] = []
    counters: dict[int, int] = {}
    colors, theme = slide.colors, slide.theme
    for paragraph in body.findall(f"{_A}p"):
        own = paragraph.find(f"{_A}pPr")
        level = max(0, min(8, _int(own.get("lvl"), 0) or 0)) if own is not None else 0
        ppr: dict[str, Any] = {}
        rpr: dict[str, Any] = {}
        for source in sources:
            if isinstance(source, dict):
                rpr.update(source)
                continue
            for element in (source.find(f"{_A}defPPr"), source.find(f"{_A}lvl{level + 1}pPr")):
                if element is None:
                    continue
                ppr.update(paragraph_props(element))
                rpr.update(run_props(element.find(f"{_A}defRPr"), colors, theme))
        if own is not None:
            ppr.update(paragraph_props(own))
            rpr.update(run_props(own.find(f"{_A}defRPr"), colors, theme))

        runs = []
        for child in paragraph:
            name = _local(child)
            if name not in ("r", "fld", "br"):
                continue
            props = {**rpr, **run_props(child.find(f"{_A}rPr"), colors, theme)}
            if name == "br":
                text = "\n"
            elif name == "fld" and child.get("type") == "slidenum":
                text = str(slide.number)
            else:
                text = child.findtext(f"{_A}t") or ""
            runs.append(_run(text, props, slide, fit.font_scale))
        if not runs:
            end = {**rpr, **run_props(paragraph.find(f"{_A}endParaRPr"), colors, theme)}
            runs.append(_run("", end, slide, fit.font_scale))
        size = runs[0]["sizePt"] or _DEFAULT_SIZE_PT
        visible = any(run["text"].strip() for run in runs)

        line = ppr.get("lnSpc")
        if line is None:
            spacing = 1.0 - fit.spacing_reduction
        elif line[0] == "pct":
            spacing = line[1] - fit.spacing_reduction
        else:
            spacing = line[1] / (size * _SINGLE_LINE)
        out: dict[str, Any] = {
            "level": level,
            "align": ppr.get("align", "left"),
            "runs": runs,
            "lineSpacing": round(max(spacing, 0.1), 4),
        }

        bullet = ppr.get("bullet")
        if bullet is not None and bullet[0] == "auto" and visible:
            start = bullet[2] if bullet[2] is not None else 1
            counters[level] = counters[level] + 1 if level in counters else start
            out["bullet"] = _autonumber(bullet[1], counters[level])
        elif visible:
            counters.pop(level, None)
        if bullet is not None and bullet[0] == "char" and visible and bullet[1]:
            out["bullet"] = bullet[1]
        for deeper in [key for key in counters if key > level]:
            del counters[deeper]

        if "marL" in ppr:
            out["marginEmu"] = max(0, int(ppr["marL"]))
        if "indent" in ppr:
            out["indentEmu"] = int(ppr["indent"])
        for key, name in (("spcBef", "spaceBeforePt"), ("spcAft", "spaceAfterPt")):
            stated = ppr.get(key)
            if stated is None:
                continue
            points = stated[1] * size if stated[0] == "pct" else stated[1]
            if points > 0:
                out[name] = round(points, 2)
        paragraphs.append(out)
    return paragraphs


def _has_text(paragraphs: list[dict[str, Any]]) -> bool:
    return any(run["text"].strip() for p in paragraphs for run in p["runs"])


# -------------------------------------------------------------------------------- paint & lines


def _shape_paint(
    props: list[tuple[Any, Any]], styles: list[Any], slide: _Slide, scope: _Scope
) -> tuple[Paint, Any]:
    """Fill of a shape and the part an image fill is resolved in."""
    for element, part in props:
        fill = fill_child(element)
        if fill is not None:
            paint = paint_of(fill, slide.colors)
            if paint.kind == "group":
                return scope.group_paint or Paint("none"), part
            return paint, part
    for style in styles:
        reference = style.find(f"{_A}fillRef") if style is not None else None
        if reference is not None:
            return style_fill(reference, slide.colors), None
    return Paint("none"), None


def _shape_stroke(props: list[Any], styles: list[Any], slide: _Slide) -> dict[str, Any] | None:
    """Outline of a shape: attributes from the nearest ``a:ln``, then the style's ``a:lnRef``."""
    lines = [p.find(f"{_A}ln") for p in props if p is not None]
    lines = [line for line in lines if line is not None]
    placeholder = None
    for style in styles:
        reference = style.find(f"{_A}lnRef") if style is not None else None
        if reference is not None:
            theme_line, placeholder = style_line(reference, slide.colors)
            if theme_line is not None:
                lines.append(theme_line)
            break
    fill = next((fill_child(line) for line in lines if fill_child(line) is not None), None)
    if fill is None:
        return None
    paint = paint_of(fill, slide.colors, placeholder)
    if paint.kind != "solid" or not paint.color:
        return None
    width = next((_int(line.get("w")) for line in lines if line.get("w") is not None), None)
    if width is None:
        width = _DEFAULT_LINE_EMU
    dash = next(
        (
            d.get("val", "solid")
            for d in (line.find(f"{_A}prstDash") for line in lines)
            if d is not None
        ),
        "solid",
    )

    def end(tag: str) -> str:
        stated = next(
            (e for e in (line.find(f"{_A}{tag}") for line in lines) if e is not None), None
        )
        kind = stated.get("type", "none") if stated is not None else "none"
        return "none" if kind == "none" else "arrow"

    return {
        "color": slide.color_value(paint.color),
        "widthEmu": width if width > 0 else _HAIRLINE_EMU,
        "dash": _DASH.get(dash, "solid"),
        "head": end("headEnd"),
        "tail": end("tailEnd"),
    }


def _outline(props: Any, placed: _Placement, where: str) -> tuple[dict[str, Any], bool]:
    """How a shape's geometry is said in SlideIR, and whether it is a line."""
    width, height = max(round(placed.cx), 1), max(round(placed.cy), 1)
    if props is None:
        return {"shape": "rect"}, False
    custom = props.find(f"{_A}custGeom")
    if custom is not None:
        try:
            geometry = custom_geometry(custom, width, height)
        except (ValueError, ZeroDivisionError, OverflowError) as error:
            logger.warning("freeform of %s cannot be traced: %s", where, error)
            return {}, False
        return ({"geometry": geometry} if geometry else {}), False
    preset = props.find(f"{_A}prstGeom")
    if preset is None:
        return {"shape": "rect"}, False
    name = preset.get("prst", "rect")
    stated = adjustments_of(preset)
    if name in _LINE_PRESETS:
        return {"shape": "line"}, True
    if name in _NAMED_PRESETS:
        return {"shape": _NAMED_PRESETS[name]}, False
    if name == "roundRect":
        adjust = _ROUND_RECT_ADJ
        for guide_name, formula in stated:
            if guide_name == "adj" and formula.startswith("val "):
                adjust = _int(formula.removeprefix("val ").strip(), adjust) or 0
        return {"shape": "round-rect", "cornerRadius": max(0.0, min(adjust / 100000, 0.5))}, False
    if name == "triangle" and all(formula == "val 50000" for _, formula in stated):
        return {"shape": "triangle"}, False
    try:
        geometry = preset_geometry(name, stated, width, height)
    except (ValueError, ZeroDivisionError, OverflowError) as error:
        logger.warning("preset %s of %s cannot be traced: %s", name, where, error)
        return {}, False
    if geometry is None:
        logger.warning("preset geometry %s of %s is not modelled; drawn as its box", name, where)
        return {}, False
    return {"geometry": geometry}, name.startswith(("bentConnector", "curvedConnector"))


# ------------------------------------------------------------------------------------- elements


def _decor(
    slide: _Slide,
    element_id: str,
    shape_id: int,
    placed: _Placement,
    look: dict[str, Any],
    outline: dict[str, Any],
    paint: Paint,
    stroke: dict[str, Any] | None,
) -> dict[str, Any]:
    decor: dict[str, Any] = {
        "kind": "decor",
        "elementId": element_id,
        "role": "decoration",
        "box": placed.box(),
        **look,
        "sourceShapeId": shape_id,
        **outline,
        "fill": slide.color_value(paint.color) if paint.kind == "solid" and paint.color else None,
        "stroke": stroke,
    }
    if paint.kind == "solid" and paint.alpha < 1 and stroke is None:
        decor["opacity"] = round(max(0.0, paint.alpha), 3)
    if placed.flip_h:
        decor["flipH"] = True
    if placed.flip_v:
        decor["flipV"] = True
    return decor


def _text(
    element_id: str,
    role: str,
    shape_id: int,
    placed: _Placement,
    look: dict[str, Any],
    paragraphs: list[dict[str, Any]],
    body: _Body,
) -> dict[str, Any]:
    return {
        "kind": "text",
        "elementId": element_id,
        "role": role,
        "box": placed.box(),
        **look,
        "sourceShapeId": shape_id,
        "paragraphs": paragraphs,
        "overflow": None,
        "verticalAlign": body.anchor,
        "textInsets": body.insets,
    }


def _image(
    slide: _Slide,
    element_id: str,
    shape_id: int,
    placed: _Placement,
    look: dict[str, Any],
    blip_fill: Any,
    part: Any,
    origin: _Origin,
) -> dict[str, Any] | None:
    blip = blip_fill.find(f"{_A}blip")
    if blip is None:
        return None
    media = _media_part(part, blip.get(f"{_R}embed"), slide)
    if media is None:
        logger.warning("picture %s on slide %d is linked or missing", element_id, slide.number)
        return None
    area = (placed.cx * placed.cy) / (slide.size[0] * slide.size[1])
    image: dict[str, Any] = {
        "kind": "image",
        "elementId": element_id,
        "role": "logo" if origin.kind != "slide" and area < _LOGO_SHARE else "image",
        "box": placed.box(),
        **look,
        "sourceShapeId": shape_id,
        "mediaPart": media,
        "fit": "contain",
    }
    source = blip_fill.find(f"{_A}srcRect")
    if source is not None:
        crop = {k: max(0.0, (_int(source.get(k), 0) or 0) / 100000) for k in ("l", "t", "r", "b")}
        if any(crop.values()) and crop["l"] + crop["r"] < 1 and crop["t"] + crop["b"] < 1:
            image["crop"] = {k: min(v, 0.999) for k, v in crop.items()}
            image["fit"] = "cover"
    return image


def _shape(el: Any, origin: _Origin, scope: _Scope, slide: _Slide) -> list[dict[str, Any]]:
    """Elements of one ``p:sp``: decoration, text, or a decoration with text on it."""
    nv = el.find(f"{_P}nvSpPr")
    c_nv = nv.find(f"{_P}cNvPr") if nv is not None else None
    shape_id = _int(c_nv.get("id"), 0) if c_nv is not None else 0
    shape_id = shape_id or 0
    c_nv_sp = nv.find(f"{_P}cNvSpPr") if nv is not None else None
    text_box = c_nv_sp is not None and _truthy(c_nv_sp.get("txBox"))
    ph = _placeholder_ref(el)
    where = f"{origin.kind} shape {shape_id} on slide {slide.number}"

    chain: list[tuple[Any, Any]] = [(el, origin.part)]
    kind = ""
    if ph is not None:
        if origin.kind != "slide":
            return []
        kind = ph.get("type", "")
        layout_ph = _match(
            kind or "obj", _int(ph.get("idx")), slide.layout.placeholders, by_idx=True
        )
        master_ph = _match(
            layout_ph.type if layout_ph else kind or "obj",
            None,
            slide.master.placeholders,
            by_idx=False,
        )
        if layout_ph is not None:
            chain.append((layout_ph.element, slide.layout.part))
        if master_ph is not None:
            chain.append((master_ph.element, slide.master.part))
        kind = kind or (layout_ph.type if layout_ph else "obj")

    props = [(c.find(f"{_P}spPr"), part) for c, part in chain]
    stated = (_xfrm(p.find(f"{_A}xfrm")) for p, _ in props if p is not None)
    placed = next((placement for placement in stated if placement is not None), None)
    if placed is None:
        logger.warning("%s has no position of its own or inherited; skipped", where)
        return []
    placed = _placed(placed, scope)
    look = _look(placed, scope, origin)
    geometry_props = next(
        (
            p
            for p, _ in props
            if p is not None
            and (p.find(f"{_A}prstGeom") is not None or p.find(f"{_A}custGeom") is not None)
        ),
        None,
    )
    styles = [c.find(f"{_P}style") for c, _ in chain]
    paint, paint_part = _shape_paint(
        [(p, part) for p, part in props if p is not None], styles, slide, scope
    )
    stroke = _shape_stroke([p for p, _ in props], styles, slide)
    outline, is_line = _outline(geometry_props, placed, where)

    tx_body = el.find(f"{_P}txBody")
    body = _body(c.find(f"{_P}txBody/{_A}bodyPr") for c, _ in chain)
    if ph is None:
        style_kind = "otherStyle"
    elif kind in _TITLE_TYPES:
        style_kind = "titleStyle"
    elif kind in _OTHER_STYLE_TYPES:
        style_kind = "otherStyle"
    else:
        style_kind = "bodyStyle"
    sources: list[Any] = [
        s for s in (slide.default_style, slide.master.styles.get(style_kind)) if s is not None
    ]
    font_ref = _font_ref(next((s for s in styles if s is not None), None), slide)
    if font_ref:
        sources.append(font_ref)
    sources.extend(
        lst
        for lst in (c.find(f"{_P}txBody/{_A}lstStyle") for c, _ in reversed(chain))
        if lst is not None
    )
    paragraphs = _paragraphs(tx_body, sources, slide, body) if tx_body is not None else []
    has_text = _has_text(paragraphs)

    base = f"{origin.kind}-{shape_id}"
    elements: list[dict[str, Any]] = []
    if paint.kind == "image":
        image = _image(
            slide,
            slide.unique(f"{base}-picture"),
            shape_id,
            placed,
            look,
            paint.blip,
            paint_part or origin.part,
            origin,
        )
        if image is not None:
            elements.append(image)
        paint = Paint("none")
    visible = paint.kind == "solid" or stroke is not None

    if is_line:
        if stroke is not None:
            elements.append(
                _decor(slide, slide.unique(base), shape_id, placed, look, outline, paint, stroke)
            )
        return elements
    role = _ROLE_OF_PLACEHOLDER.get(kind, "body")
    if ph is not None or text_box or not visible:
        if visible:
            elements.append(
                _decor(
                    slide,
                    slide.unique(f"{base}-shape"),
                    shape_id,
                    placed,
                    look,
                    outline,
                    paint,
                    stroke,
                )
            )
        if has_text:
            elements.append(
                _text(slide.unique(base), role, shape_id, placed, look, paragraphs, body)
            )
        return elements
    decor = _decor(slide, slide.unique(base), shape_id, placed, look, outline, paint, stroke)
    if has_text:
        decor["paragraphs"] = paragraphs
        decor["verticalAlign"] = body.anchor
        decor["textInsets"] = body.insets
    elements.append(decor)
    return elements


def _picture(el: Any, origin: _Origin, scope: _Scope, slide: _Slide) -> list[dict[str, Any]]:
    c_nv = el.find(f"{_P}nvPicPr/{_P}cNvPr")
    shape_id = (_int(c_nv.get("id"), 0) if c_nv is not None else 0) or 0
    chain = [el]
    ph = _placeholder_ref(el)
    if ph is not None:
        if origin.kind != "slide":
            return []
        layout_ph = _match(
            ph.get("type", "obj"), _int(ph.get("idx")), slide.layout.placeholders, by_idx=True
        )
        if layout_ph is not None:
            chain.append(layout_ph.element)
    stated = (_xfrm(c.find(f"{_P}spPr/{_A}xfrm")) for c in chain)
    placed = next((placement for placement in stated if placement is not None), None)
    blip_fill = el.find(f"{_P}blipFill")
    if placed is None or blip_fill is None:
        return []
    placed = _placed(placed, scope)
    image = _image(
        slide,
        slide.unique(f"{origin.kind}-{shape_id}"),
        shape_id,
        placed,
        _look(placed, scope, origin),
        blip_fill,
        origin.part,
        origin,
    )
    return [image] if image is not None else []


def _connector(el: Any, origin: _Origin, scope: _Scope, slide: _Slide) -> list[dict[str, Any]]:
    c_nv = el.find(f"{_P}nvCxnSpPr/{_P}cNvPr")
    shape_id = (_int(c_nv.get("id"), 0) if c_nv is not None else 0) or 0
    props = el.find(f"{_P}spPr")
    placed = _xfrm(props.find(f"{_A}xfrm")) if props is not None else None
    if placed is None:
        return []
    placed = _placed(placed, scope)
    stroke = _shape_stroke([props], [el.find(f"{_P}style")], slide)
    if stroke is None:
        return []
    where = f"{origin.kind} connector {shape_id} on slide {slide.number}"
    outline, _ = _outline(props, placed, where)
    if "geometry" not in outline:
        outline = {"shape": "line"}
    return [
        _decor(
            slide,
            slide.unique(f"{origin.kind}-{shape_id}"),
            shape_id,
            placed,
            _look(placed, scope, origin),
            outline,
            Paint("none"),
            stroke,
        )
    ]


def _frame(el: Any, origin: _Origin, scope: _Scope, slide: _Slide) -> list[dict[str, Any]]:
    """A table or a chart; any other graphic frame is reported and left out."""
    c_nv = el.find(f"{_P}nvGraphicFramePr/{_P}cNvPr")
    shape_id = (_int(c_nv.get("id"), 0) if c_nv is not None else 0) or 0
    placed = _xfrm(el.find(f"{_P}xfrm"))
    data = el.find(f"{_A}graphic/{_A}graphicData")
    where = f"{origin.kind} frame {shape_id} on slide {slide.number}"
    if placed is None or data is None:
        return []
    placed = _placed(placed, scope)
    base = {
        "elementId": slide.unique(f"{origin.kind}-{shape_id}"),
        "box": placed.box(),
        **_look(placed, scope, origin),
        "sourceShapeId": shape_id,
    }
    table = data.find(f"{_A}tbl")
    if table is not None:
        fields = table_fields(table, slide.colors, slide.color_value)
        if fields is None:
            return []
        return [{"kind": "table", "role": "table", **base, **fields}]
    chart = data.find(f"{_C}chart")
    if chart is not None:
        try:
            blob = origin.part.related_part(chart.get(f"{_R}id")).blob
        except (KeyError, AttributeError):
            logger.warning("%s names a chart part that is missing", where)
            return []
        fields, reason = chart_fields(blob)
        if fields is None:
            logger.warning("%s: chart left out, %s", where, reason)
            return []
        return [{"kind": "chart", "role": "chart", **base, **fields}]
    logger.warning("%s holds %s, which SlideIR cannot express", where, data.get("uri", "?"))
    return []


def _children(tree: Any) -> Iterator[Any]:
    """Shapes of a tree, with markup-compatibility blocks read through their fallback."""
    for child in tree:
        if not isinstance(child.tag, str):
            continue
        if child.tag == f"{_MC}AlternateContent":
            chosen = child.find(f"{_MC}Fallback")
            if chosen is None:
                chosen = child.find(f"{_MC}Choice")
            if chosen is not None:
                yield from _children(chosen)
            continue
        yield child


def _on_canvas(elements: list[dict[str, Any]], size: tuple[int, int]) -> list[dict[str, Any]]:
    """Leave out inherited graphics lying wholly past the slide edge.

    A master or a layout may keep a designer's rulers and notes off the slide; a show never paints
    them, so the slide does not carry them. A shape touching the edge is still on the slide.
    """
    width, height = size

    def seen(element: dict[str, Any]) -> bool:
        box = element.get("box")
        if not isinstance(box, dict):
            return True
        return not (
            box["x"] > width
            or box["y"] > height
            or box["x"] + box["cx"] < 0
            or box["y"] + box["cy"] < 0
        )

    return [element for element in elements if seen(element)]


def _walk(tree: Any, origin: _Origin, scope: _Scope, slide: _Slide) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    if tree is None:
        return elements
    for child in _children(tree):
        name = _local(child)
        c_nv = child.find(f"*/{_P}cNvPr")
        if c_nv is not None and _truthy(c_nv.get("hidden")):
            continue
        if name == "sp":
            elements.extend(_shape(child, origin, scope, slide))
        elif name == "pic":
            elements.extend(_picture(child, origin, scope, slide))
        elif name == "cxnSp":
            elements.extend(_connector(child, origin, scope, slide))
        elif name == "graphicFrame":
            elements.extend(_frame(child, origin, scope, slide))
        elif name == "grpSp":
            props = child.find(f"{_P}grpSpPr")
            group = _group_of(props.find(f"{_A}xfrm")) if props is not None else None
            shape_id = (_int(c_nv.get("id"), 0) if c_nv is not None else 0) or 0
            fill = fill_child(props)
            paint = paint_of(fill, slide.colors) if fill is not None else None
            if paint is not None and paint.kind == "group":
                paint = scope.group_paint
            inner = _Scope(
                groups=(*scope.groups, group) if group is not None else scope.groups,
                group_ids=(*scope.group_ids, slide.unique(f"{origin.kind}-group-{shape_id}")),
                group_paint=paint if paint is not None else scope.group_paint,
            )
            elements.extend(_walk(child, origin, inner, slide))
        elif name == "contentPart":
            logger.warning("ink on slide %d is not read", slide.number)
    return elements


# ------------------------------------------------------------------------------------ background


def _background(slide: _Slide) -> tuple[str | None, str | None]:
    """Colour and picture of the background: the first of slide, layout, master to state one."""
    for container, part in (
        (slide.element, slide.part),
        (slide.layout.element, slide.layout.part),
        (slide.master.element, slide.master.part),
    ):
        background = container.find(f"{_P}cSld/{_P}bg")
        if background is None:
            continue
        properties = background.find(f"{_P}bgPr")
        reference = background.find(f"{_P}bgRef")
        if properties is not None:
            fill = fill_child(properties)
            if fill is None:
                return None, None
            paint = paint_of(fill, slide.colors)
        elif reference is not None:
            paint = style_fill(reference, slide.colors)
            if paint.kind == "image":
                # A picture in the theme's background styles would be resolved against the theme
                # part; the colour named on the reference is what the slide reads as instead.
                color = slide.colors.first(reference)
                paint = Paint("solid", color[0]) if color else Paint("none")
        else:
            return None, None
        if paint.kind == "image":
            blip = paint.blip.find(f"{_A}blip")
            media = _media_part(part, blip.get(f"{_R}embed") if blip is not None else None, slide)
            if media is None:
                return None, None
            color = mean_color(slide.media[media])
            return (slide.color_value(color) if color else None), media
        if paint.kind == "solid" and paint.color:
            return slide.color_value(paint.color), None
        return None, None
    return None, None


# ----------------------------------------------------------------------------------------- deck


def _master_of(master: Any) -> _Master:
    try:
        theme_blob = master.part.part_related_by(RT.THEME).blob
    except (KeyError, AttributeError):
        theme_blob = None
    element = master._element
    styles = {}
    tx_styles = element.find(f"{_P}txStyles")
    if tx_styles is not None:
        for name in ("titleStyle", "bodyStyle", "otherStyle"):
            style = tx_styles.find(f"{_P}{name}")
            if style is not None:
                styles[name] = style
    return _Master(
        element=element,
        part=master.part,
        theme=read_theme(theme_blob),
        color_map=_color_map(element, override=False) or {},
        placeholders=_placeholders(_tree(element)),
        styles=styles,
    )


def _palette_lookup(palette: Iterable[dict[str, Any]]) -> Callable[[str], str]:
    roles: dict[str, str] = {}
    for entry in palette:
        roles.setdefault(str(entry["hex"]).upper(), str(entry["role"]))

    def color_value(hex_value: str) -> str:
        value = hex_value.upper()
        return roles.get(value, value)

    return color_value


def _notes(slide: Any) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
    except (AttributeError, KeyError):
        return ""
    return frame.text.strip() if frame is not None else ""


def read_deck(
    pptx_path: Path, palette: Iterable[dict[str, Any]], media_dir: Path | None = None
) -> list[dict[str, Any]]:
    """Every slide of a deck as SlideIR, colours named by the palette role they match.

    Args:
        pptx_path: The deck to read.
        palette: Palette entries (``role``, ``hex``) of the deck's own design system.
        media_dir: When given, every picture the slides draw is written there under the file name
            of its part, overwriting what is there.

    Returns:
        SlideIR payloads in presentation order, each validated against the contract.
    """
    presentation = Presentation(str(pptx_path))
    size = (int(presentation.slide_width or 0), int(presentation.slide_height or 0))
    root = presentation.part._element
    default_style = root.find(f"{_P}defaultTextStyle")
    first_number = _int(root.get("firstSlideNum"), 1) or 1
    color_value = _palette_lookup(palette)
    masters: dict[str, _Master] = {}
    layouts: dict[str, _Layout] = {}
    media: dict[str, bytes] = {}
    slides: list[dict[str, Any]] = []

    for index, slide in enumerate(presentation.slides):
        layout = slide.slide_layout
        master = layout.slide_master
        master_info = masters.get(str(master.part.partname))
        if master_info is None:
            master_info = masters.setdefault(str(master.part.partname), _master_of(master))
        layout_name = str(layout.part.partname)
        layout_info = layouts.get(layout_name)
        if layout_info is None:
            layout_info = layouts.setdefault(
                layout_name,
                _Layout(
                    element=layout._element,
                    part=layout.part,
                    placeholders=_placeholders(_tree(layout._element)),
                    color_map=_color_map(layout._element, override=True),
                ),
            )
        color_map = (
            _color_map(slide._element, override=True)
            or layout_info.color_map
            or master_info.color_map
        )
        context = _Slide(
            index=index,
            number=first_number + index,
            part=slide.part,
            element=slide._element,
            layout=layout_info,
            master=master_info,
            colors=Colors(master_info.theme, color_map),
            default_style=default_style,
            color_value=color_value,
            size=(max(size[0], 1), max(size[1], 1)),
            media=media,
        )
        elements: list[dict[str, Any]] = []
        if slide._element.get("showMasterSp") != "0":
            if layout._element.get("showMasterSp") != "0":
                elements += _on_canvas(
                    _walk(
                        _tree(master._element), _Origin("master", master.part), _Scope(), context
                    ),
                    context.size,
                )
            elements += _on_canvas(
                _walk(_tree(layout._element), _Origin("layout", layout.part), _Scope(), context),
                context.size,
            )
        elements += _walk(_tree(slide._element), _Origin("slide", slide.part), _Scope(), context)

        payload: dict[str, Any] = {
            "schemaVersion": SLIDE_IR_SCHEMA_VERSION,
            "slideId": f"imported-{index + 1}",
            "variant": "balanced",
            "index": index,
            "patternId": layout_name,
            "slideCx": size[0],
            "slideCy": size[1],
            "elements": elements,
        }
        background, picture = _background(context)
        if background is not None:
            payload["background"] = background
        if picture is not None:
            payload["backgroundImage"] = picture
        notes = _notes(slide)
        if notes:
            payload["notes"] = notes
        SlideIR.model_validate(payload)
        slides.append(payload)

    if media_dir is not None:
        media_dir.mkdir(parents=True, exist_ok=True)
        for name, blob in media.items():
            (media_dir / Path(name).name).write_bytes(blob)
    return slides
