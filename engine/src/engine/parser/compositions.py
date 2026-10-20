"""I recover reusable text compositions from actual slides without changing their geometry."""

import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from engine.parser._composition_geometry import (
    attach_image_frames,
    estimate_room,
    resolved_metrics,
    safe_room,
)
from engine.parser._text_style import inherited_shapes, text_metrics
from engine.parser.slides import read_deck

_Transform = tuple[float, float, float, float]
_IDENTITY: _Transform = (0.0, 0.0, 1.0, 1.0)


def _group_transform(group: Any, outer: _Transform) -> _Transform:
    """I map a group's child coordinate space onto the slide.

    A shape inside a group keeps its offset in the group's own child space (`chOff`/`chExt`).
    Read as is, cards of three different groups land on the same box, so reading order and the
    neighbour check compare shapes that are nowhere near each other on the slide.
    """
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    xfrm = group._element.grpSpPr.find(f"{ns}xfrm")
    if xfrm is None:
        return outer
    off, ext = xfrm.find(f"{ns}off"), xfrm.find(f"{ns}ext")
    ch_off, ch_ext = xfrm.find(f"{ns}chOff"), xfrm.find(f"{ns}chExt")
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return outer
    scale_x = int(ext.get("cx")) / max(1, int(ch_ext.get("cx")))
    scale_y = int(ext.get("cy")) / max(1, int(ch_ext.get("cy")))
    # Child point p lands at off + (p - chOff) * scale inside the parent, then the outer map.
    dx = int(off.get("x")) - int(ch_off.get("x")) * scale_x
    dy = int(off.get("y")) - int(ch_off.get("y")) * scale_y
    ox, oy, osx, osy = outer
    return (ox + dx * osx, oy + dy * osy, osx * scale_x, osy * scale_y)


def _shapes(shapes: Any, transform: _Transform = _IDENTITY) -> Any:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes, _group_transform(shape, transform))
        else:
            yield shape, transform


def _slide_box(shape: Any, transform: _Transform) -> dict[str, int]:
    dx, dy, sx, sy = transform
    return {
        "x": round(dx + (shape.left or 0) * sx),
        "y": round(dy + (shape.top or 0) * sy),
        "cx": round((shape.width or 0) * sx),
        "cy": round((shape.height or 0) * sy),
    }


_NUMBER = re.compile(r"^\s*\d{1,2}[.)]?\s*$")
_URL = re.compile(r"https?://|www\.", re.IGNORECASE)

#: Placeholders that hold copy. Picture, chart, table and media placeholders are holes for
#: content the pipeline does not generate; footer, date and number belong to the master.
_TEXT_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_TITLE,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}
_IMAGE_PLACEHOLDERS = {
    PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.BITMAP,
    PP_PLACEHOLDER.MEDIA_CLIP,
}
#: More pictures than this on one slide is an icon or logo sheet, not a slide design.
_PICTURE_SHEET = 12
#: An empty filled card covering this share of the slide, with nothing drawn on it, is a place the
#: template left for a picture; a smaller one is a pill or a badge.
_PANEL_SHARE = 0.08
#: A photo of the sample slide covering this share of the slide shows the template's own example
#: subject; a smaller picture is an icon, an avatar or a logo and stays.
_PHOTO_SHARE = 0.15
_NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_FILLS = ("solidFill", "gradFill", "pattFill", "blipFill")


def _is_filled_backdrop(shape: Any) -> bool:
    """I tell a filled shape without copy: a pill, a band or a card that text can sit on."""
    if shape.is_placeholder or (shape.has_text_frame and shape.text.strip()):
        return False
    properties = shape._element.find(f"{_NS_P}spPr")
    if properties is None:
        return False
    if any(properties.find(f"{_NS_A}{fill}") is not None for fill in _FILLS):
        return True
    if properties.find(f"{_NS_A}noFill") is not None:
        return False
    reference = shape._element.find(f"{_NS_P}style/{_NS_A}fillRef")
    return reference is not None and reference.get("idx", "0") != "0"


def _backdrops(slide: Any) -> list[dict[str, int]]:
    """I collect filled shapes drawn under slide copy, from the slide and from its layout."""
    boxes = []
    for container in (slide.slide_layout.shapes, slide.shapes):
        for shape, transform in _shapes(container):
            if _is_filled_backdrop(shape):
                boxes.append(_slide_box(shape, transform))
    return boxes


def _backdrop_width(box: dict[str, int], backdrops: list[dict[str, int]]) -> int | None:
    """I find how far a text box may run before it leaves the pill it starts on.

    Templates often set a title in a wide text box over a narrower coloured pill: the text is
    readable only on the pill. The pill is a filled shape that holds the start of the first line
    and is narrower than the box. Cards larger than the box do not limit it.
    """
    middle = box["y"] + box["cy"] // 2
    fitting = [
        b
        for b in backdrops
        if b["x"] <= box["x"] + box["cx"] // 20
        and b["x"] + b["cx"] > box["x"]
        and b["y"] <= middle <= b["y"] + b["cy"]
        and b["x"] + b["cx"] < box["x"] + box["cx"]
        and b["cy"] <= box["cy"] * 2
    ]
    if not fitting:
        return None
    narrowest = min(fitting, key=lambda b: b["x"] + b["cx"])
    return narrowest["x"] + narrowest["cx"] - box["x"]


def _centre_inside(box: dict[str, int], frame: dict[str, int]) -> bool:
    x, y = box["x"] + box["cx"] / 2, box["y"] + box["cy"] / 2
    return (
        frame["x"] <= x <= frame["x"] + frame["cx"] and frame["y"] <= y <= frame["y"] + frame["cy"]
    )


def _image_slots(placed: list[Any], slide_area: int) -> list[dict[str, Any]]:
    """I find where a picture goes: picture placeholders, large empty cards and sample photos.

    A picture placeholder says so itself. A large filled card with no copy and nothing drawn on it
    is the second kind: the template left a white frame where its author meant to put a picture,
    and a deck that keeps it shows a blank card. The third is a large photo on the sample slide
    itself: students at a laptop illustrate the template's example, not the deck's subject. Brand
    art lives on the layout and the master, which are not read here, and small pictures (icons,
    avatars, logos) stay.
    """
    boxes = [(shape, _slide_box(shape, transform)) for shape, transform in placed]
    slots = []
    for shape, box in boxes:
        # A small frame is a portrait, an icon or a QR code: nothing to illustrate.
        if box["cx"] * box["cy"] < _PANEL_SHARE * slide_area:
            continue
        if shape.is_placeholder and shape.placeholder_format.type in _IMAGE_PLACEHOLDERS:
            slots.append({"slotId": str(shape.shape_id), "kind": "placeholder", "box": box})
            continue
        if (
            shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            and box["cx"] * box["cy"] >= _PHOTO_SHARE * slide_area
            # Only a picture of the slide itself: one inside a group moves with the group.
            and shape._element.getparent().tag == f"{_NS_P}spTree"
        ):
            slots.append({"slotId": str(shape.shape_id), "kind": "photo", "box": box})
            continue
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or not _is_filled_backdrop(shape):
            continue
        drawn_on = any(
            other is not shape
            and _centre_inside(other_box, box)
            and other_box["cx"] * other_box["cy"] < box["cx"] * box["cy"]
            for other, other_box in boxes
        )
        if not drawn_on:
            slots.append({"slotId": str(shape.shape_id), "kind": "panel", "box": box})
    return slots


def _qr_candidates(placed: list[Any]) -> list[dict[str, Any]]:
    """Every picture and picture placeholder on the slide, whatever its size.

    A small picture is as likely a logo as a QR code, and this layer cannot tell them apart.
    It only names the shapes a code could replace. The caller decides which of them, if any,
    actually holds one. Panels are not candidates: a code replaces a picture, not an empty card.
    """
    candidates = []
    for shape, transform in placed:
        box = _slide_box(shape, transform)
        if shape.is_placeholder and shape.placeholder_format.type in _IMAGE_PLACEHOLDERS:
            kind = "placeholder"
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            kind = "photo"
        else:
            continue
        candidates.append({"slotId": str(shape.shape_id), "kind": kind, "box": box})
    return candidates


def _prompt(shape: Any, slide: Any) -> str | None:
    """I read the prompt an empty placeholder inherits from its layout or master.

    Many templates ship their real designs as slides of empty placeholders: the copy exists only
    as the layout's prompt ("Текст", "Заголовок"). Such a placeholder is a slot, and its prompt is
    the sample that tells the writer what the box expects. Only the first line counts: a master's
    "second level, third level" list describes outline levels, not content.
    """
    if not shape.is_placeholder or shape.placeholder_format.type not in _TEXT_PLACEHOLDERS:
        return None
    for parent in inherited_shapes(shape, slide)[1:]:
        if parent.has_text_frame and parent.text.strip():
            return parent.text.strip().splitlines()[0].strip()
    return ""


def _assign_single_title(slots: list[dict[str, Any]], slide_height: int) -> None:
    """I give a composition one title: the template's title placeholder, else its headline.

    The headline is the largest non-numeric text in the top third of the slide, set clearly larger
    than the rest of the slide's copy: templates differ in scale, so the test is relative, not a
    fixed point size. Card headers, numbers and paragraphs are body copy. Marking them as titles
    tells the writer to put a slide claim into a card label.
    """
    titles = [slot for slot in slots if slot["role"] == "title"]
    keep = None
    if titles:
        keep = min(titles, key=lambda s: (-s["fontSize"], s["box"]["y"], s["box"]["x"]))
    elif slots:
        words = [slot for slot in slots if not _NUMBER.match(slot["sample"])]

        def stands_out(slot: dict[str, Any]) -> bool:
            others = sorted(o["fontSize"] for o in words if o is not slot)
            if not others:
                return slot["fontSize"] >= 24
            return slot["fontSize"] >= 24 or slot["fontSize"] >= others[len(others) // 2] * 1.15

        candidates = [
            slot for slot in words if slot["box"]["y"] < slide_height * 0.3 and stands_out(slot)
        ]
        if candidates:
            keep = min(candidates, key=lambda s: (-s["fontSize"], s["box"]["y"], s["box"]["x"]))
    for slot in slots:
        slot["role"] = "title" if slot is keep else "body"


def _autofit(shape: Any, slide: Any) -> tuple[float, float]:
    """Export disables template shrink-to-fit; measure nominal fonts before accepting copy."""
    for item in inherited_shapes(shape, slide):
        body = item.text_frame._txBody.find(f"{_NS_A}bodyPr")
        if body is None:
            continue
        auto = next(
            (
                c
                for c in body
                if c.tag
                in {
                    f"{_NS_A}normAutofit",
                    f"{_NS_A}noAutofit",
                    f"{_NS_A}spAutoFit",
                }
            ),
            None,
        )
        if auto is None:
            continue
        if auto.tag == f"{_NS_A}normAutofit":
            return max(0.01, int(auto.get("fontScale", "100000")) / 100000), int(
                auto.get("lnSpcReduction", "0")
            ) / 100000
        break
    return 1.0, 0.0


def extract_compositions(path: Path) -> list[dict[str, Any]]:
    """I expose every editable text slot and its inherited typography to the writer."""
    presentation = Presentation(str(path))
    # Reuse the existing reader for inherited paragraph/list styles and grouped geometry.
    resolved_slides = read_deck(path, [])
    slide_height = presentation.slide_height
    if slide_height is None:
        raise ValueError("Template has no slide height.")
    patterns = []
    for index, slide in enumerate(presentation.slides):
        slots = []
        placed = list(_shapes(slide.shapes))
        shapes = [shape for shape, _ in placed]
        if any(s.has_chart or s.has_table for s in shapes):
            continue
        # Reference pages are not designs: a page of links and an icon or logo sheet would be
        # cloned with their underlined URLs and pictures.
        links = [r for r in slide.part.rels.values() if r.reltype.endswith("/hyperlink")]
        pictures = sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes)
        if links or pictures > _PICTURE_SHEET:
            continue
        if any(s.has_text_frame and _URL.search(s.text) for s in shapes):
            continue
        slide_area = (presentation.slide_width or 1) * slide_height
        backdrops = _backdrops(slide)
        resolved = resolved_slides[index]["elements"]
        image_slots = _image_slots(placed, slide_area)
        image_ids = {s["slotId"] for s in image_slots}
        for shape, transform in placed:
            if not shape.has_text_frame:
                continue
            if shape.is_placeholder and shape.placeholder_format.type in (
                PP_PLACEHOLDER.FOOTER,
                PP_PLACEHOLDER.DATE,
                PP_PLACEHOLDER.SLIDE_NUMBER,
            ):
                continue
            text = shape.text.strip()
            prompt = None if text else _prompt(shape, slide)
            if not text and prompt is None:
                continue
            box = _slide_box(shape, transform)
            # I preserve only actual page numbers, not repeated instructional/sample copy.
            if box["y"] > slide_height * 0.88 and text.isdigit():
                continue
            metrics = text_metrics(shape, slide)
            element = next(
                (
                    e
                    for e in resolved
                    if e.get("paragraphs")
                    and e.get("sourceShapeId") == shape.shape_id
                    and e["elementId"].startswith("slide-")
                ),
                None,
            )
            if element is not None:
                box = element["box"]
                font_scale, spacing_reduction = _autofit(shape, slide)
                metrics = resolved_metrics(
                    element, metrics, font_scale=font_scale, spacing_reduction=spacing_reduction
                )
            elif prompt is not None:
                # Empty placeholders have no visible copy in SlideIR. Their inherited metrics
                # still describe real room, rather than the length of the layout prompt.
                element = {
                    "elementId": f"slide-{shape.shape_id}",
                    "sourceShapeId": shape.shape_id,
                    "kind": "text",
                    "box": box,
                    "rotation": shape.rotation,
                }
            physical_chars, _ = estimate_room(metrics)
            reach = _backdrop_width(box, backdrops)
            if reach is not None:
                frame = shape.text_frame
                inner = (reach - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700
                metrics["textWidthPt"] = max(1.0, min(metrics["textWidthPt"], inner))
            obstacles: list[str] = []
            unknown = element is None
            if element is not None:
                metrics, obstacles, unknown = safe_room(element, metrics, resolved, image_ids)
            transforms = shape._element.xpath("ancestor-or-self::p:grpSp/p:grpSpPr/a:xfrm")
            unknown = unknown or any(
                t.get("rot", "0") != "0" or t.get("flipH", "0") == "1" or t.get("flipV", "0") == "1"
                for t in transforms
            )
            body = shape.text_frame._txBody.find(f"{_NS_A}bodyPr")
            unknown = unknown or (body is not None and body.get("vert", "horz") != "horz")
            paragraph_count = max(1, len(shape.text_frame.paragraphs))
            capacity, lines = estimate_room(metrics, paragraph_count)
            is_placeholder_title = shape.is_placeholder and shape.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE,
                PP_PLACEHOLDER.CENTER_TITLE,
            )
            # An unsupported transform/overlap cannot certify expanded copy. Short samples
            # alone never restrict a geometrically understood frame.
            basis = "obstacle-aware" if obstacles or reach is not None else "geometry"
            if unknown:
                basis = "sample-fallback"
                capacity = min(capacity, max(4, len(text or prompt or "")))
            assert capacity <= physical_chars
            slots.append(
                {
                    "slotId": str(shape.shape_id),
                    "role": "title" if is_placeholder_title else "body",
                    "sample": (text or prompt or "")[:220],
                    **({"prompt": True} if prompt is not None else {}),
                    "maxChars": capacity,
                    "maxLines": lines,
                    "capacity": {"safeChars": capacity, "basis": basis},
                    **metrics,
                    "box": box,
                }
            )
        _assign_single_title(slots, slide_height)
        image_count = pictures + sum(
            s.is_placeholder and s.placeholder_format.type in _IMAGE_PLACEHOLDERS for s in shapes
        )
        if len(slots) == 1 and slots[0]["fontSize"] < 32 and image_count:
            continue
        if slots and len(slots) <= 14:
            slots.sort(key=lambda s: (s["role"] != "title", s["box"]["y"], s["box"]["x"]))
            pattern = {
                "patternId": f"source-slide-{index}",
                "sourceIndex": index,
                "name": slide.slide_layout.name,
                "slots": slots,
                "imageCount": image_count,
                "imageSlots": image_slots,
                "qrCandidateSlots": _qr_candidates(placed),
            }
            attach_image_frames(pattern, resolved_slides[index])
            patterns.append(pattern)
    if not patterns:
        raise ValueError("Template contains no usable editable text compositions.")
    return patterns
