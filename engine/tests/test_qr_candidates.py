"""Small pictures are candidates, not automatically disposable content slots."""

from pathlib import Path
from typing import cast

from engine.export.compositions import export_compositions
from engine.parser.compositions import extract_compositions
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.util import Inches, Pt


def test_small_and_grouped_candidates_export_only_explicit_qr(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(7), Inches(1))
    title.text = "Contacts"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    image = tmp_path / "logo.png"
    Image.new("RGB", (30, 30), "red").save(image)
    logo = slide.shapes.add_picture(str(image), Inches(1), Inches(2), Inches(0.5), Inches(0.5))
    group = slide.shapes.add_group_shape()
    target = group.shapes.add_picture(str(image), Inches(2), Inches(3), Inches(1), Inches(1))
    group.left, group.top = Inches(4), Inches(2)
    group.width, group.height = Inches(2), Inches(2)
    source = tmp_path / "source.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    assert pattern["imageSlots"] == []
    candidates = {s["slotId"]: s for s in pattern["qrCandidateSlots"]}
    assert set(candidates) == {str(logo.shape_id), str(target.shape_id)}
    assert candidates[str(target.shape_id)]["box"] == {
        "x": Inches(4),
        "y": Inches(2),
        "cx": Inches(2),
        "cy": Inches(2),
    }
    item = {
        "patternId": pattern["patternId"],
        "texts": [{"slotId": s["slotId"], "text": "Contacts"} for s in pattern["slots"]],
    }
    export_compositions(source, tmp_path / "unchanged.pptx", [item], [pattern])
    unchanged = Presentation(str(tmp_path / "unchanged.pptx")).slides[0]
    assert any(
        s.name == logo.name and cast(Picture, s).image.blob == image.read_bytes()
        for s in unchanged.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    unchanged_group = cast(
        GroupShape, next(s for s in unchanged.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
    )
    assert cast(Picture, unchanged_group.shapes[0]).image.blob == image.read_bytes()
    (tmp_path / "media").mkdir()
    qr = tmp_path / "media" / "qr.png"
    Image.new("RGB", (64, 64), "black").save(qr)
    item["images"] = [
        {"slotId": str(target.shape_id), "purpose": "qr", "fit": "cover", "path": "media/qr.png"}
    ]
    export_compositions(source, tmp_path / "qr.pptx", [item], [pattern])
    result = Presentation(str(tmp_path / "qr.pptx")).slides[0]
    retained = cast(
        Picture, next(s for s in result.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    )
    assert retained.image.blob == image.read_bytes()
    result_group = cast(
        GroupShape, next(s for s in result.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
    )
    inserted = cast(Picture, result_group.shapes[0])
    assert inserted.image.blob == qr.read_bytes()
    assert (
        inserted.crop_left == inserted.crop_top == inserted.crop_right == inserted.crop_bottom == 0
    )


def test_small_picture_placeholder_is_a_candidate(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[8])
    title = slide.shapes.title
    assert title is not None
    title.text = "Contacts"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    picture = slide.placeholders[1]
    picture.width = picture.height = Inches(0.4)
    source = tmp_path / "placeholder.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    candidate = next(s for s in pattern["qrCandidateSlots"] if s["slotId"] == str(picture.shape_id))
    assert candidate["kind"] == "placeholder"
    assert candidate["box"]["cx"] == Inches(0.4)
    assert str(picture.shape_id) not in {s["slotId"] for s in pattern["imageSlots"]}
