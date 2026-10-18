"""A QR code keeps a rectangular frame even when the slot it replaces is masked.

An ordinary picture still inherits the slot's geometry. A QR code cannot: an ellipse or a
rounded card clips modules off the code. The picture file itself is a plain PNG made here; the
real code is produced on the TypeScript side and decoded from the exported deck there.
"""

from pathlib import Path
from typing import Any

import pytest
from engine.export.compositions import export_compositions
from engine.parser.compositions import extract_compositions
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _png(path: Path, size: tuple[int, int] = (40, 20)) -> Path:
    """A wide picture, so contain and cover land on different boxes."""
    Image.new("RGB", size, (12, 34, 56)).save(path)
    return path


def _deck(path: Path) -> tuple[Any, Any]:
    """A title and an ellipse the template left for a picture."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    title.text_frame.paragraphs[0].add_run().text = "Where the code goes"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    frame = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1.4), Inches(4), Inches(3))
    logo = _png(path.parent / "logo.png", (10, 10))
    slide.shapes.add_picture(str(logo), Inches(0.2), Inches(5), Inches(0.4), Inches(0.4))
    deck.save(str(path))
    return frame, title


def _export(tmp_path: Path, images: list[dict[str, str]]) -> Any:
    source = tmp_path / "source.pptx"
    frame, _photo = _deck(source)
    pattern = extract_compositions(source)[0]
    run = tmp_path / "run"
    (run / "media").mkdir(parents=True)
    _png(run / "media" / "code.png")
    texts = [{"slotId": slot["slotId"], "text": "Code"} for slot in pattern["slots"]]
    slot_id = next(slot["slotId"] for slot in pattern["imageSlots"])
    bound = [{**image, "slotId": slot_id} for image in images]
    export_compositions(
        source,
        run / "deck.pptx",
        [
            {
                "patternId": pattern["patternId"],
                "texts": texts,
                "speakerNotes": "",
                "images": bound,
            }
        ],
        [pattern],
    )
    return frame, Presentation(str(run / "deck.pptx")).slides[0]


def _picture(slide: Any) -> Any:
    pictures = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name.startswith("Illustration ")
    ]
    assert len(pictures) == 1
    return pictures[0]


def test_a_picture_without_a_purpose_keeps_the_slot_geometry(tmp_path: Path) -> None:
    frame, slide = _export(tmp_path, [{"slotId": "2", "path": "media/code.png", "fit": "cover"}])
    # The slot id is whatever the file assigned; the frame built above is the only image slot.
    source = extract_compositions(tmp_path / "source.pptx")[0]
    assert [slot["slotId"] for slot in source["imageSlots"]] == [str(frame.shape_id)]
    picture = _picture(slide)
    assert picture._element.spPr.find(f"{_A}prstGeom").get("prst") == "ellipse"
    assert (picture.left, picture.top, picture.width, picture.height) == (
        frame.left,
        frame.top,
        frame.width,
        frame.height,
    )
    assert picture.crop_left > 0 or picture.crop_top > 0
    assert all(shape.shape_id != frame.shape_id for shape in slide.shapes)


def test_a_qr_code_is_a_rectangle_fitted_inside_the_slot(tmp_path: Path) -> None:
    frame, slide = _export(
        tmp_path,
        [{"slotId": "2", "path": "media/code.png", "fit": "cover", "purpose": "qr"}],
    )
    picture = _picture(slide)
    geometry = picture._element.spPr.find(f"{_A}prstGeom")
    assert geometry is not None and geometry.get("prst") == "rect"
    assert picture._element.spPr.find(qn("a:custGeom")) is None
    # Cover was asked for and ignored: nothing is cropped, and the wide picture is letterboxed.
    assert picture.crop_left == picture.crop_right == picture.crop_top == picture.crop_bottom == 0
    assert picture.left == frame.left
    assert picture.width == frame.width
    assert picture.top > frame.top
    assert picture.height < frame.height
    assert picture.top + picture.height < frame.top + frame.height
    # The picture stays inside the slot, centred on it.
    centre = picture.left + picture.width / 2
    assert abs(centre - (frame.left + frame.width / 2)) <= 1
    embedded = picture.image.blob
    assert embedded == (tmp_path / "run" / "media" / "code.png").read_bytes()
    assert all(shape.shape_id != frame.shape_id for shape in slide.shapes)
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and shape.image.blob == (tmp_path / "logo.png").read_bytes()
        for shape in slide.shapes
    )


def test_qr_rejects_unknown_or_missing_shape(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    _deck(source)
    pattern = extract_compositions(source)[0]
    (tmp_path / "media").mkdir()
    _png(tmp_path / "media" / "code.png")
    item = {
        "patternId": pattern["patternId"],
        "texts": [{"slotId": s["slotId"], "text": "Code"} for s in pattern["slots"]],
        "images": [{"slotId": "999", "path": "media/code.png", "fit": "cover", "purpose": "qr"}],
    }
    with pytest.raises(ValueError, match="Unknown QR slot"):
        export_compositions(source, tmp_path / "out.pptx", [item], [pattern])
    pattern["qrCandidateSlots"].append({"slotId": "999"})
    with pytest.raises(ValueError, match="QR slot 999 was not found"):
        export_compositions(source, tmp_path / "out.pptx", [item], [pattern])


def test_a_qr_code_does_not_inherit_a_rounded_mask(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.7))
    title.text_frame.paragraphs[0].add_run().text = "Rounded card"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(1.5), Inches(3), Inches(4)
    )
    source = tmp_path / "source.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    run = tmp_path / "run"
    (run / "media").mkdir(parents=True)
    _png(run / "media" / "code.png", (30, 30))
    export_compositions(
        source,
        run / "deck.pptx",
        [
            {
                "patternId": pattern["patternId"],
                "texts": [{"slotId": slot["slotId"], "text": "Code"} for slot in pattern["slots"]],
                "speakerNotes": "",
                "images": [
                    {
                        "slotId": str(card.shape_id),
                        "path": "media/code.png",
                        "fit": "contain",
                        "purpose": "qr",
                    }
                ],
            }
        ],
        [pattern],
    )
    exported = Presentation(str(run / "deck.pptx")).slides[0]
    picture = _picture(exported)
    assert picture._element.spPr.find(f"{_A}prstGeom").get("prst") == "rect"
    # A square picture in a taller card is centred and not stretched to the card.
    assert picture.height < card.height
    assert abs(picture.width - card.width) <= 1
    assert picture.top > card.top
    centre = picture.top + picture.height / 2
    assert abs(centre - (card.top + card.height / 2)) <= 1
    assert picture.image.blob == (run / "media" / "code.png").read_bytes()
