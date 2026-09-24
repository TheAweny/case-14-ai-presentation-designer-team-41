"""Pattern catalog: reusable arrangements read off a template's layouts."""

from pathlib import Path
from typing import Any

from engine.contracts import PatternCatalog
from engine.parser import pattern_catalog_of
from pptx import Presentation


def _template(destination: Path, slides_on: tuple[int, ...] = (1, 1)) -> Path:
    """The stock python-pptx template, with slides built on the given layouts."""
    presentation = Presentation()
    for index in slides_on:
        presentation.slides.add_slide(presentation.slide_layouts[index])
    presentation.save(str(destination))
    return destination


def _without_placeholders(destination: Path) -> Path:
    """A template whose layouts place nothing, so no arrangement can be read off them."""
    presentation = Presentation()
    for layout in presentation.slide_layouts:
        tree = layout.shapes._spTree
        for element in tree.xpath("./p:sp[p:nvSpPr/p:nvPr/p:ph]"):
            tree.remove(element)
    presentation.save(str(destination))
    return destination


def _by_part(catalog: dict[str, Any], part_suffix: str) -> dict[str, Any]:
    return next(
        p for p in catalog["patterns"] if any(s.endswith(part_suffix) for s in p["sourceParts"])
    )


def test_catalog_matches_the_contract(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_template(tmp_path / "t.pptx"))

    PatternCatalog.model_validate(catalog)
    assert len(catalog["templateHash"]) == 64


def test_layouts_that_place_nothing_are_not_patterns(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_template(tmp_path / "t.pptx"))

    parts = [part for p in catalog["patterns"] for part in p["sourceParts"]]
    assert not any(part.endswith("slideLayout7.xml") for part in parts)


def test_layouts_with_the_same_arrangement_are_one_pattern(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_template(tmp_path / "t.pptx", slides_on=(1, 1)))

    shared = _by_part(catalog, "slideLayout2.xml")
    assert len(shared["sourceParts"]) >= 2
    # Every layout showing it, plus the two slides built on it.
    assert shared["occurrences"] == len(shared["sourceParts"]) + 2


def test_intents_follow_the_roles_and_their_placement(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_template(tmp_path / "t.pptx"))

    assert "title" in _by_part(catalog, "slideLayout1.xml")["intents"]
    assert "two-column" in _by_part(catalog, "slideLayout4.xml")["intents"]
    assert "two-column" not in _by_part(catalog, "slideLayout2.xml")["intents"]
    assert "image" in _by_part(catalog, "slideLayout9.xml")["intents"]


def test_text_slots_state_a_capacity_and_others_do_not(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_template(tmp_path / "t.pptx"))

    for pattern in catalog["patterns"]:
        for slot in pattern["slots"]:
            if slot["role"] in {"title", "subtitle", "body"}:
                assert slot["maxChars"] and slot["maxChars"] > 0, slot
            if slot["role"] in {"image", "chart", "table", "page-number"}:
                assert slot["maxChars"] is None, slot


def test_title_and_first_content_slot_are_required(tmp_path: Path) -> None:
    pattern = _by_part(pattern_catalog_of(_template(tmp_path / "t.pptx")), "slideLayout4.xml")

    required = [s["slotId"] for s in pattern["slots"] if not s["optional"]]
    assert required == ["title", "body-1"]


def test_catalog_is_stable_across_reads(tmp_path: Path) -> None:
    path = _template(tmp_path / "t.pptx")

    assert pattern_catalog_of(path) == pattern_catalog_of(path)


def test_template_without_placeholders_still_offers_one_pattern(tmp_path: Path) -> None:
    catalog = pattern_catalog_of(_without_placeholders(tmp_path / "t.pptx"))

    PatternCatalog.model_validate(catalog)
    (pattern,) = catalog["patterns"]
    assert [s["role"] for s in pattern["slots"]] == ["title", "body"]
