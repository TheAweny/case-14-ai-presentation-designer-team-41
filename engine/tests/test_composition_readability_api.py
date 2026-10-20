"""The composition endpoint includes file-based findings in the public audit report."""

from pathlib import Path

import pytest
from engine.api import app
from engine.export.compositions import export_compositions
from engine.parser.compositions import extract_compositions
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


@pytest.mark.parametrize("protected", [False, True])
def test_composition_audit_measures_actual_replacement(tmp_path: Path, protected: bool) -> None:
    media = tmp_path / "media"
    media.mkdir()
    Image.new("RGB", (160, 120), "#29394F").save(media / "source.png")
    Image.new("RGB", (160, 120), "#F8F4E3").save(media / "replacement.png")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(
        str(media / "source.png"), Inches(0), Inches(0), deck.slide_width, deck.slide_height
    )
    if protected:
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2), Inches(9.5), Inches(6)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor.from_string("29394F")
        panel.line.fill.background()
        panel._element.find("{*}style/{*}effectRef").set("idx", "0")
    for top, height, size in [(0.5, 0.8, 32), (2, 2.4, 20)]:
        shape = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(7), Inches(height))
        shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Исходный образец"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string("F8F4E3")
    source, output = tmp_path / "source.pptx", tmp_path / "exported.pptx"
    deck.save(str(source))
    patterns = extract_compositions(source)
    pattern = patterns[0]
    composed = {
        "patternId": pattern["patternId"],
        "texts": [
            {"slotId": slot["slotId"], "text": "Новый результат"} for slot in pattern["slots"]
        ],
        "speakerNotes": "",
        "images": [
            {
                "slotId": pattern["imageSlots"][0]["slotId"],
                "path": "media/replacement.png",
                "fit": "cover",
            }
        ],
    }
    export_compositions(source, output, [composed], patterns)
    response = TestClient(app).post(
        "/v1/composition-audit",
        json={
            "schemaVersion": 1,
            "templatePath": str(source),
            "outputPath": str(output),
            "slides": [composed],
        },
    )
    assert response.status_code == 200, response.text
    report = response.json()
    findings = [f for f in report["findings"] if f["checkId"] == "image-text-readability"]
    if protected:
        assert not findings
    else:
        assert findings
        assert all("status=fail" in f["evidence"]["measured"] for f in findings)
        assert {f["target"]["elementId"] for f in findings} <= {
            f"slot-{binding['slotId']}" for binding in composed["texts"]
        }
        assert not report["passed"]
        assert report["counts"]["major"] >= len(findings)
