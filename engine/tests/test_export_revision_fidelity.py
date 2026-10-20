"""Edited SlideIR must export to pptx and pdf with structure and typography intact."""

from __future__ import annotations

import shutil
import struct
import zlib
from pathlib import Path
from typing import cast

import pytest
from engine.export import export_pdf, export_pptx
from engine.parser import design_system_of, import_deck
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.shapes.autoshape import Shape
from pptx.util import Emu, Inches, Pt

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _png(colour: tuple[int, int, int]) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        payload = tag + data
        return (
            struct.pack(">I", len(data))
            + payload
            + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    data = zlib.compress(b"\x00" + bytes(colour))
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", data) + chunk(b"IEND", b"")
    )


def _paragraph_def_font(paragraph: object, font: str, size_pt: float) -> None:
    properties = paragraph._p.get_or_add_pPr()
    default = OxmlElement("a:defRPr")
    default.set("sz", str(int(size_pt * 100)))
    latin = OxmlElement("a:latin")
    latin.set("typeface", font)
    default.append(latin)
    properties.append(default)


def _deck(path: Path) -> Path:
    pixel = path.parent / "backdrop.png"
    pixel.write_bytes(_png((180, 40, 40)))
    presentation = Presentation()
    assert presentation.slide_width is not None and presentation.slide_height is not None
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(pixel),
        Emu(0),
        Emu(0),
        presentation.slide_width,
        presentation.slide_height,
    )
    overlay = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    overlay.text_frame.paragraphs[0].add_run().text = "ORIGINAL OVERLAY"
    overlay.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    nested = inner.shapes.add_textbox(Inches(0.3), Inches(0.3), Inches(3), Inches(0.8))
    paragraph = nested.text_frame.paragraphs[0]
    paragraph.add_run().text = "NESTED ORIGINAL"
    _paragraph_def_font(paragraph, "Synthetic Narrow", 22)

    presentation.save(str(path))
    return path


def _text_ink(image: Image.Image, box: dict[str, int], dpi: int) -> bool:
    scale = dpi / 914400
    left = max(0, int(box["x"] * scale))
    top = max(0, int(box["y"] * scale))
    right = min(image.width, int((box["x"] + box["cx"]) * scale))
    bottom = min(image.height, int((box["y"] + box["cy"]) * scale))
    for y in range(top, bottom):
        for x in range(left, right):
            pixel = image.getpixel((x, y))
            assert isinstance(pixel, tuple) and len(pixel) == 3
            red, green, blue = pixel
            if red < 80 and green < 80 and blue < 80:
                return True
    return False


def _nested_text_shape(presentation: Presentation) -> Shape:
    slide = presentation.slides[0]
    outer = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
    inner = next(shape for shape in outer.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
    return cast(Shape, next(shape for shape in inner.shapes if shape.has_text_frame))


def _run_text(element: dict) -> str:
    return element["paragraphs"][0]["runs"][0]["text"]


def _edited_revision_paths(tmp_path: Path) -> tuple[dict, dict, Path, Path]:
    source = _deck(tmp_path / "source.pptx")
    design = design_system_of(source)
    [imported] = import_deck(source)
    assert imported["index"] == 0

    overlay = next(
        element
        for element in imported["elements"]
        if element["kind"] == "text" and "ORIGINAL OVERLAY" in _run_text(element)
    )
    nested = next(
        element
        for element in imported["elements"]
        if element["kind"] == "text" and "NESTED ORIGINAL" in _run_text(element)
    )
    assert len(nested.get("groupIds", [])) == 2

    overlay["paragraphs"][0]["runs"][0]["text"] = "EDITED OVERLAY"
    nested["paragraphs"][0]["runs"][0]["text"] = "EDITED NESTED"
    assert nested["paragraphs"][0]["runs"][0].get("sizePt") == 22
    assert nested["paragraphs"][0]["runs"][0].get("font") == "Synthetic Narrow"

    pptx_path = tmp_path / "edited.pptx"
    export_pptx([imported], design, source, pptx_path)
    assert pptx_path.stat().st_size > 0

    pdf_path = tmp_path / "edited.pdf"
    if shutil.which("soffice") or shutil.which("libreoffice"):
        export_pdf(pptx_path, pdf_path)
        assert pdf_path.stat().st_size > 0

    [again] = import_deck(pptx_path)
    slide_images = [
        element
        for element in again["elements"]
        if element["kind"] == "image" and element.get("sourcePart") != "master"
    ]
    overlay_again = next(
        element
        for element in again["elements"]
        if element["kind"] == "text" and "EDITED OVERLAY" in _run_text(element)
    )
    nested_again = next(
        element
        for element in again["elements"]
        if element["kind"] == "text" and "EDITED NESTED" in _run_text(element)
    )
    assert slide_images
    assert again["elements"].index(slide_images[0]) < again["elements"].index(overlay_again)
    assert len(nested_again.get("groupIds", [])) == 2
    assert nested_again["paragraphs"][0]["runs"][0]["text"] == "EDITED NESTED"
    assert nested_again["paragraphs"][0]["runs"][0].get("sizePt") == 22
    assert nested_again["paragraphs"][0]["runs"][0].get("font") == "Synthetic Narrow"

    exported = Presentation(str(pptx_path))
    assert any(shape.shape_type == MSO_SHAPE_TYPE.GROUP for shape in exported.slides[0].shapes)
    nested_shape = _nested_text_shape(exported)
    run = nested_shape.text_frame.paragraphs[0].runs[0]
    assert run.text == "EDITED NESTED"
    assert run.font.name == "Synthetic Narrow"
    assert run.font.size == Pt(22)
    return overlay_again, nested_again, pptx_path, pdf_path


def test_export_revision_keeps_order_groups_typography_and_pdf(tmp_path: Path) -> None:
    _edited_revision_paths(tmp_path)
    pdf_path = tmp_path / "edited.pdf"
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        assert not pdf_path.exists()
        return
    assert pdf_path.stat().st_size > 0


def test_export_revision_pdf_keeps_text_visible_over_picture(tmp_path: Path) -> None:
    overlay_again, _nested, _pptx_path, pdf_path = _edited_revision_paths(tmp_path)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None or not shutil.which("pdftoppm"):
        pytest.skip("LibreOffice or pdftoppm is missing; structural checks still run.")
    assert pdf_path.stat().st_size > 0
    from engine.render import rasterize_pdf

    dpi = 72
    pngs = rasterize_pdf(pdf_path, tmp_path / "png", dpi=dpi)
    assert len(pngs) == 1
    with Image.open(pngs[0]) as image:
        rgb = image.convert("RGB")
        assert _text_ink(rgb, overlay_again["box"], dpi)
