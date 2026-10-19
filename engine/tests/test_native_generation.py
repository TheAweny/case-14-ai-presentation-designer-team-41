"""Exercise the composition API with the native builder, not a hand-drawn slide."""

import importlib
from pathlib import Path
from typing import Any, cast

import pytest
from engine.contracts import CompositionExportRequest
from engine.parser.compositions import extract_compositions
from pptx import Presentation
from pptx.util import Inches, Pt

api = importlib.import_module("engine.api.app")


@pytest.mark.parametrize("variant", ["dense", "balanced", "airy"])
def test_generation_exports_and_audits_native_kpi(
    tmp_path: Path, monkeypatch: Any, variant: str
) -> None:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.33), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for text, y, height, size in [("Pilot", 0.3, 0.8, 32), ("Revenue", 1.5, 5.5, 16)]:
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12), Inches(height))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text, run.font.name, run.font.size = text, "DejaVu Sans", Pt(size)
    template = tmp_path / "template.pptx"
    presentation.save(str(template))
    pattern = extract_compositions(template)[0]
    target = pattern["slots"][-1]["slotId"]
    block = {
        "kind": "kpi",
        "items": [
            {
                "value": {"state": "present", "literal": "18", "value": 18},
                "caption": "Revenue",
                "unit": "%",
                "unitAmbiguous": False,
                "period": "2026",
                "periodKnown": True,
                "sources": [
                    {"itemId": "brief", "quote": "Revenue 18 % in 2026", "status": "grounded"}
                ],
                "caveat": "Pilot only",
            }
        ],
    }
    request = CompositionExportRequest.model_validate(
        {
            "schemaVersion": 1,
            "templatePath": str(template),
            "outputPath": str(tmp_path / "deck.pptx"),
            "slides": [
                {
                    "patternId": pattern["patternId"],
                    "speakerNotes": "Pilot only",
                    "texts": [
                        {"slotId": s["slotId"], "text": "" if s["slotId"] == target else "Pilot"}
                        for s in pattern["slots"]
                    ],
                    "visual": {
                        "kind": "native",
                        "slotId": target,
                        "block": block,
                        "variant": variant,
                    },
                }
            ],
        }
    )
    monkeypatch.setattr(api.render, "render_to_png", lambda *a, **kw: [])
    monkeypatch.setattr(api.render, "missing_fonts", lambda *a, **kw: [])
    monkeypatch.setattr(api, "bundle_previews", lambda *a, **kw: None)
    api.composition_deck(request)
    exported = Presentation(str(tmp_path / "deck.pptx"))
    texts = "\n".join(cast(Any, s).text for s in exported.slides[0].shapes if s.has_text_frame)
    assert "18 %" in texts and "2026" in texts and "Pilot only" in texts
    report = api.composition_audit(request)
    assert report.schemaVersion == 1
