"""Three variants of one plan: different in composition, identical in what they say.

The axis is density. Dense takes the fullest layouts of the template, its smallest body step and
two columns for a long list; airy the sparsest layouts, the largest step that fits and centred
short copy. Every variant keeps every word of the plan and stays on the template.
"""

from collections import Counter
from pathlib import Path
from typing import Any

import pytest
from engine.api import app
from engine.api._profiles import deterministic_profile
from engine.checks import audit_deck
from engine.contracts import DeckPlan, LayoutVariantsResponse, SlideIR
from engine.layout import VARIANTS, build_variants, fingerprint, select_pattern
from engine.parser import template_profile_of
from fastapi.testclient import TestClient
from pptx import Presentation

_LIST = [
    "Parses layouts, palette and type scale",
    "Finds the patterns the template repeats",
    "Plans the deck once",
    "Lays out three variants from one plan",
    "Audits the result",
    "Exports PPTX, PDF and HTML",
]


def _slide(i: int, intent: str, blocks: list[dict[str, Any]], subhead: str = "") -> dict[str, Any]:
    return {
        "slideId": f"s{i}",
        "headline": "Templates are read as rules" if i else "Presentation designer",
        "subhead": subhead,
        "intent": intent,
        "blocks": blocks,
        "speakerNotes": "",
        "sources": [],
    }


def _plan(*slides: dict[str, Any]) -> DeckPlan:
    return DeckPlan.model_validate(
        {
            "schemaVersion": 1,
            "planId": "plan",
            "briefId": "brief",
            "title": "Deck",
            "language": "en",
            "slides": list(slides),
        }
    )


def _deck() -> DeckPlan:
    return _plan(
        _slide(0, "title", [], subhead="Team 41"),
        _slide(
            1, "statement", [{"kind": "paragraph", "text": "Decks are still laid out by hand."}]
        ),
        _slide(2, "bullets", [{"kind": "bullets", "items": _LIST}]),
        _slide(3, "bullets", [{"kind": "bullets", "items": _LIST[:3]}]),
        _slide(4, "section", []),
    )


@pytest.fixture(scope="module")
def profile(tmp_path_factory: pytest.TempPathFactory) -> tuple[dict[str, Any], dict[str, Any]]:
    path = tmp_path_factory.mktemp("template") / "t.pptx"
    Presentation().save(str(path))
    return template_profile_of(path)


def _words(slides: list[SlideIR]) -> Counter[str]:
    return Counter(
        run["text"]
        for slide in slides
        for element in slide.model_dump(mode="json")["elements"]
        for paragraph in element.get("paragraphs") or []
        for run in paragraph["runs"]
    )


def test_three_variants_differ_in_composition(profile: tuple[dict[str, Any], ...]) -> None:
    design_system, patterns = profile
    variants = build_variants(_deck(), design_system, patterns)
    prints = {name: fingerprint(slides) for name, slides in variants.items()}
    assert set(variants) == set(VARIANTS)
    assert len(set(prints.values())) == 3


def test_every_variant_says_exactly_the_same(profile: tuple[dict[str, Any], ...]) -> None:
    design_system, patterns = profile
    plan = _deck()
    variants = build_variants(plan, design_system, patterns)
    words = [_words(slides) for slides in variants.values()]
    assert words[0] == words[1] == words[2]
    for slide in plan.slides:
        assert words[0][slide.headline] >= 1


def test_dense_splits_a_long_list_into_two_columns(profile: tuple[dict[str, Any], ...]) -> None:
    design_system, patterns = profile
    plan = _plan(_slide(1, "bullets", [{"kind": "bullets", "items": _LIST}]))
    catalog = {"patterns": [p for p in patterns["patterns"] if p["patternId"] == "title-body"]}
    (slide,) = build_variants(plan, design_system, catalog)["dense"]
    bodies = [e for e in slide.model_dump(mode="json")["elements"] if e["role"] == "body"]
    assert len(bodies) == 2
    assert bodies[0]["box"]["y"] == bodies[1]["box"]["y"]
    assert bodies[0]["box"]["x"] < bodies[1]["box"]["x"]


def test_airy_sets_copy_no_smaller_than_dense(profile: tuple[dict[str, Any], ...]) -> None:
    design_system, patterns = profile
    variants = build_variants(_deck(), design_system, patterns)

    def body_sizes(slides: list[SlideIR]) -> list[float]:
        return [
            run["sizePt"]
            for slide in slides
            for element in slide.model_dump(mode="json")["elements"]
            if element["role"] == "body"
            for paragraph in element["paragraphs"]
            for run in paragraph["runs"]
        ]

    assert min(body_sizes(variants["airy"])) >= max(body_sizes(variants["dense"]))


def test_a_slide_never_lands_on_a_pattern_without_room_for_its_headline() -> None:
    body_only = {
        "patternId": "body-only",
        "intents": ["statement"],
        "occurrences": 9,
        "density": 0.5,
        "slots": [
            {
                "slotId": "body",
                "role": "body",
                "optional": False,
                "maxChars": None,
                "box": {"x": 0, "y": 0, "cx": 100, "cy": 100},
            }
        ],
    }
    titled = {
        **body_only,
        "patternId": "titled",
        "occurrences": 1,
        "slots": [
            {**body_only["slots"][0], "slotId": "title", "role": "title"},
            body_only["slots"][0],
        ],
    }
    plan = _plan(_slide(1, "statement", [{"kind": "paragraph", "text": "Copy."}]))
    for variant in VARIANTS:
        chosen = select_pattern(plan.slides[0], {"patterns": [body_only, titled]}, variant)
        assert chosen["patternId"] == "titled"


def test_what_the_layout_calls_fitting_the_audit_does_too(
    profile: tuple[dict[str, Any], ...],
) -> None:
    design_system, patterns = profile
    for slides in build_variants(_deck(), design_system, patterns).values():
        dumped = [s.model_dump(mode="json", exclude_unset=True) for s in slides]
        fitting = {
            e["elementId"] for s in dumped for e in s["elements"] if e.get("overflow") is None
        }
        report = audit_deck(
            {
                "deckId": "d",
                "slides": dumped,
                "designSystem": design_system,
                "patternCatalog": patterns,
            },
            "default",
            deterministic_profile("default"),
        )
        flagged = {
            f["target"]["elementId"] for f in report["findings"] if f["checkId"] == "text-overflow"
        }
        assert not fitting & flagged


def test_variants_are_deterministic(profile: tuple[dict[str, Any], ...]) -> None:
    design_system, patterns = profile
    first = build_variants(_deck(), design_system, patterns)
    second = build_variants(_deck(), design_system, patterns)
    assert {k: fingerprint(v) for k, v in first.items()} == {
        k: fingerprint(v) for k, v in second.items()
    }


def test_template_path_is_not_needed_for_variants(tmp_path: Path) -> None:
    path = tmp_path / "t.pptx"
    Presentation().save(str(path))
    design_system, _ = template_profile_of(path)
    variants = build_variants(_deck(), design_system)
    assert all(len(slides) == 5 for slides in variants.values())


def test_layouts_endpoint_returns_the_three_variants(tmp_path: Path) -> None:
    path = tmp_path / "t.pptx"
    Presentation().save(str(path))
    payload = {
        "schemaVersion": 1,
        "templatePath": str(path),
        "plan": _deck().model_dump(mode="json"),
    }
    response = TestClient(app).post("/v1/layouts", json=payload)
    assert response.status_code == 200
    body = LayoutVariantsResponse.model_validate(response.json())
    assert [str(v.variant.value) for v in body.variants] == list(VARIANTS)

    one = TestClient(app).post("/v1/layouts", json={**payload, "variants": ["airy"]})
    assert [v["variant"] for v in one.json()["variants"]] == ["airy"]


def test_layouts_endpoint_leaves_unset_optional_fields_out(tmp_path: Path) -> None:
    path = tmp_path / "t.pptx"
    Presentation().save(str(path))
    payload = {
        "schemaVersion": 1,
        "templatePath": str(path),
        "plan": _deck().model_dump(mode="json"),
    }
    body = TestClient(app).post("/v1/layouts", json=payload).json()
    element = body["variants"][0]["slides"][0]["elements"][0]
    assert "rotation" not in element and "groupIds" not in element
    assert "overflow" in element
