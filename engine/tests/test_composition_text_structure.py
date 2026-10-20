"""Synthetic round trips for paragraph structure in composition text replacement."""

import shutil
import subprocess
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from engine.export.compositions import _apply_safe_text_area, _replace_text, export_compositions
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.text import CT_TextParagraphProperties
from pptx.oxml.xmlchemy import BaseOxmlElement, OxmlElement
from pptx.shapes.autoshape import Shape
from pptx.text.text import _Paragraph
from pptx.util import Inches, Pt

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _xml(element: BaseOxmlElement | None) -> str:
    assert element is not None
    return deepcopy(element).xml


def _properties(paragraph: _Paragraph) -> CT_TextParagraphProperties:
    properties = paragraph._p.pPr
    assert properties is not None
    return properties


@pytest.fixture
def text_box() -> tuple[Any, Shape]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(6.3), Inches(4.1))
    return deck, shape


def _bullet(paragraph: Any, kind: str, **attributes: str) -> Any:
    marker = OxmlElement(f"a:{kind}")
    marker.attrib.update(attributes)
    properties = paragraph._p.get_or_add_pPr()
    default = properties.find(f"{_A}defRPr")
    properties.insert(properties.index(default) if default is not None else len(properties), marker)
    return marker


def _round_trip(deck: Any, shape: Shape, text: str) -> Shape:
    _replace_text(shape, text)
    saved = BytesIO()
    deck.save(saved)
    saved.seek(0)
    return cast(Shape, Presentation(saved).slides[0].shapes[0])


def test_paragraph_order_levels_spacing_and_list_styles_survive_export(
    text_box: tuple[Any, Shape], tmp_path: Path
) -> None:
    deck, shape = text_box
    shape.text = "Heading\nSample bullet\nSample number"
    originals = shape.text_frame.paragraphs
    for index, paragraph in enumerate(originals):
        paragraph.level = index
        paragraph.alignment = (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT)[index]
        paragraph.space_before = Pt(index * 3)
        paragraph.space_after = Pt(9 - index)
        paragraph.line_spacing = 1.2 + index / 10
        paragraph._p.get_or_add_pPr().set("marL", str(210000 + index * 31000))
        paragraph._p.get_or_add_pPr().set("indent", str(-12000 - index * 1000))
        run = paragraph.runs[0]
        run.font.name = f"Synthetic Face {index}"
        run.font.size = Pt(26 - index * 4)
        run.font.bold = index == 0
        run.font.color.rgb = RGBColor(31 + index * 23, 77, 131)
        paragraph._p.get_or_add_endParaRPr().set("lang", "ru-RU")
    _bullet(originals[0], "buNone")
    _bullet(originals[1], "buChar", char="◦")
    _bullet(originals[2], "buAutoNum", type="arabicParenR", startAt="4")
    properties = [_xml(p._p.pPr) for p in originals]
    run_properties = [_xml(p.runs[0]._r.rPr) for p in originals]
    source, output = tmp_path / "source.pptx", tmp_path / "output.pptx"
    deck.save(str(source))
    source_bytes = source.read_bytes()
    text = "Заголовок\nУсловие не изменено\nРост на 17% за 2026 год"
    slot_id = str(shape.shape_id)
    export_compositions(
        source,
        output,
        [{"patternId": "synthetic", "texts": [{"slotId": slot_id, "text": text}]}],
        [
            {
                "patternId": "synthetic",
                "sourceIndex": 0,
                "slots": [{"slotId": slot_id, "maxChars": len(text)}],
            }
        ],
    )
    exported = cast(Shape, Presentation(str(output)).slides[0].shapes[0])
    assert source.read_bytes() == source_bytes
    assert exported.text == text
    assert [_xml(p._p.pPr) for p in exported.text_frame.paragraphs] == properties
    assert [_xml(p.runs[0]._r.rPr) for p in exported.text_frame.paragraphs] == run_properties
    assert [p.level for p in exported.text_frame.paragraphs] == [0, 1, 2]


@pytest.mark.parametrize(
    "text",
    [
        "Первый\vпродолжение\nВторой\vещё строка",
        "\vНачало\v\vКонец\v\n\nПоследний\n",
        "\n\n",
        "",
    ],
)
def test_soft_breaks_and_empty_paragraphs_round_trip_without_reordering(
    text_box: tuple[Any, Shape], text: str
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Образец\vпродолжение"
    paragraph.level = 2
    paragraph.runs[0].font.size = Pt(19)
    paragraph._p.get_or_add_endParaRPr().set("lang", "ru-RU")
    _bullet(paragraph, "buChar", char="▪")
    exported = _round_trip(deck, shape, text)
    assert exported.text == text
    assert len(exported.text_frame.paragraphs) == len(text.split("\n"))
    for actual, expected in zip(exported.text_frame.paragraphs, text.split("\n"), strict=True):
        assert actual.text == expected
        assert len(actual._p.findall(f"{_A}br")) == expected.count("\v")
        assert actual.level == 2
        assert _properties(actual).find(f"{_A}buChar").get("char") == "▪"
        assert actual._p[-1].tag == f"{_A}endParaRPr"
        assert all(run.font.size == Pt(19) for run in actual.runs)


@pytest.mark.parametrize("kind", ["buChar", "buAutoNum", "inherited"])
@pytest.mark.parametrize("prefix", ["• ", "- ", "* ", "1. ", "03) ", "(2) ", "а) ", "IV. "])
def test_explicit_markers_do_not_duplicate_native_or_inherited_markers(
    text_box: tuple[Any, Shape], kind: str, prefix: str
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Образец"
    paragraph.level = 1
    paragraph.font.size = Pt(21)
    if kind == "inherited":
        style = shape.text_frame._txBody.find(f"{_A}lstStyle")
        level = OxmlElement("a:lvl2pPr")
        marker = OxmlElement("a:buChar")
        marker.set("char", "•")
        level.append(marker)
        style.append(level)
    else:
        _bullet(
            paragraph, kind, **({"char": "•"} if kind == "buChar" else {"type": "arabicPeriod"})
        )
    text = prefix + "Не менее 12 проектов"
    exported = _round_trip(deck, shape, text)
    actual = exported.text_frame.paragraphs[0]
    assert exported.text == text
    assert actual.level == 1
    assert actual.font.size == Pt(21)
    properties = _properties(actual)
    assert [
        child.tag
        for child in properties
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum", "buBlip"}
    ] == [f"{_A}buNone"]
    assert properties.index(properties.find(f"{_A}buNone")) < properties.index(
        properties.find(f"{_A}defRPr")
    )


def test_template_defined_marker_is_not_duplicated(text_box: tuple[Any, Shape]) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Образец"
    _bullet(paragraph, "buChar", char="→")
    text = "→ Следующий шаг"
    exported = _round_trip(deck, shape, text)
    assert exported.text == text
    assert _properties(exported.text_frame.paragraphs[0]).find(f"{_A}buNone") is not None


@pytest.mark.parametrize(
    "text",
    [
        "-12% за год",
        "−12% за год",
        "1.5 млн рублей",
        "2026 год",
        "I квартал",
        "а)без пробела",
        "Условие\v• Сохраняем знак внутри абзаца",
    ],
)
def test_content_prefixes_and_soft_continuations_keep_native_numbering(
    text_box: tuple[Any, Shape], text: str
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Образец"
    _bullet(paragraph, "buAutoNum", type="romanUcPeriod", startAt="3")
    before = _xml(paragraph._p.pPr)
    exported = _round_trip(deck, shape, text)
    assert exported.text == text
    assert _xml(exported.text_frame.paragraphs[0]._p.pPr) == before


def test_extra_paragraphs_reuse_last_style_without_restarting_numbering(
    text_box: tuple[Any, Shape],
) -> None:
    deck, shape = text_box
    shape.text = "Заголовок\nНачало списка"
    heading, item = shape.text_frame.paragraphs
    heading.runs[0].font.size = Pt(29)
    item.runs[0].font.size = Pt(17)
    item.level = 2
    _bullet(item, "buAutoNum", type="arabicPeriod", startAt="7")
    text = "Новый заголовок\nНачало\nПродолжение\nЗавершение"
    exported = _round_trip(deck, shape, text)
    paragraphs = exported.text_frame.paragraphs
    assert exported.text == text
    assert [p.level for p in paragraphs] == [0, 2, 2, 2]
    assert [p.runs[0].font.size for p in paragraphs] == [Pt(29), Pt(17), Pt(17), Pt(17)]
    numbers = [_properties(p).find(f"{_A}buAutoNum") for p in paragraphs[1:]]
    assert [n.get("type") for n in numbers] == ["arabicPeriod"] * 3
    assert [n.get("startAt") for n in numbers] == ["7", None, None]


def test_fewer_paragraphs_remove_unused_samples(text_box: tuple[Any, Shape]) -> None:
    deck, shape = text_box
    shape.text = "Первый\nВторой\nТретий"
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        paragraph.level = index
    exported = _round_trip(deck, shape, "Без сокращения смысла")
    assert exported.text == "Без сокращения смысла"
    assert len(exported.text_frame.paragraphs) == 1
    assert exported.text_frame.paragraphs[0].level == 0


def test_soft_lines_keep_typography_and_drop_previous_hyperlinks(
    text_box: tuple[Any, Shape],
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Старая ссылка"
    run.font.size = Pt(23)
    run.font.italic = True
    run.hyperlink.address = "https://example.org/old"
    exported = _round_trip(deck, shape, "Новая строка\vпродолжение")
    for actual in exported.text_frame.paragraphs[0].runs:
        assert actual.font.size == Pt(23)
        assert actual.font.italic is True
        assert actual.hyperlink.address is None


def test_empty_placeholder_uses_end_paragraph_typography_for_every_soft_line(
    text_box: tuple[Any, Shape],
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    end = paragraph._p.get_or_add_endParaRPr()
    end.set("sz", "2400")
    end.set("b", "1")
    exported = _round_trip(deck, shape, "Первая\vВторая")
    actual = exported.text_frame.paragraphs[0]
    assert actual.text == "Первая\vВторая"
    assert len(actual.runs) == 2
    assert all(run.font.size == Pt(24) and run.font.bold for run in actual.runs)
    assert actual._p[-1].tag == f"{_A}endParaRPr"


@pytest.mark.parametrize("newline", ["\r\n", "\r", "\n"])
def test_paragraph_endings_are_normalized_without_touching_soft_breaks(
    text_box: tuple[Any, Shape], newline: str
) -> None:
    deck, shape = text_box
    shape.text = "Первый\nВторой"
    shape.text_frame.paragraphs[1].space_before = Pt(11)
    exported = _round_trip(deck, shape, f"Один\vпродолжение{newline}Два{newline}")
    assert exported.text == "Один\vпродолжение\nДва\n"
    assert len(exported.text_frame.paragraphs) == 3
    assert exported.text_frame.paragraphs[1].space_before == Pt(11)


def test_normal_autofit_cannot_shrink_measured_font_or_line_spacing(
    text_box: tuple[Any, Shape],
) -> None:
    deck, shape = text_box
    shape.text = "Образец"
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.runs[0].font.size = Pt(27)
    paragraph.line_spacing = 1.3
    body = shape.text_frame._txBody.bodyPr
    for child in list(body):
        body.remove(child)
    shrink = OxmlElement("a:normAutofit")
    shrink.set("fontScale", "65000")
    shrink.set("lnSpcReduction", "20000")
    body.append(shrink)
    exported = _round_trip(deck, shape, "Новый текст\vпродолжение")
    assert exported.text_frame._txBody.bodyPr.find(f"{_A}normAutofit") is None
    assert exported.text_frame._txBody.bodyPr.find(f"{_A}noAutofit") is not None
    actual = exported.text_frame.paragraphs[0]
    assert all(run.font.size == Pt(27) for run in actual.runs)
    assert actual.line_spacing == 1.3


def test_missing_local_autofit_explicitly_disables_inherited_shrinking(
    text_box: tuple[Any, Shape],
) -> None:
    deck, shape = text_box
    body = shape.text_frame._txBody.bodyPr
    for child in list(body):
        body.remove(child)
    exported = _round_trip(deck, shape, "Текст без уменьшения кегля")
    assert exported.text_frame._txBody.bodyPr.find(f"{_A}noAutofit") is not None


@pytest.mark.parametrize("scale", [(1.0, 1.0), (2.0, 1.5), (0.5, 0.75)])
@pytest.mark.parametrize("indent", [-8, 6])
def test_safe_area_keeps_geometry_and_charges_indent_once(
    text_box: tuple[Any, Shape], scale: tuple[float, float], indent: int
) -> None:
    _, shape = text_box
    frame = shape.text_frame
    frame.margin_left, frame.margin_top = Pt(10), Pt(12)
    frame.margin_right, frame.margin_bottom = Pt(20), Pt(16)
    paragraph = frame.paragraphs[0]
    paragraph.text = "Sample"
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(Pt(18)))
    properties.set("indent", str(Pt(indent)))
    original_properties = _xml(properties)
    geometry = (shape.left, shape.top, shape.width, shape.height)
    sx, sy = scale
    slot = {
        "box": {"cx": round(shape.width * sx), "cy": round(shape.height * sy)},
        "textWidthPt": 100 * sx,
        "textHeightPt": 90 * sy,
    }
    _apply_safe_text_area(shape, slot)
    assert frame.margin_right == shape.width - Pt(10 + 18 + max(0, indent) + 100)
    assert frame.margin_bottom == shape.height - Pt(12 + 90)
    assert (shape.left, shape.top, shape.width, shape.height) == geometry
    assert frame.margin_left == Pt(10) and frame.margin_top == Pt(12)
    assert _xml(properties) == original_properties
    margins = frame.margin_right, frame.margin_bottom
    _apply_safe_text_area(shape, slot)
    assert (frame.margin_right, frame.margin_bottom) == margins


def test_safe_area_resolves_inherited_list_indent_and_never_expands_original_room(
    text_box: tuple[Any, Shape],
) -> None:
    _, shape = text_box
    frame = shape.text_frame
    frame.margin_left, frame.margin_right = Pt(10), Pt(20)
    style = frame._txBody.find(f"{_A}lstStyle")
    level = OxmlElement("a:lvl1pPr")
    level.set("marL", str(Pt(30)))
    level.set("indent", str(Pt(-10)))
    style.append(level)
    _apply_safe_text_area(shape, {"textWidthPt": 100})
    assert frame.margin_right == shape.width - Pt(10 + 30 + 100)
    before = frame.margin_right, frame.margin_bottom
    _apply_safe_text_area(shape, {"textWidthPt": 1000, "textHeightPt": 1000})
    assert (frame.margin_right, frame.margin_bottom) == before


def test_export_applies_safe_area_before_replacing_copy(
    text_box: tuple[Any, Shape], tmp_path: Path
) -> None:
    deck, shape = text_box
    shape.text = "Sample"
    shape.text_frame.margin_left = Pt(10)
    shape.text_frame.paragraphs[0]._p.get_or_add_pPr().set("marL", str(Pt(24)))
    geometry = (shape.left, shape.top, shape.width, shape.height)
    source, output = tmp_path / "safe-source.pptx", tmp_path / "safe-output.pptx"
    deck.save(str(source))
    slot_id = str(shape.shape_id)
    text = "Copy stays inside the measured safe area"
    slot = {
        "slotId": slot_id,
        "maxChars": len(text),
        "textWidthPt": 80,
        "textHeightPt": 70,
        "box": {"cx": shape.width, "cy": shape.height},
    }
    export_compositions(
        source,
        output,
        [{"patternId": "safe", "texts": [{"slotId": slot_id, "text": text}]}],
        [{"patternId": "safe", "sourceIndex": 0, "slots": [slot]}],
    )
    exported = cast(Shape, Presentation(str(output)).slides[0].shapes[0])
    assert exported.text == text
    assert exported.text_frame.margin_right == shape.width - Pt(10 + 24 + 80)
    assert (exported.left, exported.top, exported.width, exported.height) == geometry


@pytest.mark.skipif(
    not all(shutil.which(tool) for tool in ("soffice", "pdftoppm")),
    reason="LibreOffice and Poppler are required for the rendered list check",
)
def test_render_keeps_list_order_soft_continuation_and_single_markers(
    text_box: tuple[Any, Shape], tmp_path: Path
) -> None:
    deck, shape = text_box
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Sample"
    paragraph.runs[0].font.name = "Arial"
    paragraph.runs[0].font.size = Pt(22)
    paragraph.space_after = Pt(12)
    _bullet(paragraph, "buChar", char="•")
    _replace_text(shape, "Alpha\vcontinued\n2. Beta\nGamma")
    output = tmp_path / "structure.pptx"
    deck.save(str(output))
    office = Path(shutil.which("soffice") or "")
    # The Windows console launcher waits for LibreOffice to release its user profile.
    if office.with_suffix(".com").is_file():
        office = office.with_suffix(".com")
    subprocess.run(
        [
            str(office),
            f"-env:UserInstallation={(tmp_path / 'office-profile').as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(output),
        ],
        check=True,
        capture_output=True,
        timeout=30,
    )
    subprocess.run(
        ["pdftoppm", "-png", "-r", "110", str(output.with_suffix(".pdf")), str(tmp_path / "slide")],
        check=True,
        capture_output=True,
        timeout=15,
    )
    previews = list(tmp_path.glob("slide-*.png"))
    assert len(previews) == 1
    with Image.open(previews[0]) as image:
        assert image.width > 0 and image.height > 0
        ink = image.convert("L").point([255 if value < 128 else 0 for value in range(256)])
        rows = [ink.crop((0, y, ink.width, y + 1)).getbbox() is not None for y in range(ink.height)]
        # Three items plus a soft continuation must render as four distinct lines.
        assert sum(present and (y == 0 or not rows[y - 1]) for y, present in enumerate(rows)) == 4
