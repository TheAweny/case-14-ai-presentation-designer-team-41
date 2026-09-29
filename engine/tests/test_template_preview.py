"""I keep thumbnail rendering bounded without discarding the template profile."""

import base64
import json
import subprocess
from pathlib import Path

import pytest
from engine.api.app import app
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Emu

from engine import render


def fake_renderer(monkeypatch, source: Path) -> list:
    calls = []

    def run(command, **kwargs):
        calls.append((command, kwargs))
        if command[0] == "soffice":
            directory = Path(command[command.index("--outdir") + 1])
            (directory / f"{source.stem}.pdf").write_bytes(b"pdf")
        else:
            Path(command[-1] + "-1.png").write_bytes(b"png")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(render.subprocess, "run", run)
    return calls


def test_preview_exports_and_rasterizes_only_first_slide(tmp_path, monkeypatch):
    source = tmp_path / "large.pptx"
    source.write_bytes(b"source")
    calls = fake_renderer(monkeypatch, source)
    images = render.render_to_png(source, tmp_path / "preview", first_slide_only=True)
    conversion = calls[0][0]
    pdf_filter = conversion[conversion.index("--convert-to") + 1]
    options = json.loads(pdf_filter.split(":", 2)[2])
    assert options["PageRange"]["value"] == "1"
    assert options["ExportHiddenSlides"]["value"] == "true"
    raster = calls[1][0]
    assert raster[raster.index("-f") + 1] == "1"
    assert raster[raster.index("-l") + 1] == "1"
    assert images[0].read_bytes() == b"png"
    assert source.read_bytes() == b"source"


def test_full_deck_render_remains_unrestricted(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    source.touch()
    calls = fake_renderer(monkeypatch, source)
    render.render_to_png(source, tmp_path / "all")
    assert calls[0][0][calls[0][0].index("--convert-to") + 1] == "pdf"
    assert "-l" not in calls[1][0]
    assert [kwargs["timeout"] for _, kwargs in calls] == [120, 120]


@pytest.mark.parametrize("first_slide_only", [False, True])
def test_render_stages_artifacts_on_destination_filesystem(tmp_path, monkeypatch, first_slide_only):
    source = tmp_path / "source.pptx"
    source.touch()
    destination = tmp_path / "mounted-runs" / "previews"
    calls = fake_renderer(monkeypatch, source)
    images = render.render_to_png(source, destination, first_slide_only=first_slide_only)
    command = calls[0][0]
    workspace = Path(command[command.index("--outdir") + 1])
    assert workspace.parent == destination.resolve()
    assert not workspace.exists()
    assert (destination / "source.pdf").read_bytes() == b"pdf"
    assert images[0].parent == destination


def test_preview_processes_share_one_deadline(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    source.touch()
    calls = fake_renderer(monkeypatch, source)
    times = iter([100.0, 102.0, 108.0])
    monkeypatch.setattr(render.time, "monotonic", lambda: next(times))
    render.render_to_png(source, tmp_path / "preview", timeout_seconds=20)
    assert [kwargs["timeout"] for _, kwargs in calls] == [18, 12]


def test_expired_preview_does_not_start_another_process(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    source.touch()
    calls = fake_renderer(monkeypatch, source)
    times = iter([100.0, 101.0, 121.0])
    monkeypatch.setattr(render.time, "monotonic", lambda: next(times))
    with pytest.raises(subprocess.TimeoutExpired):
        render.render_to_png(source, tmp_path / "preview", timeout_seconds=20)
    assert len(calls) == 1


@pytest.mark.parametrize("preview_fails", [False, True])
def test_upload_keeps_all_slides_and_bounds_optional_preview(tmp_path, monkeypatch, preview_fails):
    source = tmp_path / "template.pptx"
    deck = Presentation()
    for _ in range(3):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(str(source))
    calls = []

    def preview(path, directory, **kwargs):
        calls.append(kwargs)
        if preview_fails:
            raise subprocess.TimeoutExpired("soffice", 20)
        image = directory / "slide-1.png"
        image.write_bytes(b"thumbnail")
        return [image]

    monkeypatch.setattr(render, "render_to_png", preview)
    response = TestClient(app).post("/v1/template-profile", json={"template_path": str(source)})
    assert response.status_code == 200
    assert len(response.json()["templateIr"]["slides"]) == 3
    assert calls == [{"dpi": 96, "first_slide_only": True, "timeout_seconds": 20}]
    expected = (
        None
        if preview_fails
        else "data:image/png;base64," + base64.b64encode(b"thumbnail").decode("ascii")
    )
    assert response.json()["previewDataUrl"] == expected


def test_template_previews_draw_every_slide_at_the_requested_width(tmp_path, monkeypatch):
    source = tmp_path / "template.pptx"
    deck = Presentation()
    deck.slide_width = Emu(914400 * 20)
    deck.save(source)
    seen = {}

    def draw(path, output_dir, dpi, **kwargs):
        seen["dpi"] = dpi
        output_dir.mkdir(parents=True, exist_ok=True)
        images = [output_dir / f"slide-{n}.png" for n in (1, 2)]
        for image in images:
            image.write_bytes(b"png")
        return images

    monkeypatch.setattr(render, "render_to_png", draw)
    response = TestClient(app).post(
        "/v1/template-previews",
        json={
            "schemaVersion": 1,
            "templatePath": str(source),
            "outputDir": str(tmp_path / "previews"),
            "widthPx": 640,
        },
    )
    assert response.status_code == 200
    assert response.json()["previews"] == ["slide-1.png", "slide-2.png"]
    # 640 pixels over a 20-inch slide.
    assert seen["dpi"] == 32


def test_template_previews_refuse_a_missing_template(tmp_path):
    response = TestClient(app).post(
        "/v1/template-previews",
        json={
            "schemaVersion": 1,
            "templatePath": str(tmp_path / "missing.pptx"),
            "outputDir": str(tmp_path / "previews"),
            "widthPx": 640,
        },
    )
    assert response.status_code == 404
