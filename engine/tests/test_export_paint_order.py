"""Editor export must paint what the editor shows.

A master picture that fills the slide is already behind the slide's text. Reading the deck
flattens that picture into SlideIR, and writing the deck back used to draw it a second time on
top of the text. Generation export clones the source slide and must not do that either.

Nothing here names a customer template. The deck is built in the test.
"""

from __future__ import annotations

import shutil
import struct
import zlib
from pathlib import Path
from typing import cast

import pytest
from engine.checks import audit_deck
from engine.export import export_pdf, export_pptx
from engine.export.compositions import export_compositions
from engine.parser import design_system_of, import_deck
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.parts.slide import BaseSlidePart
from pptx.util import Emu, Inches, Pt


def _png(colour: tuple[int, int, int]) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        payload = tag + data
        return (
            struct.pack(">I", len(data))
            + payload
            + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    data = zlib.compress(b"\x00" + bytes(colour))
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", data) + chunk(b"IEND", b"")
    )


def _deck(path: Path, *, front_picture: bool) -> Path:
    """A synthetic deck: a red master picture, then a text box, then maybe a blue picture."""
    pixel = path.parent / "pixel.png"
    pixel.write_bytes(_png((220, 20, 20)))
    presentation = Presentation()
    master = presentation.slide_masters[0]
    # Slide masters can hold pictures; python-pptx only exposes add_picture on slide shape trees.
    image_part, relationship_id = cast(BaseSlidePart, master.part).get_or_add_image_part(str(pixel))
    assert presentation.slide_width is not None and presentation.slide_height is not None
    shape_id = master.shapes._next_shape_id
    master.shapes._spTree.add_pic(
        shape_id,
        f"Picture {shape_id - 1}",
        image_part.desc,
        relationship_id,
        0,
        0,
        int(presentation.slide_width),
        int(presentation.slide_height),
    )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "VISIBLE"
    run.font.size = Pt(32)
    run.font.bold = True
    if front_picture:
        blue = path.parent / "front.png"
        blue.write_bytes(_png((20, 20, 220)))
        slide.shapes.add_picture(
            str(blue),
            Emu(0),
            Emu(0),
            presentation.slide_width,
            presentation.slide_height,
        )
    presentation.save(str(path))
    return path


def _kinds(slide: dict) -> list[str]:
    return [
        element["kind"] for element in slide["elements"] if element["kind"] in {"image", "text"}
    ]


def _text_ink(image: Image.Image, box: dict[str, int], dpi: int) -> bool:
    """True when the raster shows dark text ink inside the text box."""
    scale = dpi / 914400
    left = max(0, int(box["x"] * scale))
    top = max(0, int(box["y"] * scale))
    right = min(image.width, int((box["x"] + box["cx"]) * scale))
    bottom = min(image.height, int((box["y"] + box["cy"]) * scale))
    for y in range(top, bottom):
        for x in range(left, right):
            pixel = image.getpixel((x, y))
            assert isinstance(pixel, tuple) and len(pixel) == 3
            red, green, blue = pixel
            if red < 80 and green < 80 and blue < 80:
                return True
    return False


def _slide_pictures(path: Path) -> int:
    presentation = Presentation(str(path))
    return sum(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[0].shapes
    )


def test_integrator_group_text_can_move_above_a_slide_picture(tmp_path: Path) -> None:
    source = tmp_path / "grouped.pptx"
    pixel = tmp_path / "blue.png"
    pixel.write_bytes(_png((20, 20, 220)))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Grouped text"
    slide.shapes.add_picture(str(pixel), Emu(0), Emu(0), deck.slide_width, deck.slide_height)
    deck.save(str(source))
    [imported] = import_deck(source)
    imported["elements"].sort(key=lambda element: 0 if element["kind"] == "image" else 1)
    destination = tmp_path / "reordered.pptx"
    export_pptx([imported], design_system_of(source), source, destination)
    shapes = list(Presentation(str(destination)).slides[0].shapes)
    group_at = next(
        index for index, shape in enumerate(shapes) if shape.shape_type == MSO_SHAPE_TYPE.GROUP
    )
    image_at = next(
        index for index, shape in enumerate(shapes) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    assert image_at < group_at


def test_editor_export_does_not_paint_the_master_picture_over_the_text(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=False)
    design = design_system_of(source)
    [imported] = import_deck(source)
    assert imported["elements"][0]["sourcePart"] == "master"
    assert imported["elements"][0]["kind"] == "image"
    assert "text" in _kinds(imported)
    assert _kinds(imported).index("image") < _kinds(imported).index("text")

    destination = tmp_path / "edited.pptx"
    export_pptx([imported], design, source, destination)
    # The master picture stays on the master. The slide itself does not gain a second copy.
    assert _slide_pictures(destination) == 0
    [again] = import_deck(destination)
    assert _kinds(again).index("image") < _kinds(again).index("text")
    assert again["elements"][0]["sourcePart"] == "master"
    report = audit_deck(
        {
            "schemaVersion": 1,
            "deckId": "synthetic",
            "slides": [again],
            "designSystem": design,
            "patternCatalog": None,
        },
        "default",
        {"picture-covers-text": {"enabled": True, "minCoverage": 0.85}},
    )
    assert all(finding["checkId"] != "picture-covers-text" for finding in report["findings"])


def test_generation_export_keeps_the_master_picture_behind_the_text(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=False)
    [imported] = import_deck(source)
    text = next(element for element in imported["elements"] if element["kind"] == "text")
    slot = str(text["sourceShapeId"])
    output = tmp_path / "generated.pptx"
    export_compositions(
        source,
        output,
        [
            {
                "patternId": "synthetic",
                "texts": [{"slotId": slot, "text": "VISIBLE"}],
                "speakerNotes": "",
            }
        ],
        [
            {
                "patternId": "synthetic",
                "sourceIndex": 0,
                "slots": [{"slotId": slot, "maxChars": 40}],
            }
        ],
    )
    assert _slide_pictures(output) == 0
    [again] = import_deck(output)
    assert _kinds(again).index("image") < _kinds(again).index("text")
    assert any(
        element["kind"] == "text" and "VISIBLE" in element["paragraphs"][0]["runs"][0]["text"]
        for element in again["elements"]
    )


def test_a_picture_placed_in_front_on_purpose_stays_in_front(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=True)
    design = design_system_of(source)
    [imported] = import_deck(source)
    kinds = _kinds(imported)
    assert kinds.index("text") < kinds.index("image") or kinds.count("image") == 2
    images = [element for element in imported["elements"] if element["kind"] == "image"]
    assert images[-1].get("sourcePart") != "master"
    destination = tmp_path / "front.pptx"
    export_pptx([imported], design, source, destination)
    # The intentional slide picture remains. The master picture is not copied onto the slide.
    assert _slide_pictures(destination) == 1
    [again] = import_deck(destination)
    slide_images = [
        element
        for element in again["elements"]
        if element["kind"] == "image" and element.get("sourcePart") != "master"
    ]
    text = next(element for element in again["elements"] if element["kind"] == "text")
    assert again["elements"].index(slide_images[0]) > again["elements"].index(text)


def test_a_new_picture_sent_behind_text_is_exported_behind_it(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=False)
    design = design_system_of(source)
    [imported] = import_deck(source)
    width, height = imported["slideCx"], imported["slideCy"]
    behind = {
        "kind": "image",
        "elementId": "behind",
        "role": "image",
        "box": {"x": 0, "y": 0, "cx": width, "cy": height},
        "sourceShapeId": None,
        "mediaPart": "/ppt/media/image1.png",
        "fit": "cover",
    }
    text_at = next(
        index for index, element in enumerate(imported["elements"]) if element["kind"] == "text"
    )
    imported["elements"].insert(text_at, behind)
    destination = tmp_path / "behind.pptx"
    export_pptx([imported], design, source, destination)
    [again] = import_deck(destination)
    slide_images = [
        element
        for element in again["elements"]
        if element["kind"] == "image" and element.get("sourcePart") != "master"
    ]
    text = next(element for element in again["elements"] if element["kind"] == "text")
    assert slide_images
    assert again["elements"].index(slide_images[0]) < again["elements"].index(text)


def test_deleted_slide_picture_is_not_written_back(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=True)
    design = design_system_of(source)
    [imported] = import_deck(source)
    imported["elements"] = [
        element
        for element in imported["elements"]
        if not (element["kind"] == "image" and element.get("sourcePart") != "master")
    ]
    destination = tmp_path / "deleted.pptx"
    export_pptx([imported], design, source, destination)
    assert _slide_pictures(destination) == 0


def test_picture_opacity_is_written_into_the_file(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx", front_picture=True)
    design = design_system_of(source)
    [imported] = import_deck(source)
    for element in imported["elements"]:
        if element["kind"] == "image" and element.get("sourcePart") != "master":
            element["opacity"] = 0.4
    destination = tmp_path / "faded.pptx"
    export_pptx([imported], design, source, destination)
    presentation = Presentation(str(destination))
    picture = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    alpha = picture._element.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}alphaModFix"
    )
    assert alpha is not None
    assert alpha.get("amt") == "40000"


def test_pdf_render_keeps_text_visible_over_the_master_picture(tmp_path: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        pytest.skip("LibreOffice is not installed; the structural export checks still ran.")
    source = _deck(tmp_path / "source.pptx", front_picture=False)
    design = design_system_of(source)
    [imported] = import_deck(source)
    pptx_path = tmp_path / "edited.pptx"
    export_pptx([imported], design, source, pptx_path)
    pdf_path = export_pdf(pptx_path, tmp_path / "edited.pdf")
    from engine.render import rasterize_pdf

    dpi = 72
    pngs = rasterize_pdf(pdf_path, tmp_path / "png", dpi=dpi)
    assert len(pngs) == 1
    text = next(element for element in imported["elements"] if element["kind"] == "text")
    with Image.open(pngs[0]) as image:
        rgb = image.convert("RGB")
        # The box is transparent and the glyphs do not cover its corner, so the master red
        # shows there even when the word is painted on top. Ink inside the box is the word.
        assert _text_ink(rgb, text["box"], dpi)
