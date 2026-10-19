"""Reproducible gaps between SlideIR and the shared exporter.

The visuals layer does not own that exporter. These probes only record what it drops.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import cast

from engine.export import export_pptx
from engine.visuals import draw_visual
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.shapes.autoshape import Shape
from pptx.util import Emu, Inches
from visual_samples import design_system

FIXTURE = Path(__file__).parent / "fixtures" / "visuals" / "exporter-gaps.json"


def _template(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Emu(12_192_000)
    deck.slide_height = Emu(6_858_000)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(1), Inches(0.3))
    shape.fill.background()
    shape.line.fill.background()
    deck.save(str(path))


def _slide(elements: list[dict]) -> dict:
    return {
        "schemaVersion": 1,
        "slideId": "gap",
        "variant": "balanced",
        "index": 0,
        "patternId": "synthetic",
        "slideCx": 12_192_000,
        "slideCy": 6_858_000,
        "elements": elements,
    }


def _png(path: Path, width: int, height: int) -> None:
    Image.new("RGB", (width, height), (20, 40, 80)).save(path)


def test_exporter_drops_geometry_and_a_picture_without_bytes(tmp_path: Path) -> None:
    template = tmp_path / "synthetic.pptx"
    _template(template)
    design = design_system()
    geometry = {
        "kind": "decor",
        "elementId": "freeform",
        "role": "decoration",
        "box": {"x": 400_000, "y": 400_000, "cx": 800_000, "cy": 400_000},
        "sourceShapeId": None,
        "geometry": {
            "width": 10,
            "height": 10,
            "paths": [{"d": "M0 0 L10 10", "filled": False, "stroked": True}],
        },
    }
    text = {
        "kind": "text",
        "elementId": "label",
        "role": "body",
        "box": {"x": 400_000, "y": 1_200_000, "cx": 3_000_000, "cy": 400_000},
        "sourceShapeId": None,
        "paragraphs": [
            {
                "level": 0,
                "lineSpacing": 1,
                "runs": [
                    {
                        "text": "Подпись рядом с фигурой",
                        "font": "DejaVu Sans",
                        "sizePt": 16,
                        "bold": False,
                        "italic": False,
                        "color": "",
                    }
                ],
            }
        ],
        "overflow": None,
    }
    picture = tmp_path / "wide.png"
    _png(picture, 400, 80)
    image = {
        "kind": "image",
        "elementId": "photo",
        "role": "image",
        "box": {"x": 4_000_000, "y": 400_000, "cx": 2_000_000, "cy": 2_000_000},
        "sourceShapeId": None,
        "mediaPart": "upload/missing-photo.png",
        "fit": "cover",
    }
    destination = tmp_path / "gaps.pptx"
    export_pptx(
        [_slide([geometry, text, image])],
        design,
        template,
        destination,
        draw_visual,
    )
    deck = Presentation(str(destination))
    shapes = list(deck.slides[0].shapes)
    texts = [cast(Shape, shape).text_frame.text for shape in shapes if shape.has_text_frame]
    pictures = [
        shape
        for shape in shapes
        if shape.shape_type is not None and "PICTURE" in str(shape.shape_type)
    ]
    observed = [
        {
            "id": "decor-geometry-omitted",
            "observed": not any(
                getattr(shape, "has_text_frame", False) and shape.width == 800_000
                for shape in shapes
            ),
            "summary": (
                "SlideIR decor with freeform geometry is drawn by the editor renderer and "
                "skipped by export, which logs and leaves the shape out."
            ),
        },
        {
            "id": "image-without-media-omitted",
            "observed": len(pictures) == 0
            and any("Подпись рядом с фигурой" in item for item in texts),
            "summary": (
                "An image element whose media part is not in the template or the export media "
                "map is absent from the PPTX. The editor still has the element, so a picture "
                "can fail to cover or to sit beside text only after export."
            ),
        },
    ]
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    if not FIXTURE.exists():
        FIXTURE.write_text(
            json.dumps({"probes": observed}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    recorded = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert recorded["probes"] == observed
