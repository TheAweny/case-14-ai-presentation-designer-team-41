"""Unknown synthetic templates exercise geometric room and structural evidence."""

from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from engine.api import app
from engine.export.compositions import export_compositions
from engine.layout.text_fit import check_composition_text, fitting_scale, measured_height
from engine.parser._composition_geometry import composition_evidence
from engine.parser.compositions import extract_compositions
from engine.parser.slides import read_deck
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.shapes.autoshape import Shape
from pptx.util import Inches, Pt


def _text(shapes: Any, text: str, x: float, y: float, w: float, h: float, size: float) -> Any:
    shape = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = "DejaVu Sans"
    return shape


def _body(path: Path, *, obstacle: bool = False, spacing: bool = False) -> tuple[Path, str]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide.shapes, "Результаты", 0.6, 0.2, 8, 0.7, 31)
    shape = _text(slide.shapes, "Итог", 0.6, 1.4, 7.1, 4.1, 19)
    if spacing:
        shape.text_frame.margin_left = Inches(0.8)
        shape.text_frame.margin_right = Inches(0.6)
        shape.text_frame.paragraphs[0].space_before = Pt(25)
        shape.text_frame.paragraphs[0].space_after = Pt(35)
    if obstacle:
        art = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(4), Inches(1.4), Inches(3), Inches(4)
        )
        art.fill.solid()
        art.fill.fore_color.rgb = RGBColor.from_string("5A326E")
    deck.save(str(path))
    return path, str(shape.shape_id)


def test_short_sample_does_not_cap_a_large_body_and_endpoint_preserves_capacity(
    tmp_path: Path,
) -> None:
    source, slot_id = _body(tmp_path / "unknown.pptx")
    pattern = extract_compositions(source)[0]
    slot = next(s for s in pattern["slots"] if s["slotId"] == slot_id)
    text = "Команда проверила процесс на пилоте. Результаты помогают выбрать следующий этап."
    assert slot["capacity"]["basis"] == "geometry"
    assert slot["capacity"]["safeChars"] == slot["maxChars"] > len(text)
    assert not check_composition_text(pattern, [{"slotId": slot_id, "text": text}])
    response = TestClient(app).post("/v1/generation-profile", json={"template_path": str(source)})
    assert response.status_code == 200, response.text
    actual = next(s for s in response.json()["compositions"][0]["slots"] if s["slotId"] == slot_id)
    assert actual["capacity"] == slot["capacity"]


def test_partial_decoration_limits_room_without_sample_cap(tmp_path: Path) -> None:
    plain, slot_id = _body(tmp_path / "plain.pptx")
    decorated, _ = _body(tmp_path / "decorated.pptx", obstacle=True)
    before = next(s for s in extract_compositions(plain)[0]["slots"] if s["slotId"] == slot_id)
    after = next(s for s in extract_compositions(decorated)[0]["slots"] if s["slotId"] == slot_id)
    assert after["capacity"]["basis"] == "obstacle-aware"
    assert len(after["sample"]) * 5 < after["maxChars"] < before["maxChars"]
    assert after["textWidthPt"] < before["textWidthPt"]


def test_margins_and_paragraph_spacing_reduce_capacity(tmp_path: Path) -> None:
    plain, slot_id = _body(tmp_path / "plain.pptx")
    spaced, _ = _body(tmp_path / "spaced.pptx", spacing=True)
    before = next(s for s in extract_compositions(plain)[0]["slots"] if s["slotId"] == slot_id)
    after = next(s for s in extract_compositions(spaced)[0]["slots"] if s["slotId"] == slot_id)
    assert after["maxChars"] < before["maxChars"]
    assert after["spaceBeforePt"] == 25 and after["spaceAfterPt"] == 35


def test_inherited_size_bullet_indent_and_group_obstacle_are_not_the_sample(
    tmp_path: Path,
) -> None:
    def build(path: Path, *, indent: bool, grouped: bool) -> tuple[Path, str, str]:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        host = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        shape = _text(host, "Кратко", 0.4, 0.8, 8, 4.2, 18)
        run = shape.text_frame.paragraphs[0].runs[0]._r
        props = run.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
        if props is not None and "sz" in props.attrib:
            del props.attrib["sz"]
        paragraph = shape.text_frame.paragraphs[0]
        ppr = paragraph._p.get_or_add_pPr()
        default = OxmlElement("a:defRPr")
        default.set("sz", "1800")
        latin = OxmlElement("a:latin")
        latin.set("typeface", "DejaVu Sans")
        default.append(latin)
        ppr.append(default)
        if indent:
            ppr.set("marL", str(int(Inches(1.6))))
            bullet = OxmlElement("a:buChar")
            bullet.set("char", "•")
            ppr.append(bullet)
        if grouped:
            art = host.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(5), Inches(0.8), Inches(3.2), Inches(4.2)
            )
            art.fill.solid()
            art.fill.fore_color.rgb = RGBColor.from_string("5A326E")
        deck.save(str(path))
        return path, str(shape.shape_id), shape.text

    plain, plain_id, sample = build(tmp_path / "plain.pptx", indent=False, grouped=False)
    indented, indented_id, _ = build(tmp_path / "indented.pptx", indent=True, grouped=False)
    blocked, blocked_id, _ = build(tmp_path / "grouped.pptx", indent=False, grouped=True)
    plain_slot = next(s for s in extract_compositions(plain)[0]["slots"] if s["slotId"] == plain_id)
    indented_slot = next(
        s for s in extract_compositions(indented)[0]["slots"] if s["slotId"] == indented_id
    )
    blocked_slot = next(
        s for s in extract_compositions(blocked)[0]["slots"] if s["slotId"] == blocked_id
    )
    assert plain_slot["fontSize"] == 18
    assert plain_slot["capacity"]["basis"] == "geometry"
    assert plain_slot["capacity"]["safeChars"] > len(sample) * 5
    assert indented_slot["capacity"]["safeChars"] < plain_slot["capacity"]["safeChars"]
    assert indented_slot["textWidthPt"] < plain_slot["textWidthPt"]
    assert blocked_slot["capacity"]["basis"] == "obstacle-aware"
    blocked_chars = blocked_slot["capacity"]["safeChars"]
    assert plain_slot["capacity"]["safeChars"] > blocked_chars > len(sample) * 5


def test_narrow_caption_stays_short(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide.shapes, "Подпись", 1, 2, 1.4, 0.45, 17)
    deck.save(str(tmp_path / "caption.pptx"))
    slot = extract_compositions(tmp_path / "caption.pptx")[0]["slots"][0]
    assert slot["maxLines"] == 1 and slot["maxChars"] < 20


def test_rotated_text_does_not_claim_measured_room(tmp_path: Path) -> None:
    source, slot_id = _body(tmp_path / "rotated.pptx")
    deck = Presentation(str(source))
    shape = cast(Shape, next(s for s in deck.slides[0].shapes if str(s.shape_id) == slot_id))
    shape.rotation = 17
    deck.save(str(source))
    slot = next(s for s in extract_compositions(source)[0]["slots"] if s["slotId"] == slot_id)
    assert slot["capacity"]["basis"] == "sample-fallback"


def test_capacity_measures_nominal_font_when_template_shrinks_sample(tmp_path: Path) -> None:
    source, slot_id = _body(tmp_path / "shrunk.pptx")
    deck = Presentation(str(source))
    shape = cast(Shape, next(s for s in deck.slides[0].shapes if str(s.shape_id) == slot_id))
    body = shape.text_frame._txBody.bodyPr
    for child in list(body):
        body.remove(child)
    shrink = OxmlElement("a:normAutofit")
    shrink.set("fontScale", "50000")
    shrink.set("lnSpcReduction", "10000")
    body.append(shrink)
    deck.save(str(source))
    slot = next(s for s in extract_compositions(source)[0]["slots"] if s["slotId"] == slot_id)
    assert slot["fontSize"] == 19
    assert slot["lineSpacing"] == pytest.approx(1.2)


def test_native_list_and_three_grouped_heading_body_pairs_have_different_evidence(
    tmp_path: Path,
) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    native = _text(slide.shapes, "Первый\nВторой\nТретий", 0.2, 1, 2.4, 4, 15)
    for index, paragraph in enumerate(native.text_frame.paragraphs):
        paragraph.level = index % 2
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        paragraph._p.get_or_add_pPr().append(bullet)
    for index in range(3):
        group = slide.shapes.add_group_shape()
        _text(group.shapes, str(index + 1), 3 + index * 2, 1, 0.4, 0.4, 15)
        _text(group.shapes, "Этап", 3 + index * 2, 1.5, 1.7, 0.5, 18)
        _text(group.shapes, "Описание шага", 3 + index * 2, 2.2, 1.7, 2, 14)
    source = tmp_path / "structure.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    evidence = composition_evidence(read_deck(source, [])[0], pattern)
    structures = evidence["structures"]
    lists = [s for s in structures if s["listEvidence"] == "native-list"]
    assert len(lists) == 1
    assert [p["level"] for p in lists[0]["paragraphs"]] == [0, 1, 0]
    assert len({tuple(s["groupIds"]) for s in structures if s["groupIds"]}) == 3
    assert sum(s["numberingCandidate"] for s in structures) == 3
    assert all(s["semanticRole"] == "unknown" for s in structures)


@pytest.mark.parametrize("soft", [True, False])
def test_paragraph_spacing_is_charged_for_hard_breaks_only(soft: bool) -> None:
    slot = {"fontSize": 20, "textWidthPt": 300, "lineSpacing": 1.2, "spaceAfterPt": 12}
    text = "Первый" + ("\v" if soft else "\n") + "Второй"
    assert measured_height(text, slot) == 44 + (12 if soft else 24)


@pytest.mark.parametrize("separator", ["\r\n", "\r", "\n"])
def test_fit_normalizes_paragraph_breaks_like_export(separator: str) -> None:
    slot = {"fontSize": 20, "textWidthPt": 300, "lineSpacing": 1.2, "spaceAfterPt": 12}
    assert measured_height(f"Первый{separator}Второй", slot) == 68


def test_baseline_spacing_does_not_add_padding_after_last_line() -> None:
    slot = {
        "slotId": "label",
        "font": "DejaVu Sans",
        "fontSize": 32,
        "lineSpacing": 1.32,
        "textWidthPt": 270,
        "textHeightPt": 39,
        "maxLines": 1,
        "maxChars": 30,
    }
    assert not check_composition_text({"slots": [slot]}, [{"slotId": "label", "text": "2026"}])
    assert check_composition_text({"slots": [slot]}, [{"slotId": "label", "text": "2026\n2027"}])


def test_first_line_still_requires_room_for_template_font() -> None:
    slot = {
        "slotId": "label",
        "fontSize": 32,
        "lineSpacing": 1.32,
        "textWidthPt": 270,
        "textHeightPt": 20,
        "maxLines": 1,
        "maxChars": 30,
    }
    assert check_composition_text({"slots": [slot]}, [{"slotId": "label", "text": "2026"}])


def test_long_unbreakable_cyrillic_word_is_rejected() -> None:
    slot = {
        "slotId": "1",
        "font": "DejaVu Sans",
        "fontSize": 19,
        "textWidthPt": 42,
        "textHeightPt": 250,
        "maxLines": 10,
        "maxChars": 200,
    }
    findings = check_composition_text(
        {"slots": [slot]}, [{"slotId": "1", "text": "Сверхпроизводительность"}]
    )
    assert any(f["code"] == "word-break" for f in findings)


def test_paragraph_spacing_can_overflow_without_exceeding_line_count() -> None:
    slot = {
        "slotId": "1",
        "font": "DejaVu Sans",
        "fontSize": 20,
        "textWidthPt": 300,
        "textHeightPt": 80,
        "maxLines": 3,
        "maxChars": 200,
        "spaceAfterPt": 15,
    }
    good = check_composition_text({"slots": [slot]}, [{"slotId": "1", "text": "Один\nДва"}])
    bad = check_composition_text({"slots": [slot]}, [{"slotId": "1", "text": "Один\nДва\nТри"}])
    assert not good
    assert any(f["code"] == "text-overflow" for f in bad)


@pytest.mark.parametrize("foreground", [False, True])
def test_image_constraints_preserve_z_order_and_template_protection(
    tmp_path: Path, foreground: bool
) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    bitmap = BytesIO()
    Image.new("RGB", (80, 40), "#A5B8C9").save(bitmap, format="PNG")
    picture = slide.shapes.add_picture(bitmap, Inches(0), Inches(0), Inches(10), Inches(7.5))
    picture.crop_left = 0.2
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1), Inches(8), Inches(3)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string("314E5C")
    _text(slide.shapes, "Заголовок", 0.6, 0.2, 8, 0.7, 32)
    shape = _text(slide.shapes, "Текст на фотографии", 1, 1.4, 6, 2, 20)
    if foreground:
        slide.shapes._spTree.append(picture._element)
    source = tmp_path / "overlap.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    evidence = composition_evidence(read_deck(source, [])[0], pattern)
    constraint = next(
        c for c in evidence["imageConstraints"] if c["textSlotId"] == str(shape.shape_id)
    )
    assert constraint["textAboveImage"] is not foreground
    assert constraint["state"] == ("unknown" if foreground else "known")
    assert constraint["readability"] == "unknown"
    assert constraint["regionInImage"]["width"] == pytest.approx(0.6)
    assert constraint["crop"]["l"] == 0.2
    frame = next(slot["frame"] for slot in pattern["imageSlots"] if slot["kind"] == "photo")
    assert frame["crop"]["l"] == 0.2
    assert frame["textRegions"]
    if not foreground:
        assert constraint["protection"][0]["fill"] == "314E5C"
        assert frame["state"] == "known"
        assert any(region["protection"] == "solid" for region in frame["textRegions"])
    else:
        assert frame["state"] == "unknown"
        assert any(region["protection"] == "unknown" for region in frame["textRegions"])


def _export_body(tmp_path: Path, text: str) -> tuple[float, set[float], float]:
    """The scale a body's copy needs, the run sizes it is exported at, and its template size."""
    path, body_id = _body(tmp_path / "in.pptx")
    pattern = extract_compositions(path)[0]
    slot = next(s for s in pattern["slots"] if s["slotId"] == body_id)
    texts = [
        {"slotId": s["slotId"], "text": text if s["slotId"] == body_id else "Итоги"}
        for s in pattern["slots"]
    ]
    scale = fitting_scale(text, slot)
    out = export_compositions(
        path,
        tmp_path / "out.pptx",
        [{"patternId": pattern["patternId"], "texts": texts}],
        [pattern],
        allow_overflow=True,
        font_scales={(0, body_id): scale},
    )
    shape = cast(
        Shape,
        next(s for s in Presentation(str(out)).slides[0].shapes if str(s.shape_id) == body_id),
    )
    sizes = {
        run.font.size.pt
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    }
    return scale, sizes, slot["fontSize"]


def test_copy_too_long_for_its_box_is_set_smaller_until_it_fits(tmp_path: Path) -> None:
    text = " ".join(["Выручка выросла на четверть за год"] * 24)
    scale, sizes, size = _export_body(tmp_path, text)
    path, body_id = _body(tmp_path / "again.pptx")
    slot = next(s for s in extract_compositions(path)[0]["slots"] if s["slotId"] == body_id)
    assert check_composition_text({"slots": [slot]}, [{"slotId": body_id, "text": text}])
    assert 0.6 <= scale < 1
    assert measured_height(text, {**slot, "fontSize": size * scale}) <= slot["textHeightPt"]
    assert sizes == {round(size * scale * 100) / 100}


def test_copy_that_fits_keeps_its_template_size(tmp_path: Path) -> None:
    scale, sizes, size = _export_body(tmp_path, "Выручка выросла на четверть за год")
    assert scale == 1
    assert sizes == {size}


def test_shrinking_stops_at_a_readable_body_size() -> None:
    from engine.layout.text_fit import least_scale

    body = {"fontSize": 16.0, "maxLines": 6, "textWidthPt": 200, "textHeightPt": 40}
    text = " ".join(["Выручка выросла на четверть за год"] * 30)
    assert fitting_scale(text, body) == least_scale(body)
    assert body["fontSize"] * fitting_scale(text, body) >= 12


def test_a_template_slot_already_small_is_never_shrunk() -> None:
    label = {"fontSize": 9.0, "maxLines": 1, "textWidthPt": 40, "textHeightPt": 12}
    assert fitting_scale("Очень длинная подпись к карточке", label) == 1.0
