"""I keep a figure or period label of the template from reading as a fact of the new deck."""

from pathlib import Path

import pytest
from engine.export.compositions import export_compositions, figure_sample
from engine.parser.compositions import extract_compositions
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


@pytest.mark.parametrize(
    "sample",
    ["05", "2,4%", "III кв.", "Q3", "70 млн", "2025", "Июл", "22 300", "01.", "$1.2B"],
)
def test_a_figure_or_period_label_is_a_figure_sample(sample: str) -> None:
    assert figure_sample(sample)


@pytest.mark.parametrize(
    "sample",
    ["", "Фокус на удержание", "Шаг 1", "Возможности и риски", "млн", "кв.", "Иван Петров"],
)
def test_words_and_bare_units_are_not(sample: str) -> None:
    assert not figure_sample(sample)


def _template(path: Path) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(1))
    title.text = "Заголовок слайда"
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(2), Inches(1), Inches(1))
    badge.text = "05"
    label = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(5), Inches(1))
    label.text = "Фокус на удержание"
    deck.save(str(path))
    return path


def _export(tmp_path: Path, texts: dict[str, str]) -> list[str]:
    source = _template(tmp_path / "source.pptx")
    pattern = extract_compositions(source)[0]
    by_sample = {slot["sample"]: slot["slotId"] for slot in pattern["slots"]}
    item = {
        "patternId": pattern["patternId"],
        "texts": [
            {"slotId": by_sample[sample], "text": texts.get(sample, sample)} for sample in by_sample
        ],
    }
    export_compositions(source, tmp_path / "out.pptx", [item], [pattern])
    shapes = Presentation(str(tmp_path / "out.pptx")).slides[0].shapes
    return [str(getattr(shape, "text", "")) for shape in shapes]


def test_a_figure_slot_left_empty_leaves_the_slide_with_its_shape(tmp_path: Path) -> None:
    kept = _export(tmp_path, {"05": ""})
    # The badge held the sample «05»; with no fact for it, the oval is gone, not drawn blank.
    assert kept == ["Заголовок слайда", "Фокус на удержание"]


def test_a_figure_slot_with_a_fact_keeps_its_shape(tmp_path: Path) -> None:
    assert "12" in _export(tmp_path, {"05": "12"})


def test_an_empty_slot_of_words_keeps_its_shape(tmp_path: Path) -> None:
    # Only a figure sample is taken away with its shape; a text frame of words stays as drawn.
    shapes = _export(tmp_path, {"Фокус на удержание": ""})
    assert len(shapes) == 3
