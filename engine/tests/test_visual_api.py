"""Public visual endpoint, then the existing export path, then the native chart sheet."""

from __future__ import annotations

import json
from pathlib import Path
from typing import cast

from engine.api.app import app
from engine.contracts import VisualBuildResponse
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.shapes.graphfrm import GraphicFrame
from pptx.util import Emu, Inches
from visual_samples import category_chart, kpi_block, request_for, time_chart


def _template(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Emu(12_192_000)
    deck.slide_height = Emu(6_858_000)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(2), Inches(0.4))
    shape.fill.background()
    shape.line.fill.background()
    deck.save(str(path))


def _slide(elements: list[dict], variant: str = "balanced") -> dict:
    return {
        "schemaVersion": 1,
        "slideId": "visual-1",
        "variant": variant,
        "index": 0,
        "patternId": "synthetic-body",
        "slideCx": 12_192_000,
        "slideCy": 6_858_000,
        "elements": elements,
    }


def test_capabilities_are_deterministic_and_name_the_deferred_planner() -> None:
    response = TestClient(app).get("/v1/visuals/capabilities")
    assert response.status_code == 200
    body = response.json()
    assert body["callsModels"] is False
    assert body["endpoint"] == "/v1/visuals"
    assert "semantic-planner-binding" in body["deferred"]


def test_endpoint_builds_a_kpi_the_contract_accepts() -> None:
    response = TestClient(app).post("/v1/visuals", json=request_for(kpi_block()))
    assert response.status_code == 200, response.text
    body = response.json()
    VisualBuildResponse.model_validate(body)
    assert body["issues"] == []
    assert any(
        "0 шт" in run["text"]
        for element in body["elements"]
        for paragraph in element["paragraphs"]
        for run in paragraph["runs"]
    )


def test_bad_literal_is_a_client_error() -> None:
    block = kpi_block()
    block["items"][0]["value"] = {"state": "present", "literal": "1", "value": 2}
    response = TestClient(app).post("/v1/visuals", json=request_for(block))
    assert response.status_code == 422


def test_chart_export_round_trip_keeps_native_values(tmp_path: Path) -> None:
    client = TestClient(app)
    built = client.post("/v1/visuals", json=request_for(category_chart()))
    assert built.status_code == 200, built.text
    elements = built.json()["elements"]
    template = tmp_path / "synthetic.pptx"
    _template(template)
    output = tmp_path / "out"
    output.mkdir()
    exported = client.post(
        "/v1/exports",
        json={
            "schemaVersion": 1,
            "templatePath": str(template),
            "variant": "balanced",
            "slides": [_slide(elements)],
            "outputDir": str(output),
            "formats": ["pptx"],
            "previews": False,
        },
    )
    assert exported.status_code == 200, exported.text
    pptx_path = Path(exported.json()["files"][0]["path"])
    deck = Presentation(str(pptx_path))
    charts = [cast(GraphicFrame, shape).chart for shape in deck.slides[0].shapes if shape.has_chart]
    tables = [cast(GraphicFrame, shape).table for shape in deck.slides[0].shapes if shape.has_table]
    assert charts and tables
    values = [list(series.values) for series in charts[0].series]
    assert values[0] == [0, -3.5, 12]
    assert values[1][1] is None
    headers = [cell.text for cell in tables[0].rows[0].cells]
    assert "План с очень длинным названием ряда, которое нельзя обрезать в легенде" in headers
    body = [[cell.text for cell in row.cells] for row in tables[0].rows]
    assert ["Север", "0", "4"] in body
    assert any("нет данных" in cell for row in body for cell in row)


def test_unequal_time_exports_as_xy_points(tmp_path: Path) -> None:
    client = TestClient(app)
    built = client.post("/v1/visuals", json=request_for(time_chart()))
    assert built.status_code == 200, built.text
    assert built.json()["chart"]["nativeAxis"] == "xy"
    template = tmp_path / "synthetic.pptx"
    _template(template)
    output = tmp_path / "out"
    output.mkdir()
    exported = client.post(
        "/v1/exports",
        json={
            "schemaVersion": 1,
            "templatePath": str(template),
            "variant": "balanced",
            "slides": [_slide(built.json()["elements"])],
            "outputDir": str(output),
            "formats": ["pptx"],
            "previews": False,
        },
    )
    assert exported.status_code == 200, exported.text
    deck = Presentation(str(Path(exported.json()["files"][0]["path"])))
    chart = next(
        cast(GraphicFrame, shape).chart for shape in deck.slides[0].shapes if shape.has_chart
    )
    # python-pptx Point has no x/y. The editable cache is c:xVal / c:yVal on the series.
    element = chart.series[0]._element
    count = min(element.xVal.ptCount_val, element.yVal.ptCount_val)
    points = [
        (float(element.xVal.pt_v(index)), float(element.yVal.pt_v(index)))
        for index in range(count)
        if element.xVal.pt_v(index) is not None and element.yVal.pt_v(index) is not None
    ]
    assert (2020, 10) in [(int(x), int(y)) for x, y in points]
    assert all(y != 0 or x == 2021 for x, y in points)
    # The missing 2025 point is absent, not drawn as zero.
    assert all(int(x) != 2025 for x, _y in points)
    recorded = output / "roundtrip.json"
    recorded.write_text(json.dumps({"points": points}, ensure_ascii=False), encoding="utf-8")
