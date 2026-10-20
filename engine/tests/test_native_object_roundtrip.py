"""Native KPI, chart and table objects survive export and reopen unchanged."""

from __future__ import annotations

from pathlib import Path
from typing import cast

from engine.export import export_pptx
from engine.visuals import build_visual, draw_visual, style_visuals
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.shapes.graphfrm import GraphicFrame
from pptx.util import Emu, Inches
from visual_samples import (
    category_chart,
    design_system,
    kpi_block,
    present,
    request_for,
    table_block,
    time_chart,
)


def _template(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Emu(12_192_000)
    deck.slide_height = Emu(6_858_000)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(2), Inches(0.4))
    shape.fill.background()
    shape.line.fill.background()
    deck.save(str(path))


def _slide(elements: list[dict]) -> dict:
    return {
        "schemaVersion": 1,
        "slideId": "native-roundtrip",
        "variant": "balanced",
        "index": 0,
        "patternId": "synthetic",
        "slideCx": 12_192_000,
        "slideCy": 6_858_000,
        "elements": elements,
    }


def _unequal_step_chart() -> dict:
    block = time_chart()
    block["categories"] = [
        {"label": "step-0", "position": 0},
        {"label": "step-1", "position": 1},
        {"label": "step-4", "position": 4},
    ]
    block["series"] = [
        {
            "name": "Запас",
            "unit": "сутки",
            "values": [present("10", 10), present("0", 0), present("-2", -2)],
        }
    ]
    block["requestedType"] = "line"
    return block


def _xy_points(chart: object) -> list[tuple[float, float]]:
    element = chart.series[0]._element
    count = min(element.xVal.ptCount_val, element.yVal.ptCount_val)
    return [
        (float(element.xVal.pt_v(index)), float(element.yVal.pt_v(index)))
        for index in range(count)
        if element.xVal.pt_v(index) is not None and element.yVal.pt_v(index) is not None
    ]


def test_native_kpi_chart_and_table_round_trip(tmp_path: Path) -> None:
    design = design_system()
    kpi_elements = build_visual(request_for(kpi_block(), element_id="kpi"))["elements"]
    chart_built = build_visual(request_for(category_chart(), element_id="chart"))
    chart_elements = [e for e in chart_built["elements"] if e["kind"] == "chart"]
    table_built = build_visual(request_for(table_block(), element_id="table"))
    table_elements = [
        element
        for element in style_visuals([_slide(table_built["elements"])], design)[0]["elements"]
        if element["kind"] == "table"
    ]
    xy_built = build_visual(request_for(_unequal_step_chart(), element_id="xy"))
    xy_elements = [e for e in xy_built["elements"] if e["kind"] == "chart"]
    assert xy_built["chart"]["nativeAxis"] == "xy"

    elements = kpi_elements + chart_elements + table_elements + xy_elements
    template = tmp_path / "template.pptx"
    _template(template)
    destination = tmp_path / "native.pptx"
    export_pptx([_slide(elements)], design, template, destination, draw_visual)

    deck = Presentation(str(destination))
    shapes = list(deck.slides[0].shapes)
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    charts = [
        cast(GraphicFrame, shape).chart for shape in shapes if getattr(shape, "has_chart", False)
    ]
    tables = [
        cast(GraphicFrame, shape).table for shape in shapes if getattr(shape, "has_table", False)
    ]
    kpi_text = [
        shape.text_frame.text
        for shape in shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]

    assert not pictures, "charts and KPI must stay native, not pictures"
    assert len(charts) == 2
    assert len(tables) == 1
    assert kpi_text, "KPI values are native text boxes"

    category_chart_native = charts[0]
    values = [list(series.values) for series in category_chart_native.series]
    assert values[0] == [0.0, -3.5, 12.0]
    assert values[1][0] == 4.0
    assert values[1][1] is None
    assert values[1][2] == 9.0
    assert category_chart_native.value_axis.has_title
    assert "млн ₽" in category_chart_native.value_axis.axis_title.text_frame.text
    assert list(category_chart_native.plots[0].categories) == ["Север", "Юг", "Восток"]

    xy_chart = charts[1]
    points = _xy_points(xy_chart)
    xs = [int(x) for x, _y in points]
    assert sorted(set(xs)) == [0, 1, 4], "unequal category positions must stay on the XY axis"
    assert (0, 10.0) in [(int(x), y) for x, y in points]
    assert (1, 0.0) in [(int(x), y) for x, y in points]
    assert (4, -2.0) in [(int(x), y) for x, y in points]

    table = tables[0]
    body = [[cell.text for cell in row.cells] for row in table.rows]
    assert ["Остаток", "0", "шт"] in body
    missing_row = next(row for row in body if row[0].startswith("Очень длинная ячейка"))
    assert missing_row[1] not in {"0", "0.0"}
    assert "нет данных" in missing_row[1]

    joined = "\n".join(kpi_text)
    assert "2024" in joined
    assert "1 кв. 2025" in joined
    assert "шт" in joined and "%" in joined
    assert "-12" in joined or "−12" in joined
    assert "нет данных" in joined
    assert "0" in joined
