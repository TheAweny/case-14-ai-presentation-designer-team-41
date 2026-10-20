"""A missing template face warns and still leaves that face in the exported file."""

import zipfile
from pathlib import Path
from typing import cast
from xml.etree import ElementTree

from engine.api.app import app
from engine.parser.compositions import extract_compositions
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.util import Inches, Pt

from engine import render

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _source(path: Path, family: str) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1.2))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Заголовок"
    run.font.size = Pt(28)
    run.font.name = family
    deck.save(str(path))
    return path


def _request(source: Path, output: Path) -> dict:
    pattern = extract_compositions(source)[0]
    return {
        "schemaVersion": 1,
        "templatePath": str(source),
        "outputPath": str(output),
        "slides": [
            {
                "patternId": pattern["patternId"],
                "texts": [{"slotId": pattern["slots"][0]["slotId"], "text": "Новый заголовок"}],
                "speakerNotes": "",
            }
        ],
    }


def _stub_preview(monkeypatch) -> None:
    def render_to_png(pptx_path, output_dir, dpi=110, **kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        image = output_dir / "slide-1.png"
        image.write_bytes(b"png")
        return [image]

    monkeypatch.setattr(render, "render_to_png", render_to_png)


def _stub_fonts(monkeypatch, present: set[str]) -> list[set[str]]:
    seen: list[set[str]] = []

    def missing(families: set[str]) -> list[str]:
        seen.append(set(families))
        return sorted(family for family in families if family and family not in present)

    monkeypatch.setattr(render, "missing_fonts", missing)
    return seen


def _typefaces(path: Path, text: str) -> set[str]:
    found: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for paragraph in root.iter(f"{_A}p"):
                written = "".join(node.text or "" for node in paragraph.iter(f"{_A}t"))
                if text not in written:
                    continue
                for node in paragraph.iter(f"{_A}latin"):
                    face = node.get("typeface")
                    if face:
                        found.add(face)
    return found


def test_absent_font_exports_and_keeps_the_template_typeface(tmp_path: Path, monkeypatch) -> None:
    family = "Template Face"
    source = _source(tmp_path / "in.pptx", family)
    assert extract_compositions(source)[0]["slots"][0]["font"] == family
    output = tmp_path / "out.pptx"
    _stub_preview(monkeypatch)
    seen = _stub_fonts(monkeypatch, set())
    response = TestClient(app).post("/v1/composition-decks", json=_request(source, output))
    assert response.status_code == 200, response.text
    body = response.json()
    assert seen == [{family}]
    assert body["missingFonts"] == [family]
    assert (
        "Fonts not installed on the render server: Template Face. "
        "The .pptx keeps the template fonts; previews use substitutes."
    ) in body["warnings"]
    assert _typefaces(output, "Новый заголовок") == {family}


_ABSENT_WHEN_UNSET = ("bullet", "marginEmu", "indentEmu", "spaceBeforePt", "spaceAfterPt", "align")


def _paragraphs(node: object) -> list[dict]:
    found: list[dict] = []
    if isinstance(node, dict):
        if "runs" in node and "level" in node:
            found.append(node)
        for value in node.values():
            found.extend(_paragraphs(value))
    elif isinstance(node, list):
        for item in node:
            found.extend(_paragraphs(item))
    return found


def test_export_omits_unset_paragraph_fields(tmp_path: Path, monkeypatch) -> None:
    source = _source(tmp_path / "in.pptx", "Template Face")
    output = tmp_path / "out.pptx"
    _stub_preview(monkeypatch)
    _stub_fonts(monkeypatch, {"Template Face"})
    response = TestClient(app).post("/v1/composition-decks", json=_request(source, output))
    assert response.status_code == 200, response.text
    paragraphs = _paragraphs(response.json()["templateIr"])
    assert paragraphs
    for paragraph in paragraphs:
        assert "lineSpacing" in paragraph
        for key in _ABSENT_WHEN_UNSET:
            assert key not in paragraph or paragraph[key] is not None


def test_installed_font_exports_without_a_font_warning(tmp_path: Path, monkeypatch) -> None:
    family = "Template Face"
    source = _source(tmp_path / "in.pptx", family)
    output = tmp_path / "out.pptx"
    _stub_preview(monkeypatch)
    _stub_fonts(monkeypatch, {family})
    response = TestClient(app).post("/v1/composition-decks", json=_request(source, output))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["missingFonts"] == []
    assert all("Fonts not installed" not in warning for warning in body["warnings"])
    assert _typefaces(output, "Новый заголовок") == {family}


def test_peer_cards_keep_one_type_scale_and_centre_their_complete_copy(tmp_path, monkeypatch):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.8))
    title.text_frame.text = "Решения для запуска"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    cards = []
    for x in (0.5, 3.5, 6.5):
        card = slide.shapes.add_textbox(Inches(x), Inches(2), Inches(2.5), Inches(2.5))
        card.text_frame.text = "Заголовок\nТекст"
        for paragraph in card.text_frame.paragraphs:
            paragraph.runs[0].font.size = Pt(20)
        cards.append(str(card.shape_id))
    source, output = tmp_path / "cards.pptx", tmp_path / "out.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    copy = [
        "Бюджет\n900 000 ₽",
        "Владелец\nНазначить владельца пилота.",
        "Участие\nПодтвердить участие 30 операторов в трёхмесячном пилоте.",
    ]
    bindings = dict(zip(cards, copy, strict=True))
    _stub_preview(monkeypatch)
    _stub_fonts(monkeypatch, set())
    response = TestClient(app).post(
        "/v1/composition-decks",
        json={
            "schemaVersion": 1,
            "templatePath": str(source),
            "outputPath": str(output),
            "allowOverflow": True,
            "slides": [
                {
                    "patternId": pattern["patternId"],
                    "speakerNotes": "",
                    "texts": [
                        {
                            "slotId": s["slotId"],
                            "text": bindings.get(s["slotId"], "Решения для запуска"),
                        }
                        for s in pattern["slots"]
                    ],
                }
            ],
        },
    )
    assert response.status_code == 200, response.text
    exported = Presentation(str(output)).slides[0]
    frames = [cast(Shape, s).text_frame for s in exported.shapes if str(s.shape_id) in cards]
    assert len(frames) == 3
    assert {frame._txBody.bodyPr.get("anchor") for frame in frames} == {"ctr"}
    assert len({frame.paragraphs[0].runs[0].font.size for frame in frames}) == 1


def test_resizing_preserves_heading_body_hierarchy(tmp_path):
    from engine.export.compositions import export_compositions

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(4))
    shape.text_frame.text = "Заголовок\nОписание"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    shape.text_frame.paragraphs[1].runs[0].font.size = Pt(16)
    source, output = tmp_path / "source.pptx", tmp_path / "out.pptx"
    deck.save(str(source))
    pattern = extract_compositions(source)[0]
    slot = pattern["slots"][0]
    export_compositions(
        source,
        output,
        [
            {
                "patternId": pattern["patternId"],
                "texts": [{"slotId": slot["slotId"], "text": "Решение\nПроверенное пояснение"}],
            }
        ],
        [pattern],
        allow_overflow=True,
        font_scales={(0, slot["slotId"]): 1.25},
    )
    frame = cast(Shape, Presentation(str(output)).slides[0].shapes[0]).text_frame
    heading_size = frame.paragraphs[0].runs[0].font.size
    body_size = frame.paragraphs[1].runs[0].font.size
    assert heading_size is not None
    assert body_size is not None
    assert heading_size.pt == 30
    assert body_size.pt == 20
