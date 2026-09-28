"""I pin template inheritance and measurable text-fit gates before generation changes."""

from pathlib import Path

from engine.layout.text_fit import check_composition_text
from engine.parser.compositions import extract_compositions
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


def make_source(path: Path, inherited: bool = False) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0] if inherited else deck.slide_layouts[6])
    if inherited:
        shape = slide.shapes.title
        assert shape is not None
        shape.text = "Inherited title"
        layout_shape = slide.slide_layout.placeholders[0]
        style = layout_shape._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}txBody/{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle"
        )
        assert style is not None
        level = OxmlElement("a:lvl1pPr")
        props = OxmlElement("a:defRPr")
        props.set("sz", "4400")
        font = OxmlElement("a:latin")
        font.set("typeface", "Arial")
        props.append(font)
        level.append(props)
        style.append(level)
    else:
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "A short title"
        run.font.size = Pt(24)
        run.font.name = "Arial"
    deck.save(str(path))
    return path


def test_inherited_placeholder_font_is_not_assumed_18pt(tmp_path: Path) -> None:
    patterns = extract_compositions(make_source(tmp_path / "input.pptx", inherited=True))
    title = next(slot for slot in patterns[0]["slots"] if slot["role"] == "title")
    assert title["fontSize"] == 44
    assert title["font"] == "Arial"


def test_slot_geometry_and_text_insets_are_retained(tmp_path: Path) -> None:
    slot = extract_compositions(make_source(tmp_path / "input.pptx"))[0]["slots"][0]
    assert slot["box"]["x"] == Inches(1)
    assert slot["textWidthPt"] < 288
    assert slot["textHeightPt"] < 72


def test_short_content_passes_measured_fit(tmp_path: Path) -> None:
    pattern = extract_compositions(make_source(tmp_path / "input.pptx"))[0]
    assert (
        check_composition_text(
            pattern, [{"slotId": pattern["slots"][0]["slotId"], "text": "Короткий заголовок"}]
        )
        == []
    )


def test_long_content_produces_overflow_finding(tmp_path: Path) -> None:
    pattern = extract_compositions(make_source(tmp_path / "input.pptx"))[0]
    findings = check_composition_text(
        pattern, [{"slotId": pattern["slots"][0]["slotId"], "text": "Длинный русский текст " * 35}]
    )
    assert any(f["code"] == "text-overflow" for f in findings)


def test_repeated_footer_sample_is_not_treated_as_branding(tmp_path: Path) -> None:
    deck = Presentation()
    for _ in range(4):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(4), Inches(0.4))
        shape.text = "Имя Фамилия"
    path = tmp_path / "repeated.pptx"
    deck.save(str(path))
    patterns = extract_compositions(path)
    assert all(any(s["sample"] == "Имя Фамилия" for s in p["slots"]) for p in patterns)


def test_resource_sheet_is_not_offered_as_content(tmp_path: Path) -> None:
    deck = Presentation(str(make_source(tmp_path / "input.pptx")))
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i in range(4):
        shape = slide.shapes.add_textbox(Inches(i), Inches(1), Inches(1), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Download asset"
        run.hyperlink.address = f"https://example.org/{i}"
    path = tmp_path / "links.pptx"
    deck.save(str(path))
    assert len(extract_compositions(path)) == 1


def test_api_reports_overflow_before_export(tmp_path: Path) -> None:
    from engine.api.app import app
    from fastapi.testclient import TestClient

    source = make_source(tmp_path / "input.pptx")
    pattern = extract_compositions(source)[0]
    request = {
        "schemaVersion": 1,
        "templatePath": str(source),
        "outputPath": str(tmp_path / "output.pptx"),
        "slides": [
            {
                "patternId": pattern["patternId"],
                "texts": [{"slotId": pattern["slots"][0]["slotId"], "text": "Длинное слово " * 90}],
                "speakerNotes": "",
            }
        ],
    }
    client = TestClient(app)
    response = client.post("/v1/composition-check", json=request)
    assert response.status_code == 200
    assert response.json()["findings"][0]["slideIndex"] == 0
    assert client.post("/v1/composition-decks", json=request).status_code == 422
    assert not (tmp_path / "output.pptx").exists()


def test_title_capacity_stops_before_overlapping_subtitle(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for y, height, size, text in [(1, 2, 30, "Title"), (1.6, 1, 14, "Subtitle")]:
        shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(7), Inches(height))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
    path = tmp_path / "overlap.pptx"
    deck.save(str(path))
    title = extract_compositions(path)[0]["slots"][0]
    assert title["maxLines"] == 1


def test_missing_template_font_is_reported(monkeypatch) -> None:
    import subprocess

    import pytest
    from engine.render import missing_fonts, require_fonts

    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess([], 0, "Fallback Serif"),
    )
    assert missing_fonts({"OriginalFont"}) == ["OriginalFont"]
    with pytest.raises(ValueError, match="OriginalFont"):
        require_fonts({"OriginalFont"})


def test_available_template_font_is_accepted(monkeypatch) -> None:
    import subprocess

    from engine.render import missing_fonts, require_fonts

    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess([], 0, "OriginalFont"),
    )
    assert missing_fonts({"OriginalFont"}) == []
    require_fonts({"OriginalFont"})
