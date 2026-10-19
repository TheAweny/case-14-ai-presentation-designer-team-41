"""Independent counterexamples for visual identity and missing XY measurements."""

from pathlib import Path
from zipfile import ZipFile

import lxml.etree as etree
from engine.visuals import draw_visual
from engine.visuals.signature import structural_signature
from pptx import Presentation
from visual_samples import design_system


def test_integrator_nonoverlapping_element_order_is_not_a_new_design() -> None:
    elements = [
        {"kind": "text", "role": "title", "box": {"x": 0, "y": 0, "cx": 100, "cy": 20}},
        {"kind": "text", "role": "body", "box": {"x": 0, "y": 30, "cx": 100, "cy": 60}},
    ]
    assert structural_signature(elements) == structural_signature(list(reversed(elements)))


def test_integrator_xy_missing_point_survives_native_chart(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    element = {
        "kind": "chart",
        "elementId": "gap",
        "role": "chart",
        "sourceShapeId": None,
        "box": {"x": 0, "y": 0, "cx": 8_000_000, "cy": 5_000_000},
        "chartType": "line",
        "categories": ["2020", "2021", "2024"],
        "positions": [2020, 2021, 2024],
        "series": [{"name": "Revenue", "values": [10, None, 20]}],
        "xTitle": "Year",
        "yTitle": "RUB",
    }
    draw_visual(slide.shapes, element, design_system())
    output = tmp_path / "missing.pptx"
    presentation.save(str(output))
    with ZipFile(output) as archive:
        chart = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    # Keep the x position and an empty y entry. Omitting the point connects its neighbours.
    assert chart.xpath(".//c:xVal/c:numRef/c:numCache/c:ptCount/@val", namespaces=ns) == ["3"]
    assert chart.xpath(
        ".//c:yVal/c:numRef/c:numCache/c:pt[@idx='1']/c:v/text()", namespaces=ns
    ) in ([], [""])
