"""Uploaded images pass through the real export endpoint without arbitrary file access."""

from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from engine.api import app
from engine.export import export_pptx
from engine.parser import design_system_of
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Picture
from pptx.util import Inches, Pt


def _png(size: tuple[int, int] = (64, 64), *, format: str = "PNG") -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, "black").save(buffer, format=format)
    return buffer.getvalue()


@pytest.fixture
def payload(tmp_path: Path) -> dict[str, Any]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(7), Inches(1))
    title.text = "Source"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    slide.shapes.add_picture(BytesIO(_png()), Inches(1), Inches(2), Inches(1), Inches(1))
    template = tmp_path / "template.pptx"
    deck.save(str(template))
    media = tmp_path / "exports" / "media"
    media.mkdir(parents=True)
    image = media / "hash.png"
    image.write_bytes(_png())
    return {
        "schemaVersion": 1,
        "templatePath": str(template),
        "outputDir": str(media.parent),
        "variant": "balanced",
        "formats": ["pptx"],
        "previews": False,
        "media": [{"mediaPart": "upload/qr.png", "path": str(image)}],
        "slides": [
            {
                "schemaVersion": 1,
                "slideId": "s1",
                "index": 0,
                "variant": "balanced",
                "patternId": "/ppt/slideLayouts/slideLayout7.xml",
                "slideCx": 9144000,
                "slideCy": 6858000,
                "elements": [
                    {
                        "kind": "image",
                        "elementId": "qr",
                        "role": "image",
                        "sourceShapeId": None,
                        "mediaPart": "upload/qr.png",
                        "fit": "contain",
                        "box": {"x": 914400, "y": 914400, "cx": 914400, "cy": 914400},
                    }
                ],
            }
        ],
    }


@pytest.mark.parametrize("format", ["PNG", "JPEG"])
def test_export_endpoint_embeds_exact_uploaded_bytes(payload: dict, format: str) -> None:
    expected = _png(format=format)
    if format == "JPEG":
        payload["media"][0]["mediaPart"] = "upload/qr.jpg"
        payload["slides"][0]["elements"][0]["mediaPart"] = "upload/qr.jpg"
    Path(payload["media"][0]["path"]).write_bytes(expected)
    response = TestClient(app).post("/v1/exports", json=payload)
    assert response.status_code == 200, response.text
    result = Presentation(response.json()["files"][0]["path"])
    pictures = [s for s in result.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert cast(Picture, pictures[0]).image.blob == expected


@pytest.mark.parametrize(
    "case",
    [
        "outside",
        "missing",
        "unreferenced",
        "duplicate",
        "missing-file",
        "traversal",
        "oversize",
        "total",
        "count",
        "signature",
        "dimensions",
        "corrupt",
        "wrong-format",
    ],
)
def test_export_endpoint_rejects_invalid_media(payload: dict, tmp_path: Path, case: str) -> None:
    entry = payload["media"][0]
    file = Path(entry["path"])
    if case == "outside":
        outside = tmp_path / "outside.png"
        outside.write_bytes(_png())
        entry["path"] = str(outside)
    elif case == "missing":
        del payload["media"]
    elif case == "unreferenced":
        entry["mediaPart"] = "upload/unreferenced.png"
    elif case == "duplicate":
        payload["media"].append(dict(entry))
    elif case == "missing-file":
        file.unlink()
    elif case == "traversal":
        entry["mediaPart"] = "upload/../qr.png"
        payload["slides"][0]["elements"][0]["mediaPart"] = entry["mediaPart"]
    elif case == "oversize":
        file.write_bytes(_png() + b"0" * (1024 * 1024))
    elif case in ("total", "count"):
        payload["media"] = []
        payload["slides"][0]["elements"] = []
        for i in range(4 if case == "total" else 33):
            part = f"upload/qr{i}.png"
            payload["media"].append({"mediaPart": part, "path": str(file)})
            payload["slides"][0]["elements"].append(
                {
                    "kind": "image",
                    "elementId": f"qr{i}",
                    "role": "image",
                    "sourceShapeId": None,
                    "mediaPart": part,
                    "fit": "contain",
                    "box": {"x": 0, "y": 0, "cx": 100, "cy": 100},
                }
            )
        if case == "total":
            data = _png()
            file.write_bytes(data + b"0" * (800_000 - len(data)))
    elif case == "signature":
        file.write_bytes(b"not an image")
    elif case == "dimensions":
        file.write_bytes(_png((8193, 1)))
    elif case == "corrupt":
        file.write_bytes(_png()[:40])
    elif case == "wrong-format":
        file.write_bytes(_png(format="JPEG"))
    response = TestClient(app).post("/v1/exports", json=payload)
    assert response.status_code == 422, response.text
    assert not (Path(payload["outputDir"]) / "balanced.pptx").exists()


@pytest.mark.parametrize("root_link", [False, True])
def test_export_endpoint_rejects_symlink_escape(
    payload: dict, tmp_path: Path, root_link: bool
) -> None:
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "hash.png").write_bytes(_png())
    media = Path(payload["outputDir"]) / "media"
    link = media / "linked"
    if root_link:
        (media / "hash.png").unlink()
        media.rmdir()
        link = media
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("Creating symlinks is not supported for this Windows account.")
    payload["media"][0]["path"] = str(link / "hash.png")
    response = TestClient(app).post("/v1/exports", json=payload)
    assert response.status_code == 422, response.text


def test_upload_lookup_does_not_collide_with_template_basename(
    payload: dict, tmp_path: Path
) -> None:
    source = Path(payload["templatePath"])
    slide = payload["slides"][0]
    slide["elements"][0]["mediaPart"] = "image1.png"
    upload = _png((31, 31))
    export_pptx(
        [slide],
        design_system_of(source),
        source,
        tmp_path / "out.pptx",
        external_media={"upload/image1.png": upload},
    )
    picture = next(
        s
        for s in Presentation(str(tmp_path / "out.pptx")).slides[0].shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    assert cast(Picture, picture).image.blob == _png()
    slide["elements"][0]["mediaPart"] = "upload/image1.png"
    export_pptx(
        [slide],
        design_system_of(source),
        source,
        tmp_path / "uploaded.pptx",
        external_media={"upload/image1.png": upload},
    )
    picture = next(
        s
        for s in Presentation(str(tmp_path / "uploaded.pptx")).slides[0].shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    assert cast(Picture, picture).image.blob == upload


@pytest.mark.parametrize(
    "key", ["ppt/media/image1.png", "image1.png", "upload/../x.png", "upload/x\\y.png"]
)
def test_export_rejects_invalid_upload_keys(payload: dict, tmp_path: Path, key: str) -> None:
    source = Path(payload["templatePath"])
    with pytest.raises(ValueError, match="upload part"):
        export_pptx(
            payload["slides"],
            design_system_of(source),
            source,
            tmp_path / "bad.pptx",
            external_media={key: _png()},
        )
