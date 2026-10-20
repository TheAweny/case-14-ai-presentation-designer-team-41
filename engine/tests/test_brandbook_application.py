"""I verify declared rules in native files without changing slide geometry."""

import zipfile
from typing import Any

import lxml.etree as etree
import pptx
import pytest
from engine.export.brandbook import apply_brandbook
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def read_presentation(path=None) -> Any:
    return pptx.Presentation(path)


def token(kind: str, value: Any, role: str, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "kind": kind,
        "value": value,
        "role": role,
        "scope": "global",
        "authority": "explicit",
        "unit": None,
    }
    return {**base, **overrides}


def binding(*tokens, priority="brandbook"):
    return {"priority": priority, "profile": {"tokens": list(tokens)}}


@pytest.fixture
def deck(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = read_presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "A heading"
    body = slide.placeholders[1]
    body.text = "Body text"
    for shape in (slide.shapes.title, body):
        shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    extra = slide.shapes.add_shape(1, Inches(1), Inches(4), Inches(1), Inches(1))
    extra.fill.solid()
    extra.fill.fore_color.rgb = RGBColor.from_string("ABCDEF")
    presentation.save(source)
    return source, tmp_path / "branded.pptx"


def test_declared_title_and_body_fonts_override_native_runs(deck):
    source, target = deck
    apply_brandbook(
        source,
        target,
        binding(
            token("font-family", "Brand Title", "title"), token("font-family", "Brand Body", "body")
        ),
    )
    slide = read_presentation(target).slides[0]
    assert slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name == "Brand Title"
    assert slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.name == "Brand Body"


def test_font_px_converts_to_points(deck):
    source, target = deck
    apply_brandbook(source, target, binding(token("font-size", 48, "title", unit="px")))
    run = read_presentation(target).slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 36


def test_template_priority_is_byte_identical(deck):
    source, target = deck
    apply_brandbook(
        source, target, binding(token("font-family", "Brand", "body"), priority="template")
    )
    assert source.read_bytes() == target.read_bytes()


@pytest.mark.parametrize("authority", ["observed", "inferred"])
def test_non_declared_rule_does_not_override(deck, authority):
    source, target = deck
    apply_brandbook(
        source, target, binding(token("font-family", "Brand", "body", authority=authority))
    )
    assert source.read_bytes() == target.read_bytes()


def test_unknown_roles_and_scopes_are_reported_without_guessing(deck):
    source, target = deck
    warnings = apply_brandbook(
        source,
        target,
        binding(
            token("color", "#112233", "mystery"),
            token("font-family", "Brand", "body", scope="frame:99"),
        ),
    )
    assert len(warnings) == 2
    assert source.read_bytes() == target.read_bytes()


def test_theme_accent_changes_but_unknown_direct_paint_stays(deck):
    source, target = deck
    apply_brandbook(source, target, binding(token("color", "#112233", "accent1")))
    with zipfile.ZipFile(target) as archive:
        theme = etree.fromstring(archive.read("ppt/theme/theme1.xml"))
    assert theme.xpath("//*[local-name()='accent1']/*/@val") == ["112233"]
    assert str(read_presentation(target).slides[0].shapes[-1].fill.fore_color.rgb) == "ABCDEF"


def test_text_colour_and_geometry_preserved(deck):
    source, target = deck
    original = read_presentation(source).slides[0].shapes.title
    box = (original.left, original.top, original.width, original.height)
    apply_brandbook(source, target, binding(token("color", "#123456", "text")))
    shape = read_presentation(target).slides[0].shapes.title
    assert (shape.left, shape.top, shape.width, shape.height) == box
    assert str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == "123456"


def test_native_tables_keep_cells_and_apply_body_font(deck):
    source, target = deck
    presentation = read_presentation(source)
    table = presentation.slides[0].shapes.add_table(2, 2, 0, 0, Inches(3), Inches(2)).table
    table.cell(0, 0).text = "Native cell"
    presentation.save(source)
    apply_brandbook(source, target, binding(token("font-family", "Brand Body", "body")))
    table = read_presentation(target).slides[0].shapes[-1].table
    assert table.cell(0, 0).text == "Native cell"
    assert table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name == "Brand Body"


def test_declared_background_sets_editable_slide_background(deck):
    source, target = deck
    apply_brandbook(source, target, binding(token("color", "#102030", "background")))
    assert str(read_presentation(target).slides[0].background.fill.fore_color.rgb) == "102030"


def test_slide_scope_changes_only_named_slide(deck):
    source, target = deck
    original = read_presentation(source)
    slide = original.slides.add_slide(original.slide_layouts[1])
    slide.shapes.title.text = "Second heading"
    original.save(source)
    apply_brandbook(
        source, target, binding(token("font-family", "Scoped Title", "title", scope="slide:1"))
    )
    read = read_presentation(target)
    assert read.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.name == "Scoped Title"
    assert read.slides[1].shapes.title.text_frame.paragraphs[0].runs[0].font.name is None


def test_effective_template_reextracts_font_sizes_for_fit(deck, monkeypatch):
    from engine.api.app import brandbook_template
    from engine.contracts._generated import BrandbookTemplateRequest

    from engine import render

    source, target = deck
    token_value = token("font-size", 40, "title", unit="pt")
    token_value.update(
        id="title-size",
        name="Title",
        confidence=1,
        evidence=[{"sourceId": "source", "locator": "$.title", "excerpt": "40pt"}],
    )
    profile = {
        "schemaVersion": 1,
        "parserVersion": "test",
        "sourceSha256": "a" * 64,
        "status": "ready",
        "tokens": [token_value],
        "assets": [],
        "sections": [],
        "warnings": [],
        "sources": [
            {
                "id": "source",
                "path": "brand.json",
                "format": "json",
                "sha256": "a" * 64,
                "sizeBytes": 20,
                "status": "parsed",
                "warnings": [],
            }
        ],
        "coverage": {"filesTotal": 1, "filesParsed": 1, "filesUnsupported": 0, "needsOcr": False},
    }
    request = BrandbookTemplateRequest.model_validate(
        {
            "schemaVersion": 1,
            "templatePath": str(source),
            "outputPath": str(target),
            "brandbook": {"schemaVersion": 1, "priority": "brandbook", "profile": profile},
        }
    )
    monkeypatch.setattr(render, "missing_fonts", lambda fonts: [])
    result = brandbook_template(request)
    sizes = [
        slot["fontSize"]
        for pattern in result["profile"]["compositions"]
        for slot in pattern["slots"]
        if slot["sample"] == "A heading"
    ]
    assert sizes and set(sizes) == {40}
    assert result["templatePath"] == str(target)
    assert result["profile"]["designSystem"] == result["designSystem"]
