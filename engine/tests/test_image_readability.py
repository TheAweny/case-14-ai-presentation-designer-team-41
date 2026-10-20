"""Synthetic saved files exercise media pixels rather than a reconstructed SlideIR background."""

from __future__ import annotations

import json
import shutil
import subprocess
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from engine.checks import (
    DETERMINISTIC_CHECKS,
    PPTX_CHECKS,
    audit_deck,
    audit_exported_deck,
    audit_image_readability,
    check_by_id,
    image_readability,
)
from engine.contracts import AuditFinding, AuditReport
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.slide import BaseSlidePart
from pptx.util import Inches, Pt


def _image(colour: str = "white") -> Image.Image:
    return Image.new("RGB", (320, 160), colour)


def _picture(shapes: Any, image: Image.Image) -> Any:
    blob = BytesIO()
    image.save(blob, "PNG")
    blob.seek(0)
    return shapes.add_picture(blob, Inches(0), Inches(0), Inches(8), Inches(4))


def _text(shapes: Any, colour: str = "FFFFFF", *, x: float = 1) -> Any:
    shape = shapes.add_textbox(Inches(x), Inches(1), Inches(2), Inches(1))
    frame = shape.text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
    run = frame.paragraphs[0].add_run()
    run.text = "Проверяемый текст"
    run.font.size = Pt(24)
    run.font.name = "DejaVu Sans"
    run.font.color.rgb = RGBColor.from_string(colour)
    return shape


def _panel(shapes: Any, colour: str = "142532", *, width: float = 2) -> Any:
    shape = shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(width), Inches(1)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(colour)
    shape.line.fill.background()
    # Make the synthetic paint explicit, rather than accidentally relying on a theme effect.
    shape._element.find("{*}style/{*}effectRef").set("idx", "0")
    return shape


def _deck(image: Image.Image | None = None, colour: str = "FFFFFF") -> tuple[Any, Any, Any, Any]:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13), Inches(7)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = _picture(slide.shapes, image if image is not None else _image())
    text = _text(slide.shapes, colour)
    return deck, slide, picture, text


def _save(deck: Any, tmp_path: Path) -> Path:
    path = tmp_path / "synthetic.pptx"
    deck.save(str(path))
    return path


def _measured(findings: list[dict[str, Any]]) -> str:
    assert len(findings) == 1
    AuditFinding.model_validate(findings[0])
    assert findings[0]["severity"] == "major"
    assert findings[0]["proposedPatch"] is None
    return findings[0]["evidence"]["measured"]


def test_file_check_is_registered_without_claiming_a_slide_ir_pass() -> None:
    assert {check.id for check in PPTX_CHECKS} == {"image-text-readability"}
    assert check_by_id("image-text-readability") in PPTX_CHECKS
    assert not {c.id for c in PPTX_CHECKS} & {c.id for c in DETERMINISTIC_CHECKS}


@pytest.mark.parametrize(("image", "colour"), [("white", "FFFFFF"), ("black", "000000")])
def test_same_colour_text_fails_on_actual_media(tmp_path: Path, image: str, colour: str) -> None:
    deck, _, _, _ = _deck(_image(image), colour)
    findings = audit_image_readability(_save(deck, tmp_path))
    assert "status=fail; minContrastRatio=1.000" in _measured(findings)


@pytest.mark.parametrize(("image", "colour"), [("black", "F3F9FF"), ("white", "183245")])
def test_contrasting_custom_colours_pass(tmp_path: Path, image: str, colour: str) -> None:
    deck, _, _, _ = _deck(_image(image), colour)
    assert audit_image_readability(_save(deck, tmp_path)) == []


def test_no_fill_textbox_without_geometry_samples_the_picture(tmp_path: Path) -> None:
    deck, _, _, text = _deck(_image("black"))
    props = text._element.find("{*}spPr")
    props.remove(props.find("{*}prstGeom"))
    assert audit_image_readability(_save(deck, tmp_path)) == []
    text.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
    assert "status=fail" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_local_hotspot_fails_despite_dark_whole_frame(tmp_path: Path) -> None:
    image = _image("black")
    ImageDraw.Draw(image).rectangle((50, 45, 60, 55), fill="white")
    deck, _, _, _ = _deck(image)
    assert "luminanceRange=1.000" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_bright_pixels_outside_the_text_region_do_not_count(tmp_path: Path) -> None:
    image = _image("white")
    ImageDraw.Draw(image).rectangle((39, 39, 121, 81), fill="black")
    deck, _, _, _ = _deck(image)
    assert audit_image_readability(_save(deck, tmp_path)) == []


def test_heterogeneity_is_measured_even_when_minimum_contrast_passes(tmp_path: Path) -> None:
    image = _image("white")
    ImageDraw.Draw(image).rectangle((40, 40, 80, 80), fill="#888888")
    deck, _, _, _ = _deck(image, "000000")
    path = _save(deck, tmp_path)
    assert "status=fail" in _measured(audit_image_readability(path))
    assert (
        audit_image_readability(path, profile={"image-text-readability": {"maxLuminanceRange": 1}})
        == []
    )


def test_complete_existing_panel_protects_but_partial_panel_does_not(tmp_path: Path) -> None:
    deck, slide, _, text = _deck()
    panel = _panel(slide.shapes)
    text._element.addprevious(panel._element)
    assert audit_image_readability(_save(deck, tmp_path)) == []
    panel.width = Inches(1)
    assert "status=fail" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_opaque_panel_on_wrong_side_of_picture_is_not_protection(tmp_path: Path) -> None:
    deck, slide, picture, _ = _deck()
    panel = _panel(slide.shapes)
    picture._element.addprevious(panel._element)
    assert "status=fail" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_own_rectangular_fill_protects_text(tmp_path: Path) -> None:
    deck, _, _, text = _deck()
    text.fill.solid()
    text.fill.fore_color.rgb = RGBColor.from_string("243142")
    assert audit_image_readability(_save(deck, tmp_path)) == []


def test_round_panel_is_unknown_rather_than_protection(tmp_path: Path) -> None:
    deck, slide, _, text = _deck()
    panel = _panel(slide.shapes)
    panel._element.find("{*}spPr/{*}prstGeom").set("prst", "roundRect")
    text._element.addprevious(panel._element)
    assert "non-rectangular" in _measured(audit_image_readability(_save(deck, tmp_path)))


@pytest.mark.parametrize("feature", ["alpha", "theme", "effect", "theme-effect"])
def test_unsupported_protection_is_explicit_unknown(tmp_path: Path, feature: str) -> None:
    deck, slide, _, text = _deck()
    panel = _panel(slide.shapes)
    text._element.addprevious(panel._element)
    colour = panel._element.find("{*}spPr/{*}solidFill/{*}srgbClr")
    if feature == "alpha":
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "99999")
        colour.append(alpha)
    elif feature == "theme":
        scheme = OxmlElement("a:schemeClr")
        scheme.set("val", "accent1")
        colour.getparent().replace(colour, scheme)
    elif feature == "effect":
        panel._element.find("{*}spPr").append(OxmlElement("a:effectLst"))
    else:
        panel._element.find("{*}style/{*}effectRef").set("idx", "2")
    assert "status=unknown; reason=" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_gradient_protection_is_unknown_not_a_pass(tmp_path: Path) -> None:
    deck, slide, _, text = _deck()
    panel = _panel(slide.shapes)
    text._element.addprevious(panel._element)
    solid = panel._element.find("{*}spPr/{*}solidFill")
    gradient = OxmlElement("a:gradFill")
    solid.getparent().replace(solid, gradient)
    measured = _measured(audit_image_readability(_save(deck, tmp_path)))
    assert "status=unknown" in measured
    assert "status=fail" not in measured
    assert "status=measured" not in measured


def test_image_crop_changes_the_pixels_behind_text(tmp_path: Path) -> None:
    image = _image("black")
    ImageDraw.Draw(image).rectangle((160, 0, 319, 159), fill="white")
    deck, _, picture, _ = _deck(image)
    picture.crop_right = 0.5
    assert audit_image_readability(_save(deck, tmp_path)) == []
    picture.crop_right, picture.crop_left = 0, 0.5
    assert "status=fail" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_topmost_picture_and_foreground_occlusion_are_respected(tmp_path: Path) -> None:
    deck, slide, _, text = _deck()
    black = _picture(slide.shapes, _image("black"))
    text._element.addprevious(black._element)
    assert audit_image_readability(_save(deck, tmp_path)) == []
    text._element.addnext(black._element)
    assert "foreground picture occludes text" in _measured(
        audit_image_readability(_save(deck, tmp_path))
    )


@pytest.mark.parametrize("feature", ["alpha", "blip-effect", "rotation", "negative-crop"])
def test_unsupported_picture_is_unknown(tmp_path: Path, feature: str) -> None:
    image = Image.new("RGBA", (320, 160), (0, 0, 0, 128)) if feature == "alpha" else _image("black")
    deck, _, picture, _ = _deck(image)
    if feature == "blip-effect":
        picture._element.find("{*}blipFill/{*}blip").append(OxmlElement("a:alphaModFix"))
    elif feature == "rotation":
        picture.rotation = 25
    elif feature == "negative-crop":
        picture.crop_left = -0.1
    assert "status=unknown; reason=" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_existing_opaque_panel_can_hide_unsupported_picture(tmp_path: Path) -> None:
    deck, slide, picture, text = _deck(Image.new("RGBA", (320, 160), (255, 255, 255, 128)))
    picture._element.find("{*}blipFill/{*}blip").append(OxmlElement("a:alphaModFix"))
    panel = _panel(slide.shapes)
    text._element.addprevious(panel._element)
    assert audit_image_readability(_save(deck, tmp_path)) == []


def test_growing_text_frame_is_unknown_not_a_silent_pass(tmp_path: Path) -> None:
    deck, _, _, text = _deck(_image("black"))
    text.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    assert "frame can grow" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_group_scale_translation_and_rotation(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _picture(group.shapes, _image("black"))
    _text(group.shapes)
    group.left, group.top, group.width, group.height = (
        Inches(0.7),
        Inches(0.4),
        Inches(6),
        Inches(3),
    )
    assert audit_image_readability(_save(deck, tmp_path)) == []
    group.rotation = 12
    assert "group" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_only_selected_editable_text_is_audited(tmp_path: Path) -> None:
    deck, slide, _, editable = _deck()
    other = _text(slide.shapes, x=5)
    path = _save(deck, tmp_path)
    assert len(audit_image_readability(path)) == 2
    found = audit_image_readability(path, editable_text_shape_ids=[[str(editable.shape_id)]])
    assert len(found) == 1
    assert found[0]["target"]["elementId"] == f"slide-{editable.shape_id}"
    assert f"slide-{other.shape_id}" != found[0]["target"]["elementId"]
    assert audit_image_readability(path, editable_text_shape_ids=[[]]) == []
    with pytest.raises(ValueError, match="per slide"):
        audit_image_readability(path, editable_text_shape_ids=[])
    with pytest.raises(ValueError, match="Unknown editable"):
        audit_image_readability(path, editable_text_shape_ids=[["99999"]])


def test_theme_text_colour_is_unknown(tmp_path: Path) -> None:
    deck, _, _, text = _deck(_image("black"))
    fill = text.text_frame.paragraphs[0].runs[0]._r.find("{*}rPr/{*}solidFill")
    fill.remove(fill[0])
    scheme = OxmlElement("a:schemeClr")
    scheme.set("val", "accent2")
    fill.append(scheme)
    assert "theme/system colour" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_picture_background_fill_is_unknown_not_skipped(tmp_path: Path) -> None:
    deck, slide, picture, _ = _deck()
    slide.shapes._spTree.remove(picture._element)
    background = slide.background._cSld.get_or_add_bg().get_or_add_bgPr()
    background.append(OxmlElement("a:blipFill"))
    assert "unresolved slide background" in _measured(
        audit_image_readability(_save(deck, tmp_path))
    )


def test_group_effects_are_unknown(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _picture(group.shapes, _image("black"))
    _text(group.shapes)
    group._element.find("{*}grpSpPr").append(OxmlElement("a:effectLst"))
    assert "group effects" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_media_budget_is_unknown(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    deck, _, _, _ = _deck(_image("black"))
    monkeypatch.setattr(image_readability, "_MAX_REGION_PIXELS", 1)
    assert "pixel budget" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_master_media_is_also_part_of_the_actual_background(tmp_path: Path) -> None:
    deck, slide, picture, _ = _deck(_image("black"))
    slide.shapes._spTree.remove(picture._element)
    master = deck.slide_masters[0]
    for colour, passes in (("black", True), ("white", False)):
        media = tmp_path / f"{colour}.png"
        _image(colour).save(media)
        part, rid = cast(BaseSlidePart, master.part).get_or_add_image_part(str(media))
        master.shapes._spTree.add_pic(
            master.shapes._next_shape_id,
            colour,
            part.desc,
            rid,
            0,
            0,
            Inches(8),
            Inches(4),
        )
        findings = audit_image_readability(_save(deck, tmp_path))
        if passes:
            assert findings == []
        else:
            assert "status=fail" in _measured(findings)


def test_text_warp_is_unknown(tmp_path: Path) -> None:
    deck, _, _, text = _deck(_image("black"))
    text.text_frame._txBody.bodyPr.append(OxmlElement("a:prstTxWarp"))
    assert "text warp" in _measured(audit_image_readability(_save(deck, tmp_path)))


def test_stable_findings_custom_slide_id_variant_and_unchanged_file(tmp_path: Path) -> None:
    deck, _, _, _ = _deck()
    path = _save(deck, tmp_path)
    before = path.read_bytes()
    first = audit_image_readability(path, variant="airy", slide_ids=["custom"])
    assert first == audit_image_readability(path, variant="airy", slide_ids=["custom"])
    assert first[0]["target"]["variant"] == "airy"
    assert first[0]["target"]["slideId"] == "custom"
    assert path.read_bytes() == before


def test_no_image_and_profile_disabled_are_not_applicable(tmp_path: Path) -> None:
    deck, slide, picture, _ = _deck()
    slide.shapes._spTree.remove(picture._element)
    assert audit_image_readability(_save(deck, tmp_path)) == []
    _picture(slide.shapes, _image())
    assert (
        audit_image_readability(
            _save(deck, tmp_path), profile={"image-text-readability": {"enabled": False}}
        )
        == []
    )


def test_exported_audit_includes_file_findings_in_summary(tmp_path: Path) -> None:
    deck, _, _, text = _deck()
    path = _save(deck, tmp_path)
    fixture = Path(__file__).resolve().parents[2] / "packages/contracts/fixtures/DesignSystem.json"
    design_system = json.loads(fixture.read_text("utf-8"))["valid"][0]["payload"]
    request = {
        "deckId": "synthetic",
        "designSystem": design_system,
        "slides": [{"slideId": "s1", "variant": "dense", "index": 0, "elements": []}],
    }
    profile = {check.id: {"enabled": False} for check in DETERMINISTIC_CHECKS}
    assert audit_deck(request, "custom", profile)["passed"]
    report = audit_exported_deck(
        path, request, "custom", profile, editable_text_shape_ids=[[str(text.shape_id)]]
    )
    AuditReport.model_validate(report)
    assert report["passed"] is False
    assert report["score"] == 90
    assert report["counts"]["major"] == 1
    template = next(c for c in report["categories"] if c["category"] == "template")
    assert template == {"category": "template", "checks": 1, "findings": 1, "passed": False}
    assert "pptxPath" not in request
    text.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    unknown = audit_exported_deck(_save(deck, tmp_path), request, "custom", profile)
    assert unknown["passed"] is False
    assert "status=unknown" in _measured(unknown["findings"])


def test_saved_deck_render_shows_the_same_existing_panel(tmp_path: Path) -> None:
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    windows_office = Path("C:/Program Files/LibreOffice/program/soffice.exe")
    if soffice is None and windows_office.is_file():
        soffice = str(windows_office)
    raster = shutil.which("pdftoppm")
    if soffice is None or raster is None:
        pytest.skip("LibreOffice and pdftoppm required for visual evidence")
    deck, _, _, _ = _deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _picture(slide.shapes, _image())
    _panel(slide.shapes, "102030")
    _text(slide.shapes)
    path = _save(deck, tmp_path)
    findings = audit_image_readability(path)
    assert len(findings) == 1 and findings[0]["target"]["slideId"] == "s1"
    subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={(tmp_path / 'lo-profile').as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=60,
    )
    subprocess.run(
        [raster, "-png", "-r", "72", str(path.with_suffix(".pdf")), str(tmp_path / "slide")],
        check=True,
        capture_output=True,
        timeout=30,
    )
    for index, expected in ((1, (255, 255, 255)), (2, (16, 32, 48))):
        with Image.open(tmp_path / f"slide-{index}.png") as rendered:
            # A point inside the protected region, clear of glyphs and rectangle edges.
            observed = rendered.convert("RGB").getpixel((74, 137))
            assert isinstance(observed, tuple)
            assert all(abs(a - b) <= 2 for a, b in zip(observed, expected, strict=True))
    with Image.open(tmp_path / "slide-2.png") as rendered:
        crop = rendered.convert("RGB").crop((75, 74, 211, 136))
        assert max(crop.getchannel("R").tobytes()) >= 240  # white text survived export


@pytest.mark.parametrize("safe_inset", [False, True])
def test_picture_outside_reserved_text_region_needs_no_pixel_audit(
    tmp_path: Path, safe_inset: bool
) -> None:
    deck, _, _, text = _deck()
    text.left = Inches(7)
    if safe_inset:
        text.text_frame.margin_left = Inches(1.2)
    findings = audit_image_readability(_save(deck, tmp_path))
    if safe_inset:
        assert not findings
    else:
        assert findings  # The inner text region still intersects the photo.
