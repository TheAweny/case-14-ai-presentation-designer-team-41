"""Layout selection: slides are built on the template's patterns, per variant."""

from pathlib import Path
from typing import Any

from engine.contracts import DeckPlan
from engine.layout import build_slide_ir
from engine.parser import template_profile_of
from pptx import Presentation


def _profile(tmp_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    path = tmp_path / "t.pptx"
    Presentation().save(str(path))
    return template_profile_of(path)


def _plan(*slides: dict[str, Any]) -> DeckPlan:
    return DeckPlan.model_validate(
        {
            "schemaVersion": 1,
            "planId": "plan",
            "briefId": "brief",
            "title": "Deck",
            "language": "en",
            "slides": [
                {
                    "slideId": f"s{i}",
                    "headline": "Revenue grew because retention held",
                    "subhead": "",
                    "speakerNotes": "",
                    "sources": [],
                    **slide,
                }
                for i, slide in enumerate(slides)
            ],
        }
    )


def _pattern(patterns: dict[str, Any], pattern_id: str) -> dict[str, Any]:
    return next(p for p in patterns["patterns"] if p["patternId"] == pattern_id)


def test_each_slide_is_built_on_a_pattern_serving_its_intent(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan(
        {"intent": "title", "blocks": []},
        {"intent": "two-column", "blocks": [{"kind": "bullets", "items": ["a", "b", "c", "d"]}]},
    )

    slides = build_slide_ir(plan, design_system, patterns)

    for planned, built in zip(plan.slides, slides, strict=True):
        assert planned.intent in _pattern(patterns, built.patternId)["intents"]


def test_variants_differ_in_density(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan({"intent": "bullets", "blocks": [{"kind": "paragraph", "text": "Short."}]})

    dense = _pattern(patterns, build_slide_ir(plan, design_system, patterns, "dense")[0].patternId)
    airy = _pattern(patterns, build_slide_ir(plan, design_system, patterns, "airy")[0].patternId)

    assert dense["density"] > airy["density"]


def test_one_list_on_a_two_column_slide_is_split_between_columns(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan(
        {"intent": "two-column", "blocks": [{"kind": "bullets", "items": ["a", "b", "c", "d"]}]}
    )

    elements = build_slide_ir(plan, design_system, patterns)[0].model_dump()["elements"]

    bodies = [e for e in elements if e["role"] == "body"]
    assert len(bodies) >= 2
    texts = [[p["runs"][0]["text"] for p in e["paragraphs"]] for e in bodies]
    assert texts[:2] == [["a", "b"], ["c", "d"]]


def test_text_that_does_not_fit_raises_an_overflow_signal(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    long_text = "Retention held while revenue grew. " * 200
    plan = _plan({"intent": "statement", "blocks": [{"kind": "paragraph", "text": long_text}]})

    elements = build_slide_ir(plan, design_system, patterns)[0].model_dump()["elements"]

    body = next(e for e in elements if e["role"] == "body")
    assert body["overflow"] is not None
    assert body["overflow"]["lines"] > 0
    assert body["paragraphs"][0]["runs"][0]["text"] == long_text


def test_short_text_fits(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan({"intent": "statement", "blocks": [{"kind": "paragraph", "text": "Short."}]})

    elements = build_slide_ir(plan, design_system, patterns)[0].model_dump()["elements"]

    assert all(e["overflow"] is None for e in elements)


def test_subhead_goes_to_the_subtitle_slot_when_there_is_one(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan({"intent": "title", "subhead": "Q3 review", "blocks": []})

    elements = build_slide_ir(plan, design_system, patterns)[0].model_dump()["elements"]

    subtitle = next(e for e in elements if e["role"] == "subtitle")
    assert subtitle["paragraphs"][0]["runs"][0]["text"] == "Q3 review"


def test_running_chrome_is_left_to_the_template(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan({"intent": "bullets", "blocks": [{"kind": "paragraph", "text": "Short."}]})

    elements = build_slide_ir(plan, design_system, patterns)[0].model_dump()["elements"]

    assert {e["role"] for e in elements} <= {"title", "subtitle", "body"}


def test_layout_is_deterministic(tmp_path: Path) -> None:
    design_system, patterns = _profile(tmp_path)
    plan = _plan({"intent": "bullets", "blocks": [{"kind": "bullets", "items": ["a", "b"]}]})

    first = [s.model_dump() for s in build_slide_ir(plan, design_system, patterns)]
    second = [s.model_dump() for s in build_slide_ir(plan, design_system, patterns)]

    assert first == second


def test_text_is_placed_on_a_template_whose_layouts_only_have_titles(tmp_path: Path) -> None:
    deck = Presentation()
    for layout in list(deck.slide_layouts):
        if layout.name not in {"Title Only", "Blank"}:
            deck.slide_layouts.remove(layout)
    path = tmp_path / "titles-only.pptx"
    deck.save(str(path))
    design_system, patterns = template_profile_of(path)
    assert any(_has_body(p) for p in patterns["patterns"])

    plan = _plan({"intent": "statement", "blocks": [{"kind": "paragraph", "text": "Keep me"}]})
    (slide,) = build_slide_ir(plan, design_system, patterns)
    texts = [
        run["text"]
        for element in slide.model_dump(mode="json")["elements"]
        for paragraph in element.get("paragraphs") or []
        for run in paragraph["runs"]
    ]
    assert "Keep me" in " ".join(texts)


def _has_body(pattern: dict[str, Any]) -> bool:
    return any(slot["role"] == "body" for slot in pattern["slots"])
