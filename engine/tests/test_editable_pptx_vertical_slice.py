"""End-to-end smoke coverage for the minimal native editable PPTX path."""

import json
import shutil
import subprocess
from pathlib import Path
from typing import cast

from engine.api.app import app
from engine.layout import build_slide_ir
from engine.parser import design_system_of
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape
from pptx.util import Emu, Inches, Pt

FIXTURE_CONTENT = Path(__file__).parent / "fixtures" / "vertical-slice-content.json"


def build_template(destination: Path) -> Path:
    """Build a synthetic design fixture without encoding a product template name."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(13.333), Inches(0.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(17, 34, 51)
    accent.line.fill.background()
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "Template typography sample"
    run.font.name = "Aptos Display"
    run.font.size = Pt(30)
    run.font.bold = True
    presentation.save(str(destination))
    return destination


def test_parser_extracts_a_design_system_from_template_evidence(tmp_path: Path) -> None:
    design_system = design_system_of(build_template(tmp_path / "unknown-name.pptx"))

    assert design_system["grid"]["slideCx"] > 0
    assert design_system["palette"]
    assert design_system["typeScale"]


def test_api_returns_real_profile_for_an_uploaded_template_path(tmp_path: Path) -> None:
    template_path = build_template(tmp_path / "uploaded-template.pptx")

    response = TestClient(app).post(
        "/v1/template-profile",
        json={"template_path": str(template_path)},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["designSystem"]["palette"]
    assert payload["designSystem"]["typeScale"]
    assert payload["templateIr"]["slides"]
    assert len(payload["templateIr"]["hash"]) == 64
    if shutil.which("soffice") and shutil.which("pdftoppm"):
        assert payload["previewDataUrl"].startswith("data:image/png;base64,")


def test_api_returns_a_readable_error_for_a_non_presentation(tmp_path: Path) -> None:
    source = tmp_path / "not-a-presentation.pptx"
    source.write_bytes(b"not a zip")

    response = TestClient(app).post(
        "/v1/template-profile",
        json={"template_path": str(source)},
    )

    assert response.status_code == 422
    assert "OOXML" in response.json()["detail"]


def test_layout_builds_native_text_objects_from_content(tmp_path: Path) -> None:
    content = json.loads(FIXTURE_CONTENT.read_text(encoding="utf-8"))
    design_system = design_system_of(build_template(tmp_path / "layout-source.pptx"))
    from engine.contracts import DeckPlan

    plan = DeckPlan.model_validate(
        {
            "schemaVersion": 1,
            "planId": "plan-vertical-slice",
            "briefId": content["runId"],
            "title": content["slides"][0]["title"],
            "language": "en",
            "slides": [
                {
                    "slideId": item["slideId"],
                    "headline": item["title"],
                    "subhead": "",
                    "intent": "statement",
                    "blocks": [{"kind": "paragraph", "text": item["body"]}],
                    "speakerNotes": "",
                    "sources": [],
                }
                for item in content["slides"]
            ],
        }
    )

    slide_ir = build_slide_ir(plan, design_system)

    elements = slide_ir[0].model_dump()["elements"]
    texts = [
        paragraph["runs"][0]["text"]
        for element in elements
        for paragraph in element.get("paragraphs", [])
    ]
    assert texts == [
        "Growth overview",
        "Revenue increased by 28 percent while retention remained stable.",
    ]


def test_api_exports_an_editable_pptx_that_reopens_with_native_text_boxes(tmp_path: Path) -> None:
    template_path = build_template(tmp_path / "fixture-template.pptx")
    content = json.loads(FIXTURE_CONTENT.read_text(encoding="utf-8"))
    output_path = tmp_path / "output" / "vertical-slice.pptx"

    response = TestClient(app).post(
        "/v1/editable-decks",
        json={
            **content,
            "schemaVersion": 1,
            "templatePath": str(template_path),
            "outputPath": str(output_path),
        },
    )

    assert response.status_code == 200
    assert output_path.is_file()
    exported = Presentation(str(output_path))
    assert len(exported.slides) == 1
    text_shapes = [
        cast(Shape, shape)
        for shape in exported.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.has_text_frame
    ]
    assert [shape.text for shape in text_shapes] == [
        "Growth overview",
        "Revenue increased by 28 percent while retention remained stable.",
    ]
    assert response.json()["slideCount"] == 1


def test_api_rejects_a_missing_template(tmp_path: Path) -> None:
    response = TestClient(app).post(
        "/v1/editable-decks",
        json={
            "schemaVersion": 1,
            "runId": "run-missing-template",
            "templatePath": str(tmp_path / "missing.pptx"),
            "outputPath": str(tmp_path / "output.pptx"),
            "slides": [{"slideId": "slide-1", "title": "Title", "body": "Body"}],
        },
    )

    assert response.status_code == 404


def test_libreoffice_converts_the_exported_pptx(tmp_path: Path) -> None:
    template_path = build_template(tmp_path / "libreoffice-template.pptx")
    output_path = tmp_path / "vertical-slice.pptx"
    response = TestClient(app).post(
        "/v1/editable-decks",
        json={
            "schemaVersion": 1,
            "runId": "run-libreoffice-smoke",
            "templatePath": str(template_path),
            "outputPath": str(output_path),
            "slides": [{"slideId": "slide-1", "title": "Title", "body": "Body"}],
        },
    )

    assert response.status_code == 200
    soffice = shutil.which("soffice")
    if soffice is None:
        return
    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(output_path),
        ],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert output_path.with_suffix(".pdf").is_file()
